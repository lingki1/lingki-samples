# 核对「发票整理助手.html」的一次真实运行结果：拿「样例发票/标准答案.json」逐格比。
#
# 先自己用发票整理助手整理一遍样例发票（电脑上选「样例发票」文件夹，它会在里面新建一个「整理好的发票_日期_时间」），然后：
#     python check_run.py <发票文件夹> <整理好的发票_… 文件夹>
# 例如：
#     python check_run.py ../样例发票 "../样例发票/整理好的发票_2026-10-10_230914"
# 手机那条路下载到的压缩包，解开以后也可以这样核。
#
# 依赖：Python 3.8 或更新的版本，加 openpyxl（pip install openpyxl）。
import sys, io, json, os, hashlib, re, collections
import openpyxl
if len(sys.argv) < 3: sys.exit('用法：python check_run.py <发票文件夹> <整理好的发票_… 文件夹>')
src, out = sys.argv[1], sys.argv[2]
SAMPLES = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '样例发票')     # 仓库里的那份样例和标准答案
T = json.load(io.open(os.path.join(SAMPLES, '标准答案.json'), encoding='utf-8'))
truth = {t['file']: t for t in T['items']}
sha = lambda p: hashlib.sha256(open(p, 'rb').read()).hexdigest()
problems = []; say = lambda *a: problems.append(' '.join(str(x) for x in a))

# 1. 原件有没有被动过：和仓库里的那一份逐个比指纹
orig_dir = SAMPLES
src_files = sorted(f for f in os.listdir(src) if f.lower().endswith('.pdf'))
same = sum(1 for f in src_files if os.path.exists(os.path.join(orig_dir, f)) and sha(os.path.join(src, f)) == sha(os.path.join(orig_dir, f)))
print(f'原件：{len(src_files)} 个 PDF，和仓库里那份一模一样的 {same} 个；标准答案里有 {len(truth)} 个')
if set(src_files) != set(truth): say('原件文件名和标准答案不一致', set(src_files) ^ set(truth))

# 2. 输出文件夹：每个文件按内容指纹对回原件，再看它被放到哪、叫什么
by_hash = collections.defaultdict(list)
for f in src_files: by_hash[sha(os.path.join(src, f))].append(f)
placed = collections.defaultdict(list)   # 原文件名(一组) -> [(folder, newname)]
n_out = 0
for root, _, files in os.walk(out):
    for f in files:
        if f.endswith('.xlsx'): continue
        n_out += 1; h = sha(os.path.join(root, f)); folder = os.path.relpath(root, out).replace(os.sep, '/')
        if h not in by_hash: say('输出里有个文件对不回任何原件', folder, f); continue
        placed[h].append((folder, f))
print(f'输出：{n_out} 个文件（不算表格）')
folders = collections.Counter()
for h, names in by_hash.items():
    kinds = [truth[n]['kind'] for n in names]; got = placed.get(h, [])
    if len(got) != len(names): say('份数不对', names, '→', got)
    for folder, f in got: folders[folder] += 1
    t = truth[names[0]]
    if 'image-only' in kinds:
        if [g[0] for g in got] != ['需要人工检查']: say('纯图片那张没进「需要人工检查」', got)
    elif 'not-invoice' in kinds:
        if [g[0] for g in got] != ['不是发票']: say('行程单没进「不是发票」', got)
    else:
        dup_n = sum(1 for g in got if g[0] == '重复发票'); want_dup = len(names) - 1
        if dup_n != want_dup: say('重复的份数不对', names, got)
        want = f"{t['date']}_{t['seller']}_{t['total']:.2f}元"
        for folder, f in got:
            if not f.startswith(want): say('改的名字不对', names[0], '应以', want, '开头，实际', f)
print('各文件夹：', dict(folders))

# 3. 汇总表
wb = openpyxl.load_workbook(os.path.join(out, '发票汇总表.xlsx')); ws = wb['发票明细']; rows = list(ws.iter_rows(values_only=True)); head = rows[0]; col = {h: i for i, h in enumerate(head)}
print('表头：', head)
status = collections.Counter(); total_ok = 0.0; seen = set(); net_ok = tax_ok = 0.0
for r in rows[1:]:
    name = r[col['原文件名']]
    if not name:   # 合计行
        tot_row = r; continue
    seen.add(name); st = r[col['状态']]; status[st] += 1; t = truth.get(name)
    if not t: say('表里有个文件名不在标准答案里', name); continue
    if t['kind'] == 'not-invoice':
        if st != '不是发票': say('行程单的状态', st)
        continue
    if t['kind'] == 'image-only':
        if st != '需检查': say('纯图片那张的状态', st, r)
        if r[col['价税合计']] not in (None, ''): say('纯图片那张被填了金额（猜的？）', r[col['价税合计']])
        continue
    for k, tk in [('开票日期', 'date'), ('销售方', 'seller'), ('发票号码', 'number'), ('购买方', 'buyer')]:
        if str(r[col[k]] or '') != str(t[tk]): say(name, k, '表里', r[col[k]], '应为', t[tk])
    if abs((r[col['价税合计']] or 0) - t['total']) > 0.005: say(name, '价税合计 表里', r[col['价税合计']], '应为', t['total'])
    if r[col['税额']] is None or abs(r[col['税额']] - t['tax']) > 0.005: say(name, '税额 表里', r[col['税额']], '应为', t['tax'])
    if r[col['金额(不含税)']] is None or abs(r[col['金额(不含税)']] + r[col['税额']] - t['total']) > 0.011: say(name, '不含税金额加税额对不上合计', r[col['金额(不含税)']], r[col['税额']])
    if str(r[col['项目']] or '') != t['item'].split('*')[-1]: say(name, '项目 表里', r[col['项目']], '应为', t['item'])
    if st == '正常': total_ok += r[col['价税合计']]; net_ok += r[col['金额(不含税)']] or 0; tax_ok += r[col['税额']] or 0
    elif st != '重复': say(name, '状态是', st)
missing = set(truth) - seen
if missing: say('表里漏了', missing)
print('状态：', dict(status)); print(f'正常发票价税合计加起来 {total_ok:.2f}；标准答案 {T["total"]:.2f}')
print('表里的合计行：', [x for x in tot_row if x not in (None, '')])
if abs(total_ok - T['total']) > 0.005: say('合计对不上')
if abs((tot_row[col['价税合计']] or 0) - T['total']) > 0.005: say('合计行的数不对', tot_row[col['价税合计']])
if status != collections.Counter({'正常': 80, '重复': 3, '需检查': 1, '不是发票': 1}): say('状态的数量不对', dict(status))
ws2 = wb['分类汇总']; print('分类汇总：'); [print('  ', [x for x in r if x not in (None, '')]) for r in ws2.iter_rows(values_only=True) if any(x not in (None, '') for x in r)]
print('\n问题：' + ('没有' if not problems else '')); [print(' -', p) for p in problems]
