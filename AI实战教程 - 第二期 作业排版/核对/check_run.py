# 拿标准答案核那场对话做出来的软件跑出的结果
# 用法（在这一期的文件夹里）：python 核对/check_run.py <输出文件夹> [--renumber] [--skip 不要的那份的文件名] [--project 装着「学生作业」的文件夹] [--json out.json]
# 例：python 核对/check_run.py 排版结果/2026-10-11_211626 --renumber --skip "作业_梁思齐_20240324.docx"
# 需要 python-docx。标准答案在 ../样例/truth.json；--project 不写就用 ../样例（里面有「学生作业」）。
# 放进仓库时只改了这几行说明和下面两处路径，别的是核对时用的原样。
# 核四样：原件一个没动；每个学生有没有对应的输出、文件名对不对；文字有没有丢或者多；格式是不是「格式要求」写的那样。
# 格式按「实际生效」算：字上没写就看段落样式，再看样式的上一级，再看文档默认值。
import sys, os, io, json, hashlib, re, zipfile
from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__)); SAMPLES = os.path.join(HERE, '..', '样例'); T = json.load(io.open(os.path.join(SAMPLES, 'truth.json'), encoding='utf-8'))
a = sys.argv[1:]; OUT = a[0]; PROJ = a[a.index('--project') + 1] if '--project' in a else SAMPLES
CLASS = '高二（3）班'

def chain(doc, style_id):
    """样式和它的上级，一层层往上。"""
    out = []; st = {s.get(qn('w:styleId')): s for s in doc.styles.element.findall(qn('w:style'))}
    while style_id and style_id in st and len(out) < 12:
        out.append(st[style_id]); b = st[style_id].find(qn('w:basedOn')); style_id = b.get(qn('w:val')) if b is not None else None
    return out

def default_para_style(doc):
    for s in doc.styles.element.findall(qn('w:style')):
        if s.get(qn('w:type')) == 'paragraph' and s.get(qn('w:default')) == '1': return s.get(qn('w:styleId'))

def eff(doc, p, r, path, attr='val', para=False):
    """取实际生效的属性。path 例：'w:sz'、'w:rFonts'；para=True 取段落属性（w:jc、w:ind、w:spacing）。"""
    tag = 'w:pPr' if para else 'w:rPr'; srcs = []
    if not para and r is not None and r._element.find(qn('w:rPr')) is not None: srcs.append(r._element.find(qn('w:rPr')))
    if para and p._element.find(qn('w:pPr')) is not None: srcs.append(p._element.find(qn('w:pPr')))
    ppr = p._element.find(qn('w:pPr')); ps = ppr.find(qn('w:pStyle')) if ppr is not None else None
    for s in chain(doc, ps.get(qn('w:val')) if ps is not None else default_para_style(doc)):
        if s.find(qn(tag)) is not None: srcs.append(s.find(qn(tag)))
    dd = doc.styles.element.find(qn('w:docDefaults'))
    if dd is not None:
        x = dd.find(qn('w:pPrDefault' if para else 'w:rPrDefault'))
        if x is not None and x.find(qn(tag)) is not None: srcs.append(x.find(qn(tag)))
    for s in srcs:
        e = s.find(qn(path))
        if e is not None and e.get(qn('w:' + attr)) is not None: return e.get(qn('w:' + attr))
    return None

def flag(doc, p, r, path):
    """加粗、下划线这类开关：有标签且不是 0/false/none 就算开。"""
    for getter in (lambda: r._element.find(qn('w:rPr')),):
        e = getter()
        if e is not None and e.find(qn(path)) is not None:
            v = e.find(qn(path)).get(qn('w:val')); return v not in ('0', 'false', 'none')
    ppr = p._element.find(qn('w:pPr')); ps = ppr.find(qn('w:pStyle')) if ppr is not None else None
    for s in chain(doc, ps.get(qn('w:val')) if ps is not None else default_para_style(doc)):
        e = s.find(qn('w:rPr'))
        if e is not None and e.find(qn(path)) is not None:
            v = e.find(qn(path)).get(qn('w:val')); return v not in ('0', 'false', 'none')
    return False

def para_checks(doc, p, role):
    """一个段落按它的角色该是什么样；返回不对的地方。"""
    bad = []; runs = [r for r in p.runs if r.text.strip()]
    want = {'title': ('黑体', 44, 'center'), 'info': ('楷体', 24, 'center'), 'heading': ('黑体', 28, 'left'), 'body': ('宋体', 24, 'both'), 'ref': ('宋体', 24, 'both')}[role]
    for r in runs:
        ea = eff(doc, p, r, 'w:rFonts', 'eastAsia'); sz = eff(doc, p, r, 'w:sz'); col = eff(doc, p, r, 'w:color')
        if ea != want[0] and not (want[0] == '楷体' and ea in ('楷体_GB2312', 'KaiTi')) and not (want[0] == '黑体' and ea == 'SimHei') and not (want[0] == '宋体' and ea == 'SimSun'): bad.append(f'中文字体 {ea}≠{want[0]}')
        if str(sz) != str(want[1]): bad.append(f'字号 {int(sz) / 2 if sz else None}≠{want[1] / 2}')
        if col not in (None, 'auto', '000000'): bad.append(f'颜色 {col}')
        if role in ('body', 'ref') and re.search(r'[A-Za-z0-9]', r.text) and eff(doc, p, r, 'w:rFonts', 'ascii') != 'Times New Roman': bad.append(f"西文字体 {eff(doc, p, r, 'w:rFonts', 'ascii')}")
        if role in ('title', 'heading') and flag(doc, p, r, 'w:b'): bad.append('加粗了')
        if role == 'title' and flag(doc, p, r, 'w:u'): bad.append('有下划线')
        if eff(doc, p, r, 'w:highlight') not in (None, 'none'): bad.append('有底色')
    jc = eff(doc, p, None, 'w:jc', para=True) or 'left'
    if jc == 'start': jc = 'left'
    if jc != want[2] and not (want[2] == 'both' and jc == 'distribute'): bad.append(f'对齐 {jc}≠{want[2]}')
    if role == 'body':
        fc = eff(doc, p, None, 'w:ind', 'firstLineChars', para=True); fl = eff(doc, p, None, 'w:ind', 'firstLine', para=True)
        if not (fc == '200' or (fl and abs(int(fl) - 480) <= 24)): bad.append(f'首行缩进 chars={fc} twips={fl}')
        ln = eff(doc, p, None, 'w:spacing', 'line', para=True); rule = eff(doc, p, None, 'w:spacing', 'lineRule', para=True) or 'auto'
        if not (ln == '360' and rule == 'auto'): bad.append(f'行距 {ln}/{rule}')
        if p.text != p.text.strip() or p.text.startswith('\u3000'): bad.append('段首还留着空格')
    if role == 'heading':
        fc = eff(doc, p, None, 'w:ind', 'firstLineChars', para=True); fl = eff(doc, p, None, 'w:ind', 'firstLine', para=True)
        if (fc and fc != '0') or (fl and fl != '0'): bad.append('小标题没顶格')
    return sorted(set(bad))

def norm(s): return re.sub(r'\s+', '', s or '')

def check_doc(path, t):
    doc = Document(path); res = {'bad': [], 'fmt': {}}; sec = doc.sections[0]
    cm = lambda v: round(v / 360000, 2) if v is not None else None
    page = (cm(sec.page_width), cm(sec.page_height), cm(sec.top_margin), cm(sec.bottom_margin), cm(sec.left_margin), cm(sec.right_margin))
    if not (abs(page[0] - 21.0) < .06 and abs(page[1] - 29.7) < .06): res['bad'].append(f'纸张 {page[0]}×{page[1]}')
    if not all(abs(x - y) < .03 for x, y in zip(page[2:], (2.54, 2.54, 3.18, 3.18))): res['bad'].append(f'页边距 {page[2:]}')
    foot = ''.join(x.xml for s in doc.sections for x in [s.footer._element])
    if 'PAGE' not in foot: res['bad'].append('没有页码')
    # 文字：非空段落的顺序和内容
    got = [p for p in doc.paragraphs if p.text.strip()]; want = [x for x in t['paras'] if x['role'] in ('title', 'info', 'heading', 'body', 'ref')]
    if '--renumber' in a:   # 第 4 轮改了需求：小标题编号统一成「一、二、三」，按出现的先后重新编；「参考资料」不编号
        k = 0; w2 = []
        for x in want:
            if x['role'] == 'heading' and not x.get('ref_heading'):
                base = re.sub(r'^(（[一二三四五六七八九十]+）|[一二三四五六七八九十]+、|[0-9]+[.．、] *)', '', x['text']); x = dict(x, text='一二三四五六七八九十'[k] + '、' + base); k += 1
            w2.append(x)
        want = w2
    if t.get('info_where') in ('header', 'none') and t.get('name'):   # 文档里没写的，输出里应该补上一行
        k = 1 if want and want[0]['role'] == 'title' else 0; want = want[:k] + [{'role': 'info', 'text': f"{CLASS}  {t['name']}  {t['id']}", 'added': True}] + want[k:]
    empties = sum(1 for p in doc.paragraphs if not p.text.strip() and not p._element.findall('.//' + qn('w:drawing')))
    if empties: res['bad'].append(f'还有 {empties} 个空段')
    gi = 0; missing = []; role_of = {}
    for w in want:
        target = norm(w['text']) if w['role'] != 'info' else None
        hit = None
        for k in range(gi, len(got)):
            g = norm(got[k].text)
            if (target is not None and g == target) or (w['role'] == 'info' and t.get('name') and t['name'] in got[k].text and t['id'] in got[k].text): hit = k; break
        if hit is None: missing.append((w['role'], w['text'][:24])); continue
        role_of[hit] = w; gi = hit + 1
    extra = [got[k].text[:24] for k in range(len(got)) if k not in role_of]
    if missing: res['bad'].append(f'少了 {len(missing)} 段：{missing[:3]}')
    if extra: res['bad'].append(f'多了 {len(extra)} 段：{extra[:3]}')
    for k, w in role_of.items():
        if w['role'] == 'info' and norm(got[k].text) != norm(f"{CLASS}{t['name']}{t['id']}"): res['fmt'].setdefault('姓名学号那行的写法没统一', []).append(got[k].text)
        for b in para_checks(doc, got[k], w['role']): res['fmt'].setdefault(f"{ {'title': '题目', 'info': '姓名学号', 'heading': '小标题', 'body': '正文', 'ref': '参考资料'}[w['role']] }：{b}", []).append(got[k].text[:16])
    if any(x['role'] == 'table' for x in t['paras']) and not doc.tables: res['bad'].append('表格没了')
    if any(x['role'] == 'image' for x in t['paras']) and not doc.element.findall('.//' + qn('w:drawing')): res['bad'].append('图片没了')
    return res

def main():
    rep = {'originals': [], 'files': {}, 'unmatched_outputs': [], 'no_output': []}
    for t in T:   # 原件
        p = os.path.join(PROJ, '学生作业', t['file'])
        if not os.path.exists(p): rep['originals'].append(('不见了', t['file']))
        elif hashlib.sha256(open(p, 'rb').read()).hexdigest() != t['sha256']: rep['originals'].append(('被改了', t['file']))
    outs = {}
    for root, _, fs in os.walk(OUT):
        for f in fs:
            if f.lower().endswith('.docx') and not f.startswith('~$'): outs[os.path.relpath(os.path.join(root, f), OUT)] = os.path.join(root, f)
    used = set()
    skip = [a[i + 1] for i, x in enumerate(a) if x == '--skip']   # 老师说不要的那份（重交的两份里的一份）
    for t in T:
        if t['kind'] == 'lock' or t['file'] in skip: continue
        want_name = f"{t['id']}-{t['name']}.docx"
        cands = [k for k in outs if os.path.basename(k) == want_name and k not in used] or [k for k in outs if t['name'] in os.path.basename(k) and k not in used]
        if not cands and t['paras']:   # 文件名认不出，就按题目加正文第一段找
            first = next((norm(x['text']) for x in t['paras'] if x['role'] == 'body'), None)
            for k, p in outs.items():
                if k in used: continue
                try:
                    if first and any(norm(q.text) == first for q in Document(p).paragraphs): cands = [k]; break
                except Exception: pass
        key = f"{t['n']:>2} {t['file']}"
        if not cands: rep['no_output'].append({'file': t['file'], 'kind': t['kind'], 'traps': t['traps'], 'info': t.get('info_where')}); continue
        k = cands[0]; used.add(k); r = check_doc(outs[k], t); r['out'] = k; r['name_ok'] = os.path.basename(k) == want_name; r['traps'] = t['traps']; r['kind'] = t['kind']; rep['files'][key] = r
    rep['unmatched_outputs'] = sorted(k for k in outs if k not in used)
    n = len(rep['files']); clean = [k for k, r in rep['files'].items() if not r['bad'] and not r['fmt'] and r['name_ok']]
    print(f'原件：{"38 个都没动" if not rep["originals"] else rep["originals"]}')
    print(f'输出的 docx {len(outs)} 个；对上学生的 {n} 个；文字、格式、文件名全对的 {len(clean)} 个；没有输出的 {len(rep["no_output"])} 个；对不上人的输出 {len(rep["unmatched_outputs"])} 个')
    agg = {}
    for k, r in rep['files'].items():
        for b in r['bad']: agg.setdefault(re.sub(r'\d+(\.\d+)?', '#', b)[:40], []).append(k.split()[0])
        for b in r['fmt']: agg.setdefault(b, []).append(k.split()[0])
        if not r['name_ok']: agg.setdefault('文件名不是 学号-姓名', []).append(k.split()[0] + ':' + os.path.basename(r['out']))
    for b, ks in sorted(agg.items(), key=lambda x: -len(x[1])): print(f'  {len(ks):>2} 份  {b}  ← {", ".join(ks[:12])}{"…" if len(ks) > 12 else ""}')
    for x in rep['no_output']: print('  没有输出：', x)
    for x in rep['unmatched_outputs']: print('  对不上人的输出：', x)
    if '--json' in a: io.open(a[a.index('--json') + 1], 'w', encoding='utf-8', newline='\n').write(json.dumps(rep, ensure_ascii=False, indent=1))

if __name__ == '__main__': main()
