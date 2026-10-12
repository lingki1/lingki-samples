"""排版的检查。每条学校要求至少对应一个测试；"原来有多乱"都照着学生实际交上来的样子造。"""

import io

import docx
import pytest
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from conftest import PROJECT_ROOT
from docx_inspect import (
    body_paragraphs,
    find,
    footer_paragraphs,
    has_page_number_field,
    header_texts,
    open_bytes,
    page_setup,
    paragraph_format,
    paragraph_text,
    run_format,
    save_bytes,
    style_problems,
    text_runs,
)
from homework_formatter.formatter import format_document
from homework_formatter.rules import EMU_PER_TWIP, load_rules

RULES = load_rules(PROJECT_ROOT / "规则.toml")

TITLE = "食堂排队太久，我想这样解决"
INFO = "高二（3）班  张三  20240301"
HEADING = "一、我发现的问题"
BODY_1 = "每天中午十二点下课，全校的人几乎同时往食堂走。我连着记了 5 天，最长的一次等了 26 分钟。"
BODY_2 = "排队的时间一长，午休就只剩下不到半小时，下午第一节课很多人都在打瞌睡。"


def add(document, text, font="微软雅黑", size=11, **look):
    """加一段"学生风格"的文字：随便的字体、字号、颜色。"""
    paragraph = document.add_paragraph()
    style_run(paragraph.add_run(text), font, size, **look)
    return paragraph


def style_run(run, font="微软雅黑", size=11, bold=False, color=None, underline=False, highlight=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.underline = underline
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if highlight:
        run.font.highlight_color = highlight
    return run


def basic_document(body=(BODY_1, BODY_2)):
    document = docx.Document()
    add(document, TITLE, size=20, bold=True, color="C00000", underline=True)
    add(document, "姓名：张三    学号：20240301")
    add(document, HEADING, bold=True)
    for text in body:
        add(document, text, color="1F4E79")
    return document


def run_formatter(document, name="张三", student_id="20240301"):
    result = format_document(save_bytes(document), name, student_id, RULES)
    return open_bytes(result.docx_bytes), result.changes


def texts(document):
    return [paragraph_text(p) for p in body_paragraphs(document)]


# ---------- 第一条：页面 ----------


def test_rule_1_letter_landscape_with_narrow_margins_becomes_a4_portrait():
    document = basic_document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(27.94), Cm(21.59)
    section.left_margin = section.right_margin = section.top_margin = Cm(1.27)

    output, changes = run_formatter(document)

    assert page_setup(output) == [
        {
            "width": "11906", "height": "16838", "orient": None,
            "top": "1440", "bottom": "1440", "left": "1803", "right": "1803",
        }
    ]  # fmt: skip
    assert any("横向" in c for c in changes)


def test_rule_1_every_section_of_a_multi_section_document_is_fixed():
    document = basic_document()
    document.add_section()
    add(document, BODY_2)

    output, _ = run_formatter(document)

    setups = page_setup(output)
    assert len(setups) == 2
    assert all(s["width"] == "11906" and s["left"] == "1803" for s in setups)


# ---------- 第二条：题目 ----------


def test_rule_2_title_loses_red_bold_underline_and_becomes_heiti_erhao_centered():
    output, _ = run_formatter(basic_document())

    title = body_paragraphs(output)[0]
    assert paragraph_text(title) == TITLE
    assert style_problems(title, RULES.title, RULES) == []
    assert run_format(text_runs(title)[0])["east_asian"] == "黑体"
    assert run_format(text_runs(title)[0])["size_pt"] == 22


def test_rule_2_title_written_with_words_title_style_loses_the_style_and_its_underline_border():
    document = docx.Document()
    document.add_paragraph(TITLE, style="Title")
    add(document, INFO)
    add(document, BODY_1)

    output, _ = run_formatter(document)

    assert style_problems(body_paragraphs(output)[0], RULES.title, RULES) == []


# ---------- 第三条：班级、姓名、学号 ----------


@pytest.mark.parametrize(
    "original",
    ["姓名：张三    学号：20240301", "班级：高二3班 姓名:张三 学号:20240301", "张三 20240301", INFO],
)
def test_rule_3_info_line_is_rewritten_the_same_way_whatever_the_student_wrote(original):
    document = docx.Document()
    add(document, TITLE)
    add(document, original)
    add(document, BODY_1)

    output, changes = run_formatter(document)

    info = body_paragraphs(output)[1]
    assert paragraph_text(info) == INFO
    assert style_problems(info, RULES.info, RULES) == []
    assert run_format(text_runs(info)[0])["east_asian"] == "楷体"
    assert any("姓名行" in c for c in changes) == (original != INFO)


def test_rule_3_info_line_is_added_under_the_title_when_the_student_wrote_none():
    document = docx.Document()
    add(document, TITLE)
    add(document, HEADING)
    add(document, BODY_1)

    output, changes = run_formatter(document)

    assert texts(output)[:3] == [TITLE, INFO, HEADING]
    assert style_problems(body_paragraphs(output)[1], RULES.info, RULES) == []
    assert any("补上" in c for c in changes)


def test_rule_3_info_written_in_the_header_moves_under_the_title_and_header_is_emptied():
    document = docx.Document()
    add(document, TITLE)
    add(document, BODY_1)
    document.sections[0].header.paragraphs[0].text = INFO

    output, changes = run_formatter(document)

    assert texts(output)[:2] == [TITLE, INFO]
    assert header_texts(output) == []
    assert any("页眉" in c for c in changes)


def test_rule_3_info_line_separated_from_the_title_by_blank_lines_ends_up_right_below_it():
    document = docx.Document()
    add(document, TITLE)
    document.add_paragraph("")
    add(document, INFO)
    document.add_paragraph("")
    add(document, BODY_1)

    output, _ = run_formatter(document)

    assert texts(output) == [TITLE, INFO, BODY_1]


def test_rule_3_two_lines_of_info_become_one():
    document = docx.Document()
    add(document, TITLE)
    add(document, "姓名：张三")
    add(document, "学号：20240301")
    add(document, BODY_1)

    output, _ = run_formatter(document)

    assert texts(output) == [TITLE, INFO, BODY_1]


# ---------- 第四条：小标题 ----------


@pytest.mark.parametrize("heading", [HEADING, "1. 我发现的问题", "（一）我发现的问题", "我发现的问题"])
def test_rule_4_heading_becomes_heiti_sihao_flush_left_not_bold_whatever_its_numbering(heading):
    document = docx.Document()
    add(document, TITLE)
    add(document, INFO)
    add(document, heading, bold=True)
    add(document, BODY_1)

    output, _ = run_formatter(document)

    paragraph = find(output, HEADING)
    assert style_problems(paragraph, RULES.heading, RULES) == []
    assert run_format(text_runs(paragraph)[0])["size_pt"] == 14


def test_rule_4_heading_written_with_words_blue_heading_style_loses_the_style():
    document = docx.Document()
    add(document, TITLE)
    add(document, INFO)
    document.add_paragraph(HEADING, style="Heading 2")
    add(document, BODY_1)

    output, _ = run_formatter(document)

    assert style_problems(find(output, HEADING), RULES.heading, RULES) == []


# ---------- 老师追加：小标题编号统一成「一、二、三」 ----------


def document_with_headings(*headings):
    document = docx.Document()
    add(document, TITLE)
    add(document, INFO)
    for heading in headings:
        add(document, heading, bold=True)
        add(document, BODY_1)
    return document


def format_full(document):
    return format_document(save_bytes(document), "张三", "20240301", RULES)


def headings_of(result):
    return [t for t in texts(open_bytes(result.docx_bytes)) if t not in (TITLE, INFO, BODY_1)]


@pytest.mark.parametrize(
    "original, was",
    [
        (["1. 我发现的问题", "2. 原因分析", "3. 我的办法"], "1."),
        (["（一）我发现的问题", "（二）原因分析", "（三）我的办法"], "（一）"),
        (["我发现的问题", "原因分析", "我的办法"], "没有编号"),
    ],
)
def test_headings_are_renumbered_the_same_way_and_the_change_is_reported(original, was):
    result = format_full(document_with_headings(*original, "参考资料"))

    assert headings_of(result) == ["一、我发现的问题", "二、原因分析", "三、我的办法", "参考资料"]
    assert result.headings == headings_of(result)
    assert result.title == TITLE
    (note,) = [c for c in result.changes if "编号" in c]
    assert was in note
    assert result.warnings == []


def test_headings_already_numbered_the_right_way_are_left_alone():
    result = format_full(document_with_headings("一、我发现的问题", "二、原因分析"))

    assert headings_of(result) == ["一、我发现的问题", "二、原因分析"]
    assert not any("编号" in c for c in result.changes)


def test_a_number_typed_in_pieces_with_a_tab_is_still_replaced_cleanly():
    document = document_with_headings()
    paragraph = document.add_paragraph()
    style_run(paragraph.add_run("1"))
    style_run(paragraph.add_run(".")).add_tab()
    style_run(paragraph.add_run("我发现"), bold=True)
    style_run(paragraph.add_run("的问题"))
    add(document, BODY_1)

    result = format_full(document)

    output = open_bytes(result.docx_bytes)
    assert headings_of(result) == [HEADING]
    assert style_problems(find(output, HEADING), RULES.heading, RULES) == []
    assert find(output, HEADING).find(".//" + qn("w:tab")) is None


@pytest.mark.parametrize(
    "headings",
    [
        ["一、我发现的问题", "1. 现象", "2. 数据", "二、我的办法"],
        ["一、我发现的问题", "原因分析", "三、我的办法"],
        ["第一部分 我发现的问题", "第二部分 我的办法"],
    ],
)
def test_when_the_numbering_is_ambiguous_it_is_left_alone_and_the_teacher_is_told(headings):
    result = format_full(document_with_headings(*headings))

    assert headings_of(result) == headings
    assert any("编号" in w or "第一部分" in w for w in result.warnings)
    output = open_bytes(result.docx_bytes)
    for heading in headings:  # 格式照样统一
        assert style_problems(find(output, heading), RULES.heading, RULES) == []


def test_headings_using_words_automatic_numbering_keep_it_and_get_no_typed_number():
    document = document_with_headings()
    document.add_paragraph("我发现的问题", style="List Number")
    add(document, BODY_1)

    result = format_full(document)

    output = open_bytes(result.docx_bytes)
    heading = find(output, "我发现的问题")
    assert heading.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None  # 自动编号还在
    assert any("自动编号" in w for w in result.warnings)


def test_numbering_that_shows_no_number_does_not_count_as_automatic_numbering():
    # LibreOffice 转出来的标题常带一个"不显示编号"的编号设置
    document = document_with_headings()
    paragraph = add(document, "（一）我发现的问题")
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_numId().val = 0
    add(document, BODY_1)

    result = format_full(document)

    output = open_bytes(result.docx_bytes)
    assert result.warnings == []
    assert style_problems(find(output, HEADING), RULES.heading, RULES) == []
    assert find(output, HEADING).find(qn("w:pPr") + "/" + qn("w:numPr")) is None


def test_renumbering_can_be_switched_off_in_the_rules():
    import dataclasses

    rules = dataclasses.replace(RULES, unify_heading_numbers=False)
    document = document_with_headings("1. 我发现的问题", "2. 原因分析")

    result = format_document(save_bytes(document), "张三", "20240301", rules)

    assert headings_of(result) == ["1. 我发现的问题", "2. 原因分析"]


# ---------- 第五条：正文 ----------


def test_rule_5_body_font_size_indent_spacing_alignment_and_colour():
    document = basic_document()
    for paragraph in document.paragraphs[3:]:
        paragraph.paragraph_format.line_spacing = 2.0
        paragraph.alignment = 1  # 居中

    output, _ = run_formatter(document)

    for text in (BODY_1, BODY_2):
        paragraph = find(output, text)
        assert style_problems(paragraph, RULES.body, RULES) == []
        fmt = paragraph_format(paragraph)
        assert (fmt["first_line_chars"], fmt["first_line"]) == ("200", "480")  # 两个字
        assert (fmt["line"], fmt["jc"]) == ("360", "both")  # 1.5 倍，两端对齐
        run = run_format(text_runs(paragraph)[0])
        assert (run["east_asian"], run["ascii"], run["size_pt"]) == ("宋体", "Times New Roman", 12)


@pytest.mark.parametrize("indent", ["　　", "    ", "\t", "  \t", "　 "])
def test_rule_5_hand_typed_indentation_is_replaced_by_a_real_indent(indent):
    output, changes = run_formatter(basic_document(body=(indent + BODY_1, indent + BODY_2)))

    assert texts(output)[3:] == [BODY_1, BODY_2]
    assert any("2 段" in c and "缩进" in c for c in changes)


def test_rule_5_indentation_made_with_the_tab_key_is_removed():
    document = basic_document(body=())
    paragraph = document.add_paragraph()
    paragraph.add_run().add_tab()
    paragraph.add_run(BODY_1 + "  ")

    output, _ = run_formatter(document)

    assert texts(output)[-1] == BODY_1
    assert find(output, BODY_1).find(".//" + qn("w:tab")) is None


def test_rule_5_blank_lines_between_paragraphs_are_removed():
    document = docx.Document()
    for text in (TITLE, "", INFO, "", HEADING, BODY_1, "", "   ", BODY_2, ""):
        add(document, text)

    output, changes = run_formatter(document)

    assert texts(output) == [TITLE, INFO, HEADING, BODY_1, BODY_2]
    assert any("5 个空段落" in c for c in changes)


def test_rule_5_soft_line_breaks_become_real_paragraphs():
    document = basic_document(body=())
    paragraph = document.add_paragraph()
    run = style_run(paragraph.add_run("　　" + BODY_1))
    run.add_break()
    run.add_text("　　" + BODY_2)

    output, changes = run_formatter(document)

    assert texts(output)[3:] == [BODY_1, BODY_2]
    for text in (BODY_1, BODY_2):
        assert style_problems(find(output, text), RULES.body, RULES) == []
    assert output.element.body.find(".//" + qn("w:br")) is None
    assert any("软回车" in c for c in changes)


def test_rule_5_paragraph_spacing_inherited_from_the_document_defaults_is_gone():
    output, _ = run_formatter(basic_document())

    defaults = output.styles.element.find(qn("w:docDefaults"))
    assert defaults.find(".//" + qn("w:spacing")) is None


def test_students_own_bold_and_highlight_are_removed_and_reported():
    document = basic_document(body=())
    emphasised = "把值日表改成按星期固定，比如每个星期二都是同一组人。"
    add(document, emphasised, bold=True, highlight=WD_COLOR_INDEX.YELLOW)
    add(document, BODY_2)

    output, changes = run_formatter(document)

    assert style_problems(find(output, emphasised), RULES.body, RULES) == []
    (note,) = [c for c in changes if "加粗" in c]
    assert "黄色高亮" in note and "把值日表改成按星期固定" in note


def test_superscript_survives_because_removing_it_would_change_the_meaning():
    document = basic_document(body=())
    paragraph = document.add_paragraph()
    style_run(paragraph.add_run("教室的面积大约是 60 m"))
    style_run(paragraph.add_run("2")).font.superscript = True
    style_run(paragraph.add_run("，一共有 12 盏灯。"))

    output, _ = run_formatter(document)

    runs = [run_format(r) for r in text_runs(body_paragraphs(output)[-1])]
    assert [r["vert_align"] for r in runs] == [None, "superscript", None]
    assert all(r["east_asian"] == "宋体" for r in runs)


def test_references_heading_is_a_heading_and_entries_are_body():
    document = basic_document()
    add(document, "参考资料", bold=True)
    add(document, "[1] 学校总务处. 2024 年春季学期后勤工作小结（校内资料）.")

    output, _ = run_formatter(document)

    assert style_problems(find(output, "参考资料"), RULES.heading, RULES) == []
    entry = find(output, "[1] 学校总务处. 2024 年春季学期后勤工作小结（校内资料）.")
    assert style_problems(entry, RULES.body, RULES) == []


# ---------- 第六条：表格和图片 ----------


def document_with_table(column_width_cm=None):
    document = basic_document()
    table = document.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for index, text in enumerate(["日期", "最短（分钟）", "最长（分钟）"]):
        table.cell(0, index).text = text
        table.cell(1, index).text = str(index + 9)
    if column_width_cm:
        for grid_col in table._tbl.tblGrid.findall(qn("w:gridCol")):
            grid_col.set(qn("w:w"), str(Cm(column_width_cm).twips))
        for cell in table._cells:
            cell.width = Cm(column_width_cm)
    add(document, BODY_2)
    return document


def test_rule_6_table_is_kept_centered_and_its_text_uses_the_body_font():
    output, _ = run_formatter(document_with_table())

    (table,) = output.tables
    assert [c.text for c in table.rows[0].cells] == ["日期", "最短（分钟）", "最长（分钟）"]
    assert table._tbl.tblPr.find(qn("w:jc")).get(qn("w:val")) == "center"
    for paragraph in table._tbl.iter(qn("w:p")):
        assert style_problems(paragraph, RULES.table_text, RULES) == []
    # 表格还在原来的位置：两段正文之间
    kinds = [child.tag.split("}")[1] for child in output.element.body.iterchildren()]
    assert kinds == ["p", "p", "p", "p", "p", "tbl", "p", "sectPr"]


def test_rule_6_table_wider_than_the_new_page_is_shrunk_to_fit():
    output, changes = run_formatter(document_with_table(column_width_cm=6))

    (table,) = output.tables
    widths = [int(g.get(qn("w:w"))) for g in table._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert sum(widths) <= RULES.text_width_twips
    assert max(widths) - min(widths) <= 1  # 各列比例不变
    assert all(int(w.get(qn("w:w"))) <= max(widths) for w in table._tbl.iter(qn("w:tcW")))
    assert any("表格" in c and "缩小" in c for c in changes)


def test_rule_6_table_that_already_fits_is_not_resized():
    document = document_with_table(column_width_cm=3)
    before = [g.get(qn("w:w")) for g in document.tables[0]._tbl.tblGrid]

    output, changes = run_formatter(document)

    assert [g.get(qn("w:w")) for g in output.tables[0]._tbl.tblGrid] == before
    assert not any("表格" in c for c in changes)


def document_with_picture(tiny_png, width_cm):
    document = basic_document()
    document.add_picture(io.BytesIO(tiny_png), width=Cm(width_cm))
    add(document, BODY_2)
    return document


def test_rule_6_picture_is_kept_and_centered_without_indent(tiny_png):
    output, changes = run_formatter(document_with_picture(tiny_png, width_cm=10))

    (shape,) = output.inline_shapes
    assert shape.width == Cm(10)
    picture_paragraph = next(
        p for p in body_paragraphs(output) if p.find(".//" + qn("w:drawing")) is not None
    )
    fmt = paragraph_format(picture_paragraph)
    assert (fmt["jc"], fmt["first_line"], fmt["left"]) == ("center", "0", "0")
    assert fmt["line_rule"] == "auto"  # 固定行距会把图片切掉
    assert not any("图片" in c for c in changes)


def test_rule_6_picture_wider_than_the_new_page_is_shrunk_keeping_its_proportions(tiny_png):
    output, changes = run_formatter(document_with_picture(tiny_png, width_cm=17))

    (shape,) = output.inline_shapes
    assert shape.width <= RULES.text_width_twips * EMU_PER_TWIP
    assert abs(shape.width - shape.height) <= 1  # 原图是正方形
    assert any("图片" in c and "缩小" in c for c in changes)


# ---------- 第七条：页码 ----------


def test_rule_7_page_number_is_added_at_the_bottom_center():
    output, _ = run_formatter(basic_document())

    (footer,) = footer_paragraphs(output)
    assert has_page_number_field(footer)
    assert paragraph_format(footer)["jc"] == "center"


def test_rule_7_an_existing_footer_and_header_are_replaced():
    document = basic_document()
    document.sections[0].footer.paragraphs[0].text = "学生自己写的页脚"
    document.sections[0].header.paragraphs[0].text = "研究性学习小报告"

    output, changes = run_formatter(document)

    (footer,) = footer_paragraphs(output)
    assert has_page_number_field(footer)
    assert "学生自己写的页脚" not in paragraph_text(footer)
    assert header_texts(output) == []
    assert any("页眉" in c and "研究性学习小报告" in c for c in changes)


# ---------- 整体 ----------


def test_nothing_but_the_info_line_changes_in_the_words():
    output, _ = run_formatter(basic_document())

    assert texts(output) == [TITLE, INFO, HEADING, BODY_1, BODY_2]


def test_a_tidy_document_reports_no_changes():
    document = docx.Document()
    for text in (TITLE, INFO, HEADING, BODY_1, BODY_2):
        add(document, text)

    _, changes = run_formatter(document)

    assert changes == []


def test_formatting_twice_gives_the_same_words_and_no_further_changes():
    first = format_document(save_bytes(basic_document()), "张三", "20240301", RULES)

    second = format_document(first.docx_bytes, "张三", "20240301", RULES)

    assert texts(open_bytes(second.docx_bytes)) == texts(open_bytes(first.docx_bytes))
    assert second.changes == []


def test_a_document_without_a_title_gets_its_info_line_first():
    document = docx.Document()
    add(document, "姓名：张三    学号：20240301")
    add(document, "我发现的问题")
    add(document, BODY_1)

    output, _ = run_formatter(document)

    assert texts(output) == [INFO, "一、我发现的问题", BODY_1]
    assert style_problems(body_paragraphs(output)[1], RULES.heading, RULES) == []
