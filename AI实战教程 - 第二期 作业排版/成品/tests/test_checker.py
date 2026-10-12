"""质检的检查：先确认好文件能通过，再故意把好文件弄坏，每种坏法都必须被抓出来。"""

import io

import docx
import pytest
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from conftest import PROJECT_ROOT
from docx_inspect import body_paragraphs, find, open_bytes, save_bytes
from homework_formatter.checker import check_document, check_filename
from homework_formatter.formatter import format_document
from homework_formatter.rules import load_rules

RULES = load_rules(PROJECT_ROOT / "规则.toml")
TITLE = "食堂排队太久，我想这样解决"
INFO = "高二（3）班  张三  20240301"
HEADING = "一、我发现的问题"
HEADING_2 = "二、原因分析"
BODY = "每天中午十二点下课，全校的人几乎同时往食堂走。我连着记了 5 天。"
BODY_2 = "排队的时间一长，午休就只剩下不到半小时。"
REFERENCE = "[1] 学校总务处. 2024 年春季学期后勤工作小结（校内资料）."


@pytest.fixture
def good(tiny_png):
    document = docx.Document()
    for text in (TITLE, "张三 20240301", "1. 我发现的问题", BODY):
        document.add_paragraph(text)
    document.add_picture(io.BytesIO(tiny_png), width=Cm(5))
    document.add_paragraph("2. 原因分析")
    document.add_paragraph(BODY_2)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "日期"
    table.cell(0, 1).text = "分钟"
    document.add_paragraph("参考资料")
    document.add_paragraph(REFERENCE)
    return format_document(save_bytes(document), "张三", "20240301", RULES).docx_bytes


def check(data, info_line=INFO, **kwargs):
    return check_document(data, RULES, info_line, **kwargs)


def spoil(data, change):
    document = open_bytes(data)
    change(document)
    return save_bytes(document)


def first_run(document, text):
    return docx.text.paragraph.Paragraph(find(document, text), None).runs[0]


def paragraph(document, text):
    return docx.text.paragraph.Paragraph(find(document, text), None)


def test_a_properly_formatted_document_passes(good):
    assert check(good) == []


# 每一行：（怎么弄坏, 报告里应该出现的话）。按学校的 8 条排。
def landscape(d):
    d.sections[0].orientation = WD_ORIENT.LANDSCAPE


def letter_paper(d):
    d.sections[0].page_width, d.sections[0].page_height = Cm(21.59), Cm(27.94)


def narrow_margins(d):
    d.sections[0].left_margin = d.sections[0].right_margin = Cm(1.27)


def short_top_margin(d):
    d.sections[0].top_margin = Cm(2.0)


def set_east_asian(run, font):
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def remove_info(d):
    element = find(d, INFO)
    element.getparent().remove(element)


def move_info_down(d):
    find(d, HEADING).addnext(find(d, INFO))


def add_blank_line(d):
    find(d, BODY).addnext(d.add_paragraph("")._p)


def add_soft_break(d):
    first_run(d, BODY).add_break()


def type_indent(d):
    run = first_run(d, BODY)
    run.text = "　　" + run.text


def tab_indent(d):
    run = first_run(d, BODY)._element
    run.insert(1, docx.oxml.OxmlElement("w:tab"))


def uncenter_table(d):
    d.tables[0]._tbl.tblPr.find(qn("w:jc")).set(qn("w:val"), "left")


def widen_table(d):
    for column in d.tables[0]._tbl.tblGrid.findall(qn("w:gridCol")):
        column.set(qn("w:w"), "6000")


def table_font(d):
    set_east_asian(d.tables[0].cell(0, 0).paragraphs[0].runs[0], "楷体")


def picture_paragraph(d):
    return next(p for p in body_paragraphs(d) if p.find(".//" + qn("w:drawing")) is not None)


def uncenter_picture(d):
    picture_paragraph(d).find(qn("w:pPr") + "/" + qn("w:jc")).set(qn("w:val"), "left")


def widen_picture(d):
    d.inline_shapes[0].width = Cm(18)


def remove_footer(d):
    d.sections[0].footer.is_linked_to_previous = True


def footer_left(d):
    d.sections[0].footer.paragraphs[0].alignment = 0


def footer_without_number(d):
    footer = d.sections[0].footer._element
    for child in list(footer):
        footer.remove(child)
    d.sections[0].footer.add_paragraph("第一页")


def add_header(d):
    d.sections[0].header.paragraphs[0].text = "研究性学习小报告"


def set_heading_text(text):
    def change(d):
        first_run(d, HEADING).text = text

    return change


SPOILED = [
    # 第一条
    (landscape, "第一条：纸张不是 A4 竖向"),
    (letter_paper, "第一条：纸张不是 A4 竖向"),
    (narrow_margins, "第一条：左右页边距"),
    (short_top_margin, "第一条：上下页边距"),
    # 第二条
    (lambda d: set_east_asian(first_run(d, TITLE), "微软雅黑"), "第二条：题目「食堂排队太久，我想这…」：中文字体应为黑体，实际是微软雅黑"),
    (lambda d: setattr(first_run(d, TITLE).font, "size", Pt(16)), "第二条：题目「食堂排队太久，我想这…」：字号应为 22 磅"),
    (lambda d: setattr(first_run(d, TITLE).font, "bold", True), "有加粗"),
    (lambda d: setattr(first_run(d, TITLE).font, "underline", True), "有下划线"),
    (lambda d: setattr(paragraph(d, TITLE), "alignment", 0), "对齐应为居中，实际是靠左"),
    (lambda d: setattr(first_run(d, TITLE).font.color, "rgb", RGBColor(0xC0, 0, 0)), "字不是黑色"),
    # 第三条
    (lambda d: setattr(first_run(d, INFO), "text", "张三 20240301"), "第三条：姓名行应为「高二（3）班  张三  20240301」"),
    (lambda d: setattr(first_run(d, INFO), "text", "高二（3）班 张三 20240301"), "第三条：姓名行应为"),
    (remove_info, "第三条：姓名行应为「高二（3）班  张三  20240301」，实际是没有"),
    (move_info_down, "第三条：姓名行不在题目下面一行"),
    (lambda d: set_east_asian(first_run(d, INFO), "宋体"), "第三条：姓名行「高二（3）班  张三…」：中文字体应为楷体"),
    (lambda d: setattr(paragraph(d, INFO), "alignment", 0), "第三条"),
    # 第四条
    (lambda d: setattr(first_run(d, HEADING).font, "size", Pt(12)), "第四条：小标题「一、我发现的问题」：字号应为 14 磅"),
    (lambda d: setattr(first_run(d, HEADING).font, "bold", True), "第四条：小标题「一、我发现的问题」：有加粗"),
    (lambda d: setattr(paragraph(d, HEADING).paragraph_format, "first_line_indent", Pt(24)), "第四条：小标题「一、我发现的问题」：首行应该不缩进"),
    (lambda d: setattr(paragraph(d, HEADING).paragraph_format, "left_indent", Pt(24)), "整段有多余的缩进"),
    (lambda d: set_east_asian(first_run(d, HEADING), "宋体"), "第四条"),
    # 老师追加：小标题编号
    (set_heading_text("1. 我发现的问题"), "小标题编号：「1. 我发现的问题」的编号应为「一、」"),
    (set_heading_text("（一）我发现的问题"), "小标题编号"),
    (set_heading_text("我发现的问题"), "小标题编号：「我发现的问题」的编号应为「一、」"),
    (set_heading_text("二、我发现的问题"), "小标题编号：「二、我发现的问题」的编号应为「一、」"),
    (lambda d: setattr(first_run(d, "参考资料"), "text", "三、参考资料"), "小标题编号：「三、参考资料」不该有编号"),
    # 第五条
    (lambda d: set_east_asian(first_run(d, BODY), "微软雅黑"), "第五条：正文「每天中午十二点下课，…」：中文字体应为宋体，实际是微软雅黑"),
    (lambda d: setattr(first_run(d, BODY).font, "name", "Arial"), "英文数字的字体应为 Times New Roman，实际是 Arial"),
    (lambda d: setattr(first_run(d, BODY).font, "size", Pt(10.5)), "字号应为 12 磅，实际是 10.5"),
    (lambda d: setattr(first_run(d, BODY).font.color, "rgb", RGBColor(0x1F, 0x4E, 0x79)), "字不是黑色"),
    (lambda d: setattr(first_run(d, BODY).font, "highlight_color", WD_COLOR_INDEX.YELLOW), "有高亮"),
    (lambda d: setattr(paragraph(d, BODY).paragraph_format, "first_line_indent", Pt(0)), "首行应该缩进 2 个字"),
    (lambda d: setattr(paragraph(d, BODY).paragraph_format, "line_spacing", 1.0), "行距应为 1.5 倍"),
    (lambda d: setattr(paragraph(d, BODY).paragraph_format, "line_spacing", Pt(20)), "行距应为 1.5 倍"),
    (lambda d: setattr(paragraph(d, BODY), "alignment", 0), "对齐应为两端对齐，实际是靠左"),
    (lambda d: setattr(paragraph(d, BODY).paragraph_format, "space_after", Pt(10)), "段前或段后有多余的间距"),
    (lambda d: setattr(find(d, BODY).get_or_add_pPr().get_or_add_pStyle(), "val", "Heading2"), "还带着 Word 的样式"),
    (lambda d: set_east_asian(first_run(d, REFERENCE), "黑体"), "第五条：正文「[1] 学校总务处.…」"),
    (add_blank_line, "第五条：段落之间有空行"),
    (add_soft_break, "里面有软回车"),
    (type_indent, "开头或结尾有手敲的空格或 Tab"),
    (tab_indent, "开头或结尾有手敲的空格或 Tab"),
    # 第六条
    (uncenter_table, "第六条：第 1 个表格没有居中"),
    (widen_table, "第六条：第 1 个表格比页面的版心宽"),
    (table_font, "第六条：第 1 个表格里的字：中文字体应为宋体"),
    (uncenter_picture, "第六条：图片没有居中"),
    (widen_picture, "第六条：图片比页面的版心宽"),
    # 第七条
    (remove_footer, "第七条：页面底端没有页码"),
    (footer_left, "第七条：页码没有居中"),
    (footer_without_number, "第七条：页面底端没有页码"),
    (add_header, "页眉里还有字"),
]


@pytest.mark.parametrize("change, expected", SPOILED, ids=[e for _, e in SPOILED])
def test_every_way_of_spoiling_a_good_document_is_caught(good, change, expected):
    problems = check(spoil(good, change))

    assert any(expected in p for p in problems), problems


def test_each_fault_is_reported_once_not_once_per_run(good):
    def change(d):
        for run in paragraph(d, BODY).runs:
            set_east_asian(run, "微软雅黑")

    problems = check(spoil(good, change))

    assert len(problems) == 1


def test_missing_title_is_a_problem_unless_already_known(good):
    def remove_title(d):
        element = find(d, TITLE)
        element.getparent().remove(element)

    spoiled = spoil(good, remove_title)

    assert "第二条：没有题目" in check(spoiled)
    assert check(spoiled, expect_title=False) == []


def test_when_the_owner_is_unknown_the_info_line_wording_is_not_checked_but_its_look_is():
    document = docx.Document()
    for text in (TITLE, "姓名：某某    学号：20249999", HEADING, BODY):
        document.add_paragraph(text)
    data = format_document(save_bytes(document), None, None, RULES).docx_bytes

    assert check(data, info_line=None) == []

    spoiled = spoil(data, lambda d: set_east_asian(first_run(d, "姓名：某某    学号：20249999"), "宋体"))
    assert any("第三条" in p for p in check(spoiled, info_line=None))


@pytest.mark.parametrize(
    "filename, expected",
    [
        ("20240301-张三.docx", None),
        ("张三20240301.docx", "不是「学号-姓名.docx」的写法"),
        ("20240301-张三.doc", "不是「学号-姓名.docx」的写法"),
        ("20240302-张三.docx", "对不上"),
        ("20240301-李四.docx", "对不上"),
    ],
)
def test_rule_8_filename(filename, expected):
    problems = check_filename(filename, INFO, RULES)

    assert problems == [] if expected is None else expected in problems[0]
