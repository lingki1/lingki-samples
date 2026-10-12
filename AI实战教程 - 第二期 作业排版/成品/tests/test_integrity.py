import io

import docx
from docx.oxml.ns import qn

from conftest import PROJECT_ROOT
from docx_inspect import open_bytes, save_bytes
from homework_formatter.formatter import format_document
from homework_formatter.integrity import verify_content
from homework_formatter.rules import load_rules

RULES = load_rules(PROJECT_ROOT / "规则.toml")
INFO = RULES.info_line("张三", "20240301")
BODY = "每天中午十二点下课，全校的人几乎同时往食堂走。我连着记了 5 天。"


def original(tiny_png=None, with_table=False):
    document = docx.Document()
    document.add_paragraph("食堂排队太久，我想这样解决")
    document.add_paragraph("姓名：张三    学号：20240301")
    document.add_paragraph("")
    document.add_paragraph("一、我发现的问题")
    document.add_paragraph("　　" + BODY)
    if tiny_png:
        document.add_picture(io.BytesIO(tiny_png))
    if with_table:
        document.add_table(rows=1, cols=2).cell(0, 0).text = "日期"
    return save_bytes(document)


def formatted(data):
    return format_document(data, "张三", "20240301", RULES).docx_bytes


def tamper(data, change):
    document = open_bytes(data)
    change(document)
    return save_bytes(document)


def test_a_properly_formatted_document_passes(tiny_png):
    data = original(tiny_png, with_table=True)

    assert verify_content(data, formatted(data), INFO, "张三", "20240301") == []


def test_rewritten_info_line_removed_blank_lines_and_indentation_are_not_differences():
    data = original()

    assert verify_content(data, formatted(data), INFO, "张三", "20240301") == []


def numbered(prefix):
    document = docx.Document()
    document.add_paragraph("食堂排队太久，我想这样解决")
    document.add_paragraph("张三 20240301")
    document.add_paragraph(prefix + "我发现的问题")
    document.add_paragraph(BODY)
    return save_bytes(document)


def test_renumbered_headings_are_not_differences():
    for prefix in ("1. ", "（一）", "", "一、"):
        data = numbered(prefix)

        assert verify_content(data, formatted(data), INFO, "张三", "20240301") == []


def test_a_heading_whose_words_changed_is_still_caught():
    data = numbered("1. ")

    def change(document):
        document.paragraphs[2].runs[0].text = "一、我发现的毛病"

    (problem,) = verify_content(data, tamper(formatted(data), change), INFO, "张三", "20240301")

    assert "毛病" in problem


def test_a_lost_sentence_is_caught_and_shown():
    data = original()
    output = tamper(formatted(data), lambda d: d.paragraphs[-1].clear().add_run("每天中午十二点下课。"))

    (problem,) = verify_content(data, output, INFO, "张三", "20240301")

    assert "文字" in problem
    assert "全校的人" in problem


def test_a_single_changed_character_is_caught():
    data = original()

    def change(document):
        run = document.paragraphs[-1].runs[0]
        run.text = run.text.replace("5 天", "6 天")

    (problem,) = verify_content(data, tamper(formatted(data), change), INFO, "张三", "20240301")

    assert "5" in problem and "6" in problem


def test_added_text_is_caught():
    data = original()
    output = tamper(formatted(data), lambda d: d.add_paragraph("多出来的一段"))

    assert any("多出来的一段" in p for p in verify_content(data, output, INFO, "张三", "20240301"))


def test_a_lost_table_is_caught():
    data = original(with_table=True)

    def drop_table(document):
        table = document.tables[0]._tbl
        table.getparent().remove(table)

    problems = verify_content(data, tamper(formatted(data), drop_table), INFO, "张三", "20240301")

    assert any("表格" in p for p in problems)


def test_a_lost_picture_is_caught(tiny_png):
    data = original(tiny_png)

    def drop_picture(document):
        drawing = document.element.body.find(".//" + qn("w:drawing"))
        drawing.getparent().remove(drawing)

    problems = verify_content(data, tamper(formatted(data), drop_picture), INFO, "张三", "20240301")

    assert any("图片" in p for p in problems)


def test_a_missing_or_wrong_info_line_is_caught():
    data = original()

    def spoil(document):
        document.paragraphs[1].runs[0].text = "高二（3）班  李四  20240301"

    problems = verify_content(data, tamper(formatted(data), spoil), INFO, "张三", "20240301")

    assert any("姓名行" in p for p in problems)
