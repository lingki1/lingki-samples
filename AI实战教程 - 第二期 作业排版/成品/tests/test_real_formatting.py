"""第 3 步：把真实作业完整排一遍，逐份、逐段对照 8 条要求。"""

import io

import docx
import pytest
from docx.oxml.ns import qn

from conftest import PROJECT_ROOT, REAL_INPUT, REAL_ROSTER, needs_real_roster, needs_real_samples
from docx_inspect import (
    footer_paragraphs,
    has_page_number_field,
    header_texts,
    page_setup,
    paragraph_format,
    style_problems,
)
from homework_formatter import pipeline
from homework_formatter.converter import find_soffice
from homework_formatter.decisions import load_decisions
from homework_formatter.inventory import Outcome
from homework_formatter.numbering import split_number
from homework_formatter.roster import load_roster
from homework_formatter.rules import load_rules
from homework_formatter.scanner import fingerprint
from homework_formatter.structure import Role, classify

REAL_DECISIONS = PROJECT_ROOT / "老师的决定.xlsx"
RULES = load_rules(PROJECT_ROOT / "规则.toml")

pytestmark = [
    needs_real_samples,
    needs_real_roster,
    pytest.mark.skipif(not REAL_DECISIONS.is_file(), reason="没有「老师的决定.xlsx」"),
    pytest.mark.skipif(find_soffice() is None, reason="没有 LibreOffice，读不了旧版 .doc"),
]

A4_PAGE = {
    "width": "11906", "height": "16838", "orient": None,
    "top": "1440", "bottom": "1440", "left": "1803", "right": "1803",
}  # fmt: skip
STYLES = {
    Role.TITLE: RULES.title,
    Role.INFO: RULES.info,
    Role.HEADING: RULES.heading,
    Role.BODY: RULES.body,
}


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    before = fingerprint(REAL_INPUT)
    outcome = pipeline.run(
        REAL_INPUT,
        tmp_path_factory.mktemp("排版结果"),
        RULES,
        load_roster(REAL_ROSTER),
        load_decisions(REAL_DECISIONS),
    )
    assert fingerprint(REAL_INPUT) == before
    return outcome


@pytest.fixture(scope="module")
def outputs(result):
    return {
        path.name: docx.Document(str(path))
        for path in sorted((result.run_dir / "已完成").iterdir())
    }


def roster_names():
    return {f"{student_id}-{name}.docx" for student_id, name in load_roster(REAL_ROSTER).entries}


def body_texts(document, role=Role.BODY):
    return [b.text for b in classify(document) if b.role is role]


def test_where_all_38_files_went(result):
    inventory = result.inventory
    assert len(inventory.records) == 38
    assert inventory.count(Outcome.DONE) == 33
    assert inventory.count(Outcome.READY) == 0
    not_done = {r.source.path.name: r.outcome for r in inventory.records if r.outcome is not Outcome.DONE}
    assert not_done == {
        "高二3班 安然 小报告.docx": Outcome.REVIEW,  # 没有题目，第 4 步处理
        "研究性学习 齐悦.pdf": Outcome.UNPROCESSED,
        "尚可.docx": Outcome.UNPROCESSED,
        "作业_梁思齐_20240324.docx": Outcome.IGNORED,  # 老师决定不要
        "~$林知夏.docx": Outcome.IGNORED,
    }
    assert result.originals_untouched


def test_rule_8_file_names(outputs):
    missing = {"20240325-安然.docx", "20240332-齐悦.docx", "20240333-尚可.docx"}

    assert set(outputs) == roster_names() - missing


def test_rule_1_every_page_is_a4_portrait_with_the_right_margins(outputs):
    wrong = {name: page_setup(d) for name, d in outputs.items() if page_setup(d) != [A4_PAGE]}

    assert wrong == {}


def test_rules_2_to_5_every_paragraph_of_every_assignment(outputs):
    wrong = {}
    for name, document in outputs.items():
        for block in classify(document):
            if block.role in STYLES:
                problems = style_problems(block.element, STYLES[block.role], RULES)
                if problems:
                    wrong.setdefault(name, []).append((block.text[:12], problems))
            elif block.role not in (Role.IMAGE, Role.TABLE):
                wrong.setdefault(name, []).append((block.role.value, "不该出现"))
    assert wrong == {}


def test_every_assignment_has_title_then_info_line_then_content(outputs):
    for name, document in outputs.items():
        roles = [b.role for b in classify(document)]
        student_id, student = name[:-5].split("-")
        assert roles[:2] == [Role.TITLE, Role.INFO], name
        assert roles.count(Role.TITLE) == 1 and roles.count(Role.INFO) == 1, name
        assert body_texts(document, Role.INFO) == [RULES.info_line(student, student_id)], name
        # 每份作业三个小节、各两段正文；有参考资料的再多一个小标题和两条
        headings, bodies = roles.count(Role.HEADING), roles.count(Role.BODY)
        assert (headings, bodies) in [(3, 6), (4, 8)], name


def test_rule_5_no_hand_typed_indentation_or_soft_breaks_are_left(outputs):
    for name, document in outputs.items():
        for text in body_texts(document) + body_texts(document, Role.HEADING):
            assert text == text.strip(" \t　 "), (name, text[:12])
        assert document.element.body.find(".//" + qn("w:br")) is None, name
        assert document.element.body.find(".//" + qn("w:tab")) is None, name


def test_rule_6_the_picture_and_the_table_survive_centered(outputs):
    with_picture = outputs["20240308-白鹿鸣.docx"]
    (shape,) = with_picture.inline_shapes
    picture = next(b for b in classify(with_picture) if b.role is Role.IMAGE)
    assert paragraph_format(picture.element)["jc"] == "center"
    assert paragraph_format(picture.element)["first_line"] == "0"
    assert shape.width == 3600000  # 10 厘米，原样

    (table,) = outputs["20240303-陈一诺.docx"].tables
    assert table._tbl.tblPr.find(qn("w:jc")).get(qn("w:val")) == "center"
    assert [c.text for c in table.rows[0].cells] == ["日期", "最短（分钟）", "最长（分钟）"]
    assert [[c.text for c in row.cells] for row in table.rows[1:]] == [
        ["周一", "9", "21"],
        ["周二", "11", "18"],
        ["周三", "8", "24"],
    ]
    widths = [int(g.get(qn("w:w"))) for g in table._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert sum(widths) <= RULES.text_width_twips
    for paragraph in table._tbl.iter(qn("w:p")):
        assert style_problems(paragraph, RULES.table_text, RULES) == []

    others = [d for n, d in outputs.items() if n not in ("20240308-白鹿鸣.docx", "20240303-陈一诺.docx")]
    assert all(len(d.inline_shapes) == 0 and len(d.tables) == 0 for d in others)


def test_rule_7_every_assignment_has_a_centered_page_number_and_no_header(outputs):
    for name, document in outputs.items():
        (footer,) = footer_paragraphs(document)
        assert has_page_number_field(footer), name
        assert paragraph_format(footer)["jc"] == "center", name
        assert header_texts(document) == [], name


def test_step_5_every_written_file_passed_quality_control(result):
    written = [r for r in result.inventory.records if r.output_name]

    assert len(written) == 34  # 已完成 33 + 需人工确认 1
    assert {r.source.path.name: r.quality for r in written if r.quality != []} == {}


UNIFIED = ["一、我发现的问题", "二、原因分析", "三、我的办法"]


def test_teachers_addition_every_assignment_numbers_its_headings_the_same_way(result, outputs):
    for name, document in outputs.items():
        headings = body_texts(document, Role.HEADING)
        assert headings in (UNIFIED, UNIFIED + ["参考资料"]), name
    # 报告里列的也是排版后的小标题
    for record in result.inventory.records:
        if record.outcome is Outcome.DONE:
            assert record.structure.headings[:3] == UNIFIED, record.source.path.name


def test_the_report_says_how_the_numbering_was_changed(result):
    was_digits = ["20240302-周子衿.docx", "20240301-林知夏.docx", "20240324-梁思齐.docx"]
    was_parentheses = ["20240330-骆一鸣.docx", "20240327-邱实.docx", "20240315-苏禾.docx"]
    was_missing = ["20240304-许嘉树.docx", "20240305-沈星河.docx", "20240335-池鱼.docx"]
    untouched = ["20240308-白鹿鸣.docx", "20240322-乔木.docx", "20240311-夏以宁.docx"]

    for filename in was_digits:
        assert "原来是「1.」这种写法" in changes_of(result, filename)
    for filename in was_parentheses:
        assert "原来是「（一）」这种写法" in changes_of(result, filename)
    for filename in was_missing:
        assert "原来没有编号" in changes_of(result, filename)
    for filename in untouched:
        assert "编号" not in changes_of(result, filename)


def changes_of(result, filename):
    (record,) = [r for r in result.inventory.records if r.new_filename == filename and r.outcome is Outcome.DONE]
    return "\n".join(record.changes)


def test_the_special_cases_seen_by_hand_are_reported(result):
    assert "原来写在页眉里" in changes_of(result, "20240302-周子衿.docx")
    assert "原来写在页眉里" in changes_of(result, "20240316-陆景明.docx")
    assert "原文里没有" in changes_of(result, "20240313-叶知秋.docx")
    assert "原文里没有" in changes_of(result, "20240321-童年.docx")
    assert "横向" in changes_of(result, "20240317-温言.docx")
    assert "3 处软回车" in changes_of(result, "20240311-夏以宁.docx")
    assert "8 个空段落" in changes_of(result, "20240314-方可.docx")
    assert "8 个空段落" in changes_of(result, "20240305-沈星河.docx")
    assert "空段落" in changes_of(result, "20240327-邱实.docx")
    assert "表格" in changes_of(result, "20240303-陈一诺.docx")
    for filename in ("20240330-骆一鸣.docx", "20240309-江晚晴.docx"):
        assert "加粗、黄色高亮" in changes_of(result, filename)


def test_the_version_the_teacher_chose_is_the_one_that_was_formatted(outputs):
    text = "".join(body_texts(outputs["20240324-梁思齐.docx"]))

    assert "其中有 10 天" in text
    assert "其中有 9 天" not in text


def test_words_are_identical_before_and_after(result, outputs):
    """独立于程序自带的比对再核一遍：正文和小标题的文字，逐段相同。"""
    for record in result.inventory.records:
        if record.outcome is not Outcome.DONE:
            continue
        before = docx.Document(io.BytesIO(record.docx_bytes))
        after = outputs[record.new_filename]

        def words(document):
            lines = []
            for block in classify(document):
                if block.role is Role.HEADING:
                    lines.append(split_number(block.text)[1])  # 编号按要求会变，比编号后面的字
                elif block.role in (Role.TITLE, Role.BODY):
                    lines += [part.strip(" \t　 ") for part in block.text.split("\n")]
            return [line for line in lines if line]

        assert words(before) == words(after), record.source.path.name


# ---------- 第 4 步：需要确认的作业也排好版，放进「需人工确认」 ----------


def review_folder(outcome):
    folder = outcome.run_dir / "需人工确认"
    return {p.name: docx.Document(str(p)) for p in sorted(folder.iterdir())}


def test_the_assignment_without_a_title_is_formatted_and_waits_in_the_review_folder(result):
    files = review_folder(result)

    assert list(files) == ["20240325-安然.docx"]
    document = files["20240325-安然.docx"]
    roles = [b.role for b in classify(document)]
    assert Role.TITLE not in roles  # 题目没有替学生编
    assert roles[0] is Role.INFO
    assert body_texts(document, Role.INFO) == ["高二（3）班  安然  20240325"]
    assert body_texts(document, Role.HEADING) == UNIFIED
    assert page_setup(document) == [A4_PAGE]
    for block in classify(document):
        assert style_problems(block.element, STYLES[block.role], RULES) == []
    (record,) = [r for r in result.inventory.records if r.output_name == "20240325-安然.docx"]
    assert record.outcome is Outcome.REVIEW
    assert any("题目" in n for n in record.notes)


def test_without_the_teachers_decisions_the_undecided_files_wait_in_the_review_folder(tmp_path_factory):
    outcome = pipeline.run(
        REAL_INPUT, tmp_path_factory.mktemp("没有决定表"), RULES, load_roster(REAL_ROSTER)
    )

    assert sorted(review_folder(outcome)) == [
        "20240324-梁思齐（作业_梁思齐_20240324(1)）.docx",
        "20240324-梁思齐（作业_梁思齐_20240324）.docx",
        "20240325-安然.docx",
        "新建 Microsoft Word 文档.docx",
    ]
    assert outcome.inventory.count(Outcome.DONE) == 31
    anonymous = review_folder(outcome)["新建 Microsoft Word 文档.docx"]
    assert Role.INFO not in [b.role for b in classify(anonymous)]  # 没有替它编姓名行
    assert body_texts(anonymous, Role.HEADING) == UNIFIED


# ---------- 代码审查后补的：不经过程序自己"认结构"的独立核对（CR-29） ----------


def test_every_long_sentence_of_every_original_survives_word_for_word(result, outputs):
    """只用最笨的办法：原文里每一句长一点的话（正文的句子），必须原封不动、按原来的顺序
    出现在排好的文件里。不用程序的"认结构"和"认编号"，所以它们认错时这里照样能发现。"""
    for record in result.inventory.records:
        if record.outcome is not Outcome.DONE:
            continue

        def long_lines(document):
            lines = []
            for paragraph in document.element.body.iter(qn("w:p")):
                text = "".join(
                    "\n" if node.tag == qn("w:br") else (node.text or "")
                    for node in paragraph.iter(qn("w:t"), qn("w:br"))
                )
                lines += [" ".join(part.split()) for part in text.split("\n")]
            # 姓名行按规定要改写，不在比对之列：凡是带着这位学生学号的行都跳过
            return [line for line in lines if len(line) > 25 and record.student_id not in line]

        before = long_lines(docx.Document(io.BytesIO(record.docx_bytes)))
        after = long_lines(outputs[record.new_filename])
        assert len(before) >= 4, record.source.path.name  # 确实比到了东西
        assert before == after, record.source.path.name
