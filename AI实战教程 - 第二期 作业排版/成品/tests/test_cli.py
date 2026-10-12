import openpyxl
import pytest

from conftest import PROJECT_ROOT, make_docx, make_roster
from homework_formatter.__main__ import main
from homework_formatter.scanner import fingerprint

RULES = str(PROJECT_ROOT / "规则.toml")


@pytest.fixture(autouse=True)
def away_from_the_real_project(tmp_path, monkeypatch):
    """换到临时文件夹里运行，免得程序按默认位置读到真实的名单和决定表。"""
    monkeypatch.chdir(tmp_path)


def setup_project(tmp_path):
    source = tmp_path / "学生作业"
    source.mkdir()
    make_docx(source / "a.docx", ["题目", "高二（3）班  张三  20240301", "正文正文。"])
    roster = make_roster(tmp_path / "名单.xlsx", [("20240301", "张三"), ("20240302", "李四")])
    return source, roster, tmp_path / "排版结果"


def run(source, roster, output):
    return main(
        ["--input", str(source), "--roster", str(roster), "--output", str(output), "--rules", RULES]
    )


def test_run_formats_the_assignment(tmp_path):
    source, roster, output = setup_project(tmp_path)

    assert run(source, roster, output) == 0

    (run_dir,) = output.iterdir()
    assert [p.name for p in (run_dir / "已完成").iterdir()] == ["20240301-张三.docx"]


def test_missing_rules_file_is_explained(tmp_path, capsys):
    source, roster, output = setup_project(tmp_path)

    code = main(["--input", str(source), "--output", str(output), "--rules", "没有这个文件.toml"])

    assert code == 2
    assert "找不到规则表" in capsys.readouterr().out
    assert not output.exists()


def test_run_writes_a_report_into_a_new_dated_folder(tmp_path, capsys):
    source, roster, output = setup_project(tmp_path)

    assert run(source, roster, output) == 0

    (run_dir,) = output.iterdir()
    workbook = openpyxl.load_workbook(str(run_dir / "处理报告.xlsx"))
    assert workbook.sheetnames == ["总览", "明细"]
    printed = capsys.readouterr().out
    assert "李四" in printed
    assert "完全一致" in printed


def test_originals_are_untouched_and_earlier_results_are_kept(tmp_path):
    source, roster, output = setup_project(tmp_path)
    before = fingerprint(source)

    run(source, roster, output)
    run(source, roster, output)

    assert fingerprint(source) == before
    assert len(list(output.iterdir())) == 2


def test_refuses_to_write_results_inside_the_input_folder(tmp_path, capsys):
    source, roster, _ = setup_project(tmp_path)
    before = fingerprint(source)

    assert run(source, roster, source / "排版结果") == 2
    assert fingerprint(source) == before


def test_missing_input_folder_is_explained(tmp_path, capsys):
    arguments = ["--input", str(tmp_path / "不存在"), "--output", str(tmp_path / "out")]

    assert main([*arguments, "--rules", RULES]) == 2
    assert "找不到作业文件夹" in capsys.readouterr().out


def test_teacher_decisions_file_is_applied_and_shown_in_the_report(tmp_path, capsys):
    source, roster, output = setup_project(tmp_path)
    make_docx(source / "新建 Microsoft Word 文档.docx", ["题目", "正文正文。"])
    decisions = tmp_path / "老师的决定.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["原文件名", "姓名", "学号", "不要这份", "备注"])
    workbook.active.append(["新建 Microsoft Word 文档.docx", "李四", "20240302", None, None])
    workbook.active.append(["写错的文件名.docx", None, None, "是", None])
    workbook.save(str(decisions))

    code = main(
        ["--input", str(source), "--roster", str(roster), "--output", str(output),
         "--decisions", str(decisions), "--rules", RULES]
    )

    assert code == 0
    (run_dir,) = output.iterdir()
    workbook = openpyxl.load_workbook(str(run_dir / "处理报告.xlsx"))
    overview = {r[0]: r[1] for r in workbook["总览"].iter_rows(values_only=True)}
    assert "2 条" in overview["老师的决定"]
    assert "写错的文件名.docx" in overview["决定表里写了、但文件夹里没有的文件"]
    details = list(workbook["明细"].iter_rows(values_only=True))
    row = next(r for r in details if r[0] == "新建 Microsoft Word 文档.docx")
    assert "李四" in row and "老师指定" in row
    assert "写错的文件名.docx" in capsys.readouterr().out


def test_open_flag_opens_the_result_folder_when_done(tmp_path, monkeypatch):
    import os

    opened = []
    monkeypatch.setattr(os, "startfile", opened.append, raising=False)
    source, roster, output = setup_project(tmp_path)

    code = main(
        ["--input", str(source), "--roster", str(roster), "--output", str(output),
         "--rules", RULES, "--open"]
    )

    (run_dir,) = output.iterdir()
    assert code == 0
    assert [str(p) for p in opened] == [str(run_dir)]


def test_without_the_open_flag_nothing_is_opened(tmp_path, monkeypatch):
    import os

    opened = []
    monkeypatch.setattr(os, "startfile", opened.append, raising=False)
    source, roster, output = setup_project(tmp_path)

    run(source, roster, output)

    assert opened == []


def test_summary_tells_the_teacher_where_things_went(tmp_path, capsys):
    source, roster, output = setup_project(tmp_path)
    make_docx(source / "新建 Microsoft Word 文档.docx", ["题目", "正文正文。"])

    run(source, roster, output)

    printed = capsys.readouterr().out
    assert "已完成　1" in printed
    assert "需要您确认　1" in printed
    assert "需人工确认" in printed  # 提醒老师去看那个文件夹
