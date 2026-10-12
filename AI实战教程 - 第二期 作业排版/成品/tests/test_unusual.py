"""第 4 步：这批作业里没出现、但别的作业里可能有的东西。程序要认得出来、原样保留、提醒老师。"""

import docx
import pytest
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from conftest import PROJECT_ROOT
from docx_inspect import find, open_bytes, paragraph_format, save_bytes, style_problems
from homework_formatter.formatter import format_document
from homework_formatter.integrity import verify_content
from homework_formatter.rules import load_rules

RULES = load_rules(PROJECT_ROOT / "规则.toml")
TITLE = "食堂排队太久，我想这样解决"
INFO = "高二（3）班  张三  20240301"
HEADING = "一、我发现的问题"
BODY = "每天中午十二点下课，全校的人几乎同时往食堂走。我连着记了 5 天。"


def base_document():
    document = docx.Document()
    for text in (TITLE, INFO, HEADING, BODY):
        document.add_paragraph(text)
    return document


def run(document, name="张三", student_id="20240301"):
    data = save_bytes(document)
    result = format_document(data, name, student_id, RULES)
    info = RULES.info_line(name, student_id) if name else None
    assert verify_content(data, result.docx_bytes, info if result.info_rewritten else None, name, student_id) == []
    return result


def with_comment(document):
    document.add_comment(document.paragraphs[-1].runs, text="这里写得好", author="老师")


def with_tracked_change(document):
    paragraph = document.paragraphs[-1]._p
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "1")
    inserted.set(qn("w:author"), "某人")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "修订时加的一句。"
    run.append(text)
    inserted.append(run)
    paragraph.append(inserted)


def with_footnote(document):
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), "2")
    document.paragraphs[-1].add_run()._r.append(reference)


def with_text_box(document):
    box = parse_xml(
        f'<w:pict {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml">'
        "<v:shape><v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>文本框里的字</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict>"
    )
    document.add_paragraph().add_run()._r.append(box)


def with_floating_picture(document):
    drawing = parse_xml(
        f'<w:drawing {nsdecls("w", "wp")}><wp:anchor><wp:extent cx="100" cy="100"/></wp:anchor></w:drawing>'
    )
    document.add_paragraph().add_run()._r.append(drawing)


def with_two_columns(document):
    section = document.sections[0]._sectPr
    columns = section.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        section.append(columns)
    columns.set(qn("w:num"), "2")


def with_bulleted_list(document):
    document.add_paragraph("第一条办法是错开下课的时间，这样人就不会挤在一起。", style="List Bullet")


def with_content_control(document):
    control = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>控件里的字</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    )
    document.paragraphs[-1]._p.addnext(control)


@pytest.mark.parametrize(
    "add, keyword",
    [
        (with_comment, "批注"),
        (with_tracked_change, "修订"),
        (with_footnote, "脚注"),
        (with_text_box, "文本框"),
        (with_floating_picture, "浮动"),
        (with_two_columns, "分栏"),
        (with_bulleted_list, "自动编号"),
        (with_content_control, "内容控件"),
    ],
)
def test_unusual_content_is_reported_and_the_rest_is_still_formatted(add, keyword):
    document = base_document()
    add(document)

    result = run(document)

    assert any(keyword in w for w in result.warnings), result.warnings
    output = open_bytes(result.docx_bytes)
    assert style_problems(find(output, BODY), RULES.body, RULES) == []
    assert style_problems(find(output, HEADING), RULES.heading, RULES) == []


def test_an_ordinary_document_raises_no_warnings():
    assert run(base_document()).warnings == []


def test_bulleted_paragraph_keeps_its_bullet_and_its_own_indent():
    document = base_document()
    with_bulleted_list(document)

    output = open_bytes(run(document).docx_bytes)

    bullet = find(output, "第一条办法是错开下课的时间，这样人就不会挤在一起。")
    assert bullet.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None
    assert paragraph_format(bullet)["first_line"] is None  # 缩进交给项目符号，不硬套首行缩进


def test_words_added_in_revision_mode_are_kept():
    document = base_document()
    with_tracked_change(document)

    output = open_bytes(run(document).docx_bytes)

    assert "修订时加的一句。" in "".join(t.text for t in output.element.body.iter(qn("w:t")))


# ---------- 拿不准是谁的作业：不改姓名行 ----------


def test_without_a_known_owner_the_students_own_info_line_is_left_as_written():
    document = docx.Document()
    for text in (TITLE, "姓名：某某    学号：20249999", HEADING, BODY):
        document.add_paragraph(text)

    result = run(document, name=None, student_id=None)

    output = open_bytes(result.docx_bytes)
    info = find(output, "姓名：某某    学号：20249999")
    assert style_problems(info, RULES.info, RULES) == []  # 格式照样统一
    assert not any("姓名行" in c for c in result.changes)


def test_without_a_known_owner_no_info_line_is_invented():
    document = docx.Document()
    for text in (TITLE, HEADING, BODY):
        document.add_paragraph(text)

    result = run(document, name=None, student_id=None)

    texts = [p.text for p in open_bytes(result.docx_bytes).paragraphs]
    assert texts == [TITLE, HEADING, BODY]


def test_without_a_known_owner_a_header_holding_a_name_is_not_wiped():
    document = docx.Document()
    for text in (TITLE, HEADING, BODY):
        document.add_paragraph(text)
    document.sections[0].header.paragraphs[0].text = "高二（3）班  某某  20249999"

    result = run(document, name=None, student_id=None)

    output = open_bytes(result.docx_bytes)
    assert output.sections[0].header.paragraphs[0].text == "高二（3）班  某某  20249999"
    assert any("页眉" in w for w in result.warnings)
