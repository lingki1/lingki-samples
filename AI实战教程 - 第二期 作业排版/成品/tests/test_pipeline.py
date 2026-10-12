import dataclasses
import io

import docx
import openpyxl
import pytest

from conftest import PROJECT_ROOT, make_docx
from docx_inspect import body_paragraphs, page_setup, paragraph_text
from homework_formatter import pipeline
from homework_formatter.decisions import Decisions

from homework_formatter.inventory import Outcome
from homework_formatter.roster import Roster
from homework_formatter.rules import load_rules
from homework_formatter.scanner import fingerprint

RULES = load_rules(PROJECT_ROOT / "规则.toml")
ROSTER = Roster([("20240301", "张三"), ("20240302", "李四"), ("20240303", "王五")])
BODY = ["一、问题", "　　这是正文的第一段，写了一些内容。", "", "这是正文的第二段。"]


@pytest.fixture
def source(tmp_path):
    folder = tmp_path / "学生作业"
    folder.mkdir()
    make_docx(folder / "随便起的名字.docx", ["我的题目", "姓名：张三    学号：20240301", *BODY])
    make_docx(folder / "李四.docx", ["李四的题目", *BODY])
    return folder


def run(source, tmp_path, **kwargs):
    return pipeline.run(source, tmp_path / "排版结果", RULES, roster=ROSTER, **kwargs)


def by_name(result):
    return {r.source.path.name: r for r in result.inventory.records}


def test_ready_assignments_are_formatted_into_the_done_folder_under_their_new_names(source, tmp_path):
    result = run(source, tmp_path)

    done = result.run_dir / "已完成"
    assert sorted(p.name for p in done.iterdir()) == ["20240301-张三.docx", "20240302-李四.docx"]
    assert result.inventory.count(Outcome.DONE) == 2
    document = docx.Document(str(done / "20240301-张三.docx"))
    assert page_setup(document)[0]["width"] == "11906"
    assert [paragraph_text(p) for p in body_paragraphs(document)] == [
        "我的题目",
        "高二（3）班  张三  20240301",
        "一、问题",
        "这是正文的第一段，写了一些内容。",
        "这是正文的第二段。",
    ]


def test_changes_are_recorded_for_the_report(source, tmp_path):
    record = by_name(run(source, tmp_path))["李四.docx"]

    assert any("补上了姓名行" in c for c in record.changes)
    assert any("空段落" in c for c in record.changes)


def test_originals_are_untouched(source, tmp_path):
    before = fingerprint(source)

    result = run(source, tmp_path)

    assert fingerprint(source) == before
    assert result.originals_untouched
    assert result.file_count == 2


def review_files(result):
    folder = result.run_dir / "需人工确认"
    return sorted(p.name for p in folder.iterdir()) if folder.is_dir() else None


def texts_of(path):
    return [paragraph_text(p) for p in body_paragraphs(docx.Document(str(path)))]


def test_when_nothing_needs_review_there_is_no_review_folder(source, tmp_path):
    assert review_files(run(source, tmp_path)) is None


def test_unprocessed_files_are_written_nowhere(source, tmp_path):
    (source / "王五.pdf").write_bytes(b"%PDF-1.7")

    result = run(source, tmp_path)

    assert len(list((result.run_dir / "已完成").iterdir())) == 2
    assert review_files(result) is None
    assert by_name(result)["王五.pdf"].outcome is Outcome.UNPROCESSED
    assert by_name(result)["王五.pdf"].output_name is None


def test_anonymous_assignment_is_formatted_into_the_review_folder_under_its_own_name(source, tmp_path):
    make_docx(source / "新建 Microsoft Word 文档.docx", ["匿名的题目", "1. 问题", *BODY[1:]])

    result = run(source, tmp_path)

    record = by_name(result)["新建 Microsoft Word 文档.docx"]
    assert record.outcome is Outcome.REVIEW
    assert review_files(result) == ["新建 Microsoft Word 文档.docx"]
    assert record.output_name == "新建 Microsoft Word 文档.docx"
    path = result.run_dir / "需人工确认" / record.output_name
    # 格式排好了，但没有替它编一个姓名行
    assert texts_of(path) == ["匿名的题目", "一、问题", "这是正文的第一段，写了一些内容。", "这是正文的第二段。"]
    assert page_setup(docx.Document(str(path)))[0]["width"] == "11906"
    assert len(list((result.run_dir / "已完成").iterdir())) == 2


def test_assignment_without_a_title_is_formatted_into_the_review_folder_under_its_new_name(source, tmp_path):
    make_docx(source / "王五的作业.docx", ["姓名：王五 学号：20240303", "我发现的问题", *BODY[1:]])

    result = run(source, tmp_path)

    record = by_name(result)["王五的作业.docx"]
    assert record.outcome is Outcome.REVIEW
    assert review_files(result) == ["20240303-王五.docx"]
    assert texts_of(result.run_dir / "需人工确认" / "20240303-王五.docx")[:2] == [
        "高二（3）班  王五  20240303",
        "一、我发现的问题",
    ]
    assert any("题目" in n for n in record.notes)


def test_two_versions_from_one_student_both_go_to_review_with_names_that_tell_them_apart(source, tmp_path):
    make_docx(source / "张三(1).docx", ["我的题目", "姓名：张三    学号：20240301", "一、问题", "不一样的正文。"])

    result = run(source, tmp_path)

    assert review_files(result) == [
        "20240301-张三（张三(1)）.docx",
        "20240301-张三（随便起的名字）.docx",
    ]
    assert [p.name for p in (result.run_dir / "已完成").iterdir()] == ["20240302-李四.docx"]


def test_identity_in_doubt_keeps_the_original_name_and_the_students_own_info_line(source, tmp_path):
    make_docx(source / "对不上.docx", ["一个题目", "高二（3）班  王五  20240301", *BODY])

    result = run(source, tmp_path)

    record = by_name(result)["对不上.docx"]
    assert record.output_name == "对不上.docx"
    assert texts_of(result.run_dir / "需人工确认" / "对不上.docx")[1] == "高二（3）班  王五  20240301"


def test_old_doc_in_review_is_saved_as_docx(source, tmp_path):
    (source / "旧版.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
    converted = make_docx(tmp_path / "converted.docx", ["匿名的题目", *BODY])

    def fake_convert(paths):
        from homework_formatter.converter import ConversionResult

        return ConversionResult(converted={p: converted.read_bytes() for p in paths}, failed={})

    result = run(source, tmp_path, convert=fake_convert)

    assert review_files(result) == ["旧版.docx"]


def test_teacher_decisions_flow_through_to_the_output(source, tmp_path):
    make_docx(source / "新建 Microsoft Word 文档.docx", ["匿名的题目", *BODY])
    decisions = Decisions(assigned={"新建 Microsoft Word 文档.docx": ("王五", "20240303")})

    result = run(source, tmp_path, decisions=decisions)

    document = docx.Document(str(result.run_dir / "已完成" / "20240303-王五.docx"))
    assert paragraph_text(body_paragraphs(document)[1]) == "高二（3）班  王五  20240303"


def test_a_file_whose_words_changed_during_formatting_is_held_back(source, tmp_path, monkeypatch):
    real = pipeline.format_document

    def losing_formatter(docx_bytes, name, student_id, rules):
        result = real(docx_bytes, name, student_id, rules)
        if name != "李四":
            return result
        document = docx.Document(io.BytesIO(result.docx_bytes))
        last = document.paragraphs[-1]._element
        last.getparent().remove(last)
        buffer = io.BytesIO()
        document.save(buffer)
        return dataclasses.replace(result, docx_bytes=buffer.getvalue())

    monkeypatch.setattr(pipeline, "format_document", losing_formatter)

    result = run(source, tmp_path)

    records = by_name(result)
    assert records["李四.docx"].outcome is Outcome.REVIEW
    assert any("文字不一致" in n for n in records["李四.docx"].notes)
    assert any("没有生成" in n for n in records["李四.docx"].notes)
    assert review_files(result) is None  # 内容对不上的，哪个文件夹都不放
    assert records["李四.docx"].output_name is None
    assert [p.name for p in (result.run_dir / "已完成").iterdir()] == ["20240301-张三.docx"]
    assert records["随便起的名字.docx"].outcome is Outcome.DONE


def test_formatter_warnings_send_the_file_to_review(source, tmp_path):
    make_docx(
        source / "王五.docx",
        ["王五的题目", "王五 20240303", "一、问题", "1. 现象", "正文一段，有标点。", "二、办法", "正文二段，有标点。"],
    )

    result = run(source, tmp_path)
    record = by_name(result)["王五.docx"]

    assert record.outcome is Outcome.REVIEW
    assert any("编号" in n for n in record.notes)
    assert review_files(result) == ["20240303-王五.docx"]
    assert "20240303-王五.docx" not in [p.name for p in (result.run_dir / "已完成").iterdir()]


def test_report_lists_the_headings_as_they_are_after_formatting(source, tmp_path):
    make_docx(source / "王五.docx", ["王五的题目", "王五 20240303", "1. 问题", "正文一段，有标点。"])

    record = by_name(run(source, tmp_path))["王五.docx"]

    assert record.structure.headings == ["一、问题"]
    assert any("编号" in c for c in record.changes)


def test_a_file_that_fails_quality_control_never_reaches_the_done_folder(source, tmp_path, monkeypatch):
    real = pipeline.format_document

    def sloppy_formatter(docx_bytes, name, student_id, rules):
        result = real(docx_bytes, name, student_id, rules)
        if name != "李四":
            return result
        document = docx.Document(io.BytesIO(result.docx_bytes))
        document.paragraphs[-1].runs[0].font.bold = True  # 内容没变，格式不对
        buffer = io.BytesIO()
        document.save(buffer)
        return dataclasses.replace(result, docx_bytes=buffer.getvalue())

    monkeypatch.setattr(pipeline, "format_document", sloppy_formatter)

    result = run(source, tmp_path)

    records = by_name(result)
    assert records["李四.docx"].outcome is Outcome.REVIEW
    assert any("质检没通过" in n and "加粗" in n for n in records["李四.docx"].notes)
    assert review_files(result) == ["20240302-李四.docx"]
    assert [p.name for p in (result.run_dir / "已完成").iterdir()] == ["20240301-张三.docx"]
    assert records["随便起的名字.docx"].quality == []


def test_quality_control_result_is_recorded_and_reported(source, tmp_path):
    make_docx(source / "王五的作业.docx", ["姓名：王五 学号：20240303", "我发现的问题", *BODY[1:]])
    make_docx(
        source / "混编.docx",
        ["混编的题目", "李四 20240302", "一、问题", "1. 现象", "正文一段，有标点。", "二、办法", "正文二段，有标点。"],
    )

    result = run(source, tmp_path)

    records = by_name(result)
    assert records["李四.docx"].quality == []
    assert records["王五的作业.docx"].quality == []  # 缺题目是已知的，其余格式合格
    assert records["混编.docx"].quality is None  # 有程序没动的地方，不下质检结论
    workbook = openpyxl.load_workbook(str(result.report_path))
    rows = list(workbook["明细"].iter_rows(values_only=True))
    quality = {r[0]: r[rows[0].index("质检")] for r in rows[1:]}
    assert quality["随便起的名字.docx"] == "通过"
    assert quality["王五的作业.docx"] == "通过"
    assert quality["混编.docx"] is None
    overview = {r[0]: r[1] for r in workbook["总览"].iter_rows(values_only=True)}
    assert "全部通过" in overview["质检"]


def test_one_file_blowing_up_does_not_stop_the_others(source, tmp_path, monkeypatch):
    real = pipeline.format_document

    def fragile_formatter(docx_bytes, name, student_id, rules):
        if name == "李四":
            raise ValueError("boom")
        return real(docx_bytes, name, student_id, rules)

    monkeypatch.setattr(pipeline, "format_document", fragile_formatter)

    result = run(source, tmp_path)

    records = by_name(result)
    assert records["李四.docx"].outcome is Outcome.UNPROCESSED
    assert any("排版时出错" in n for n in records["李四.docx"].notes)
    assert records["随便起的名字.docx"].outcome is Outcome.DONE


def test_each_run_gets_its_own_folder_and_a_report(source, tmp_path):
    first = run(source, tmp_path)
    second = run(source, tmp_path)

    assert first.run_dir != second.run_dir
    workbook = openpyxl.load_workbook(str(second.report_path))
    rows = list(workbook["明细"].iter_rows(values_only=True))
    header = rows[0]
    row = dict(zip(header, next(r for r in rows if r[0] == "李四.docx")))
    assert row["去向"] == "已完成"
    assert "补上了姓名行" in row["做了哪些改动"]
    overview = {r[0]: r[1] for r in workbook["总览"].iter_rows(values_only=True)}
    assert overview["已完成"] == 2
    assert "可以排版" not in overview


def test_results_cannot_go_inside_the_input_folder(source):
    with pytest.raises(pipeline.PipelineError, match="作业文件夹"):
        pipeline.run(source, source / "排版结果", RULES)
    assert sorted(p.name for p in source.iterdir()) == ["李四.docx", "随便起的名字.docx"]
