import io

import docx
import pytest

from homework_formatter.structure import Role, analyze, classify, is_heading_text

TITLE = "食堂排队太久，我想这样解决"
INFO = "高二（3）班  张三  20240301"
BODY_1 = "每天中午十二点下课，全校的人几乎同时往食堂走。我连着记了五天。"
BODY_2 = "排队的时间一长，午休就只剩下不到半小时。"


def roles(document):
    return [(b.role, b.text) for b in classify(document)]


def doc_with(*paragraphs):
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    return document


def test_plain_layout_title_info_heading_body():
    document = doc_with(TITLE, INFO, "一、我发现的问题", BODY_1, BODY_2)

    assert roles(document) == [
        (Role.TITLE, TITLE),
        (Role.INFO, INFO),
        (Role.HEADING, "一、我发现的问题"),
        (Role.BODY, BODY_1),
        (Role.BODY, BODY_2),
    ]


@pytest.mark.parametrize(
    "text",
    [
        "一、我发现的问题",
        "（一）我发现的问题",
        "(二) 原因分析",
        "1. 我发现的问题",
        "2、原因分析",
        "３．我的办法",
        "第一部分 我发现的问题",
        "我发现的问题",
        "原因分析",
        "参考资料",
        "我的办法：",
    ],
)
def test_lines_that_are_headings(text):
    assert is_heading_text(text)


@pytest.mark.parametrize(
    "text",
    [
        BODY_1,
        "表贴在教室最后面，平时没有人会特意走过去看。",
        "第一，在地上重新画斜向的停车线，斜着停比横着停省地方。",
        "[1] 本人的现场观察记录表（见附页）.",
        "[2] 班主任提供的本班考勤与扣分记录, 2024.",
        "1. 三个年级错开十分钟下课，高一先走，高三最后，这样食堂门口就不会堵住。",
        "还有不少同学用的是一升以上的大水壶一个人接满就要将近一分钟",  # 没有标点，但超过 20 个字
        "",
    ],
)
def test_lines_that_are_not_headings(text):
    assert not is_heading_text(text)


def test_word_heading_style_marks_a_heading_whatever_the_text():
    document = doc_with(TITLE, INFO)
    document.add_paragraph("这一行很长而且有逗号，但是学生用了标题样式。", style="Heading 2")
    document.add_paragraph(BODY_1)

    assert roles(document)[2][0] is Role.HEADING


def test_title_written_with_a_heading_style_is_still_the_title():
    document = docx.Document()
    document.add_paragraph(TITLE, style="Heading 1")
    document.add_paragraph(INFO)

    assert roles(document)[0] == (Role.TITLE, TITLE)


def test_blank_lines_are_marked_and_do_not_hide_the_info_line():
    document = doc_with(TITLE, "", INFO, "", "我发现的问题", BODY_1, "", BODY_2)

    assert [r for r, _ in roles(document)] == [
        Role.TITLE,
        Role.EMPTY,
        Role.INFO,
        Role.EMPTY,
        Role.HEADING,
        Role.BODY,
        Role.EMPTY,
        Role.BODY,
    ]


def test_first_line_being_the_info_line_means_no_title():
    document = doc_with(INFO, "我发现的问题", BODY_1)

    assert [r for r, _ in roles(document)] == [Role.INFO, Role.HEADING, Role.BODY]
    assert analyze_document(document).title is None


def test_first_line_being_a_numbered_heading_or_a_long_sentence_means_no_title():
    numbered = doc_with("一、我发现的问题", BODY_1)
    sentence = doc_with(BODY_1, BODY_2)

    assert roles(numbered)[0][0] is Role.HEADING
    assert roles(sentence)[0][0] is Role.BODY


def test_no_info_line_at_all():
    document = doc_with(TITLE, "一、我发现的问题", BODY_1)

    assert [r for r, _ in roles(document)] == [Role.TITLE, Role.HEADING, Role.BODY]


def test_info_line_deep_in_the_text_is_just_body():
    document = doc_with(TITLE, "一、问题", BODY_1, BODY_2, BODY_1, BODY_2, "联系人 20240301")

    assert roles(document)[-1][0] is not Role.INFO


def test_paragraph_with_a_soft_line_break_is_body_even_if_short():
    document = doc_with(TITLE, INFO)
    paragraph = document.add_paragraph()
    run = paragraph.add_run("原因分析")
    run.add_break()
    run.add_text("我的办法")

    assert roles(document)[2][0] is Role.BODY


def test_tables_and_pictures(tiny_png):
    document = doc_with(TITLE, INFO, BODY_1)
    document.add_picture(io.BytesIO(tiny_png))
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "日期"
    document.add_paragraph(BODY_2)

    assert [r for r, _ in roles(document)] == [
        Role.TITLE,
        Role.INFO,
        Role.BODY,
        Role.IMAGE,
        Role.TABLE,
        Role.BODY,
    ]


def test_summary_lists_title_and_headings_in_order():
    document = doc_with(TITLE, INFO, "一、问题", BODY_1, "二、办法", BODY_2, "参考资料", "[1] 某资料.")

    structure = analyze_document(document)

    assert structure.title == TITLE
    assert structure.headings == ["一、问题", "二、办法", "参考资料"]


def analyze_document(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return analyze(buffer.getvalue())
