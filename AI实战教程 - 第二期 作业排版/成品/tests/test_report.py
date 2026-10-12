import openpyxl

from conftest import make_docx
from homework_formatter.inventory import take_inventory
from homework_formatter.report import DETAIL_SHEET, OVERVIEW_SHEET, write_report
from homework_formatter.roster import Roster

ROSTER = Roster([("20240301", "张三"), ("20240302", "李四"), ("20240303", "王五")])


def build(tmp_path, originals_untouched=True, roster=ROSTER):
    source = tmp_path / "作业"
    source.mkdir()
    make_docx(
        source / "b.docx",
        ["我的题目", "高二（3）班  张三  20240301", "一、问题", "正文正文。", "参考资料", "[1] 某资料."],
    )
    make_docx(source / "新建 Microsoft Word 文档.docx", ["题目", "正文正文。"])
    (source / "李四.pdf").write_bytes(b"%PDF-1.7")
    (source / "~$b.docx").write_bytes(b"\x05x" + b"\0" * 100)
    inventory = take_inventory(source, roster)
    path = write_report(
        tmp_path / "处理报告.xlsx",
        inventory,
        input_dir=source,
        roster=roster,
        originals_untouched=originals_untouched,
        file_count=4,
    )
    return openpyxl.load_workbook(str(path))


def detail_rows(workbook):
    rows = list(workbook[DETAIL_SHEET].iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, row)) for row in rows[1:]]


def overview(workbook):
    return {row[0]: row[1] for row in workbook[OVERVIEW_SHEET].iter_rows(values_only=True) if row[0]}


def test_one_row_per_file_with_problems_listed_first(tmp_path):
    rows = detail_rows(build(tmp_path))

    assert [r["原文件名"] for r in rows] == [
        "新建 Microsoft Word 文档.docx",
        "李四.pdf",
        "b.docx",
        "~$b.docx",
    ]
    assert [r["去向"] for r in rows] == ["需要您确认", "未处理", "可以排版", "已忽略"]


def test_row_contents(tmp_path):
    rows = {r["原文件名"]: r for r in detail_rows(build(tmp_path))}

    ready = rows["b.docx"]
    assert (ready["姓名"], ready["学号"]) == ("张三", "20240301")
    assert (ready["姓名来自"], ready["学号来自"]) == ("正文", "正文")
    assert ready["新文件名"] == "20240301-张三.docx"
    assert ready["类型"] == "Word 文档"
    assert ready["认出的题目"] == "我的题目"
    assert ready["认出的小标题"] == "一、问题 | 参考资料"
    assert "PDF" in rows["李四.pdf"]["说明"]
    assert rows["李四.pdf"]["学号"] == "20240302"
    assert rows["新建 Microsoft Word 文档.docx"]["新文件名"] is None


def test_student_ids_are_stored_as_text_so_excel_does_not_mangle_them(tmp_path):
    workbook = build(tmp_path)
    sheet = workbook[DETAIL_SHEET]
    header = [c.value for c in sheet[1]]
    id_cells = [row[header.index("学号")] for row in sheet.iter_rows(min_row=2)]

    assert all(isinstance(c.value, str) for c in id_cells if c.value is not None)


def test_overview_counts_and_missing_students(tmp_path):
    summary = overview(build(tmp_path))

    assert summary["文件总数"] == 4
    assert summary["可以排版"] == 1
    assert summary["需要您确认"] == 1
    assert summary["未处理"] == 1
    assert summary["已忽略"] == 1
    assert "完全一致" in summary["原件核对"]
    assert "3 人" in summary["班级名单"]
    assert "王五" in summary["名单上有、但没找到作业的学生"]
    assert "李四" not in summary["名单上有、但没找到作业的学生"]
    assert "李四" in summary["交了但没法处理的学生"]


def test_changed_originals_are_shouted_about(tmp_path):
    summary = overview(build(tmp_path, originals_untouched=False))

    assert "不一致" in summary["原件核对"]


def test_without_roster_the_overview_says_so(tmp_path):
    summary = overview(build(tmp_path, roster=None))

    assert "没有提供" in summary["班级名单"]
    assert "名单上有、但没找到作业的学生" not in summary
