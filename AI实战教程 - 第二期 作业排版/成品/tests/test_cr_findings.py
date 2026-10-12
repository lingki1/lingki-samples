"""代码审查（CR）发现的问题，每条一组复现用的检查。编号对应 文档/4-代码审查.md。

这些检查在修复前的代码上不通过，修复后通过。以后改代码时它们防止老问题回来。
"""

import dataclasses
import io
import zipfile

import docx
import openpyxl
import pytest
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt

from conftest import PROJECT_ROOT, make_docx, make_roster
from docx_inspect import find, open_bytes, paragraph_format, run_format, save_bytes, text_runs
from homework_formatter import converter, pipeline
from homework_formatter.__main__ import main
from homework_formatter.checker import check_document
from homework_formatter.decisions import DecisionsError, load_decisions
from homework_formatter.formatter import format_document
from homework_formatter.identity import info_line_kind, name_from_filename, parse_info_line
from homework_formatter.integrity import verify_content
from homework_formatter.inventory import Outcome, take_inventory
from homework_formatter.numbering import plan_numbering
from homework_formatter.roster import Roster, load_roster
from homework_formatter.rules import RulesError, load_rules
from homework_formatter.scanner import FileKind, fingerprint, scan_folder
from homework_formatter.structure import Role, classify

RULES_FILE = PROJECT_ROOT / "规则.toml"
RULES = load_rules(RULES_FILE)
ROSTER = Roster(
    [("20240301", "张三"), ("20240302", "李四"), ("20240303", "王五"), ("20240321", "童年"), ("20240325", "安然")]
)
TITLE = "食堂排队太久，我想这样解决"
INFO = "高二（3）班  张三  20240301"
HEADING = "一、我发现的问题"
BODY = "每天中午十二点下课，全校的人几乎同时往食堂走。我连着记了 5 天。"
BODY_2 = "排队的时间一长，午休就只剩下不到半小时。"
MATH_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
FORMULA = f"<m:oMathPara {MATH_NS}><m:oMath><m:r><m:t>v=s/t=120/26</m:t></m:r></m:oMath></m:oMathPara>"


def doc(*paragraphs):
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    return document


def normal_doc():
    return doc(TITLE, "张三 20240301", HEADING, BODY, BODY_2)


def fmt(document, name="张三", student_id="20240301", rules=RULES):
    return format_document(save_bytes(document), name, student_id, rules)


def verify(document, result, name="张三", student_id="20240301"):
    info = RULES.info_line(name, student_id) if result.info_rewritten else None
    return verify_content(save_bytes(document), result.docx_bytes, info, name, student_id)


def all_text(result):
    body = open_bytes(result.docx_bytes).element.body
    return "".join(t.text or "" for t in body.iter(qn("w:t")))


def paragraph_texts(result):
    return [p.text for p in open_bytes(result.docx_bytes).paragraphs]


def run_pipeline(tmp_path, files: dict, roster=ROSTER, **kwargs):
    source = tmp_path / "学生作业"
    source.mkdir(exist_ok=True)
    for name, content in files.items():
        if isinstance(content, bytes):
            (source / name).write_bytes(content)
        elif isinstance(content, list):
            make_docx(source / name, content)
        else:
            content.save(str(source / name))
    return pipeline.run(source, tmp_path / "排版结果", RULES, roster=roster, **kwargs)


def records(result):
    return {r.source.path.name: r for r in result.inventory.records}


def folder(result, name):
    path = result.run_dir / name
    return sorted(p.name for p in path.iterdir()) if path.is_dir() else []


# ====================== CR-01 独占一行的公式、符号被当成空行删掉 ======================


def add_xml_paragraph(document, inner_xml):
    paragraph = document.add_paragraph()
    paragraph._p.append(parse_xml(inner_xml))
    return paragraph


def test_cr01_a_formula_on_its_own_line_is_kept():
    document = normal_doc()
    add_xml_paragraph(document, FORMULA)
    document.add_paragraph(BODY_2)

    result = fmt(document)

    body = open_bytes(result.docx_bytes).element.body
    assert sum(1 for _ in body.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")) == 1
    assert not any("空段落" in c for c in result.changes)
    assert any("公式" in w for w in result.warnings)  # 提醒老师，转人工确认
    assert verify(document, result) == []


def test_cr01_a_line_holding_only_an_inserted_symbol_or_a_footnote_mark_is_kept():
    for inner in (
        f'<w:r {nsdecls("w")}><w:sym w:font="Wingdings" w:char="F0FC"/></w:r>',
        f'<w:r {nsdecls("w")}><w:footnoteReference w:id="2"/></w:r>',
    ):
        document = normal_doc()
        add_xml_paragraph(document, inner)

        result = fmt(document)

        assert not any("空段落" in c for c in result.changes), inner
        assert verify(document, result) == []


def test_cr01_the_comparison_notices_a_lost_formula_or_symbol():
    document = normal_doc()
    add_xml_paragraph(document, FORMULA)
    add_xml_paragraph(document, f'<w:r {nsdecls("w")}><w:sym w:font="Symbol" w:char="F061"/></w:r>')
    result = fmt(document)

    damaged = open_bytes(result.docx_bytes)
    for tag in ("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara", qn("w:sym")):
        element = next(damaged.element.body.iter(tag))
        element.getparent().remove(element)
    problems = verify_content(save_bytes(document), save_bytes(damaged), INFO, "张三", "20240301")

    assert any("公式" in p for p in problems)
    assert any("符号" in p for p in problems)


def test_cr01_text_typed_entirely_in_revision_mode_is_not_treated_as_an_empty_line():
    document = normal_doc()
    paragraph = document.add_paragraph()
    inserted = parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="x"><w:r><w:t>修订时整段加的话。</w:t></w:r></w:ins>'
    )
    paragraph._p.append(inserted)

    result = fmt(document)

    assert "修订时整段加的话。" in all_text(result)
    assert verify(document, result) == []


# ====================== CR-02 姓名行认得太宽，整行换掉时丢了别的内容 ======================


@pytest.mark.parametrize(
    "line",
    [
        "调查的班级：高二3班和高二5班，共90人。",
        "联系电话：88886666，欢迎老师同学提意见。",
        "完成日期：20241008",
        "调查班级：高二（1）班到高二（6）班",
    ],
)
def test_cr02_an_ordinary_line_near_the_top_is_not_mistaken_for_the_info_line(line):
    document = doc(TITLE, "张三 20240301", line, HEADING, BODY)

    result = fmt(document)

    assert line in paragraph_texts(result)
    assert paragraph_texts(result)[:2] == [TITLE, INFO]
    assert result.warnings == []
    assert verify(document, result) == []


def test_cr02_info_line_with_extra_words_is_not_rewritten_and_goes_to_review(tmp_path):
    line = "姓名：张三 学号：20240301 指导老师：王建国"

    result = run_pipeline(tmp_path, {"a.docx": [TITLE, line, HEADING, BODY]})

    record = records(result)["a.docx"]
    assert (record.name, record.student_id) == ("张三", "20240301")  # 人认对了
    assert record.outcome is Outcome.REVIEW
    assert any("指导老师" in n for n in record.notes)
    assert folder(result, "已完成") == []
    kept = docx.Document(str(result.run_dir / "需人工确认" / record.output_name))
    assert line in [p.text for p in kept.paragraphs]  # 一个字没动


def test_cr02_a_second_line_naming_a_group_member_is_not_deleted(tmp_path):
    result = run_pipeline(
        tmp_path, {"a.docx": [TITLE, "张三 20240301", "同组：李四 20240302", HEADING, BODY]}
    )

    record = records(result)["a.docx"]
    assert record.outcome is Outcome.REVIEW
    assert any("不止一个学号" in n for n in record.notes)
    kept = docx.Document(str(result.run_dir / "需人工确认" / record.output_name))
    assert "同组：李四 20240302" in [p.text for p in kept.paragraphs]


def test_cr02_two_students_on_one_line_are_not_silently_reduced_to_the_first(tmp_path):
    result = run_pipeline(
        tmp_path, {"a.docx": [TITLE, "张三 20240301 李四 20240302", HEADING, BODY]}
    )

    record = records(result)["a.docx"]
    assert record.outcome is Outcome.REVIEW
    assert folder(result, "已完成") == []
    kept = docx.Document(str(result.run_dir / "需人工确认" / record.output_name))
    assert "张三 20240301 李四 20240302" in [p.text for p in kept.paragraphs]


def test_cr02_a_title_that_starts_with_the_class_name_stays_the_title():
    for title in ("高二（3）班睡眠调查", "高二（3）班 睡眠情况调查"):
        document = doc(title, "张三 20240301", HEADING, BODY)

        assert [b.role for b in classify(document)][:2] == [Role.TITLE, Role.INFO], title
        assert info_line_kind(title) is None


def test_cr02_title_and_info_line_squeezed_into_one_paragraph_lose_nothing():
    document = docx.Document()
    run = document.add_paragraph().add_run(TITLE)
    run.add_break()
    run.add_text("张三 20240301")
    document.add_paragraph(HEADING)
    document.add_paragraph(BODY)

    result = fmt(document)

    assert paragraph_texts(result)[:3] == [TITLE, INFO, HEADING]
    assert result.title == TITLE
    assert verify(document, result) == []


def test_cr02_extra_words_in_the_header_keep_the_header():
    document = doc(TITLE, HEADING, BODY)
    document.sections[0].header.paragraphs[0].text = "高二（3）班  张三  20240301  指导老师：王芳"

    result = fmt(document)

    output = open_bytes(result.docx_bytes)
    assert "指导老师：王芳" in output.sections[0].header.paragraphs[0].text
    assert any("页眉" in w and "指导老师" in w for w in result.warnings)


def test_cr02_the_comparison_itself_refuses_an_info_line_swallowing_other_words():
    """比对不再依赖"认结构"：就算排版把夹着别的字的一行整行换掉，比对也要报出来。"""
    original = doc(TITLE, "姓名：张三 学号：20240301 指导老师：王建国", HEADING, BODY)
    swallowed = doc(TITLE, INFO, HEADING, BODY)

    problems = verify_content(save_bytes(original), save_bytes(swallowed), INFO, "张三", "20240301")

    assert any("指导老师" in p for p in problems)


# ====================== CR-03 文件名里碰巧含有名单上的名字就认人 ======================


@pytest.mark.parametrize(
    "filename, expected",
    [
        ("我的童年趣事调查.docx", None),
        ("安然无恙的秘密 研究报告.docx", None),
        ("童年的报告.docx", "童年"),
        ("高二3班 安然 小报告.docx", "安然"),
        ("20240321童年.docx", "童年"),
        ("作业_安然_20240325(1).docx", "安然"),
    ],
)
def test_cr03_a_roster_name_must_stand_alone_in_the_filename(filename, expected):
    assert name_from_filename(filename, ROSTER) == expected


def test_cr03_a_roster_name_inside_a_longer_name_is_not_that_student():
    roster = Roster([("20240301", "张伟")])

    assert parse_info_line("姓名：张伟明", roster) == ("张伟明", None)
    assert parse_info_line("张伟 20240301", roster) == ("张伟", "20240301")


def test_cr03_an_essay_about_childhood_is_not_filed_under_the_student_named_tongnian(tmp_path):
    result = run_pipeline(tmp_path, {"我的童年趣事调查.docx": [TITLE, HEADING, BODY]})

    record = records(result)["我的童年趣事调查.docx"]
    assert record.name is None
    assert record.outcome is Outcome.REVIEW
    assert folder(result, "已完成") == []
    assert ("20240321", "童年") in result.inventory.missing_students


def test_cr03_roster_names_typed_with_a_space_in_the_middle_still_match(tmp_path):
    path = make_roster(tmp_path / "名单.xlsx", [("20240301", "张　三"), ("20240302", "李 四")])

    roster = load_roster(path)

    assert roster.id_of("张三") == "20240301" and roster.id_of("李四") == "20240302"


# ====================== CR-04 "谁没交"算错 ======================


def test_cr04_a_wrong_id_does_not_mark_its_real_owner_as_having_submitted(tmp_path):
    # 张三把学号写成了李四的；李四根本没交
    result = run_pipeline(tmp_path, {"x.docx": [TITLE, "张三 20240302", HEADING, BODY]})

    inventory = result.inventory
    assert inventory.missing_students == ROSTER.entries  # 身份没确认，谁都不算交了
    assert inventory.possible_files["20240301"] == ["x.docx"]
    assert inventory.possible_files["20240302"] == ["x.docx"]
    overview = {r[0]: r[1] for r in openpyxl.load_workbook(str(result.report_path))["总览"].iter_rows(values_only=True)}
    assert "张三（20240301，可能是「x.docx」，待确认）" in overview["名单上有、但没找到作业的学生"]


# ====================== CR-05 一个坏文件让整批中断 ======================


def test_cr05_a_file_that_breaks_the_reader_does_not_stop_the_others(tmp_path, monkeypatch):
    from homework_formatter import inventory

    real = inventory.read_docx

    def fragile(data):
        if b"BROKEN" in data:
            raise KeyError("rId999")
        return real(data)

    monkeypatch.setattr(inventory, "read_docx", fragile)
    good = io.BytesIO()
    normal_doc().save(good)
    broken = good.getvalue() + b"BROKEN"

    result = run_pipeline(tmp_path, {"好的.docx": normal_doc(), "坏的.docx": broken})

    assert records(result)["坏的.docx"].outcome is Outcome.UNPROCESSED
    assert any("出错" in n for n in records(result)["坏的.docx"].notes)
    assert records(result)["好的.docx"].outcome is Outcome.DONE
    assert result.report_path.is_file()


def test_cr05_a_docx_whose_header_part_is_missing_is_reported_not_fatal(tmp_path):
    document = normal_doc()
    document.sections[0].header.paragraphs[0].text = "页眉"
    stream = io.BytesIO()
    document.save(stream)
    damaged = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source, zipfile.ZipFile(damaged, "w") as target:
        for item in source.infolist():
            if not item.filename.startswith("word/header"):
                target.writestr(item, source.read(item.filename))

    result = run_pipeline(tmp_path, {"好的.docx": normal_doc(), "缺页眉.docx": damaged.getvalue()})

    assert records(result)["缺页眉.docx"].outcome is Outcome.UNPROCESSED
    assert records(result)["好的.docx"].outcome is Outcome.DONE


def test_cr05_a_strange_zip_is_just_not_a_word_file(tmp_path, monkeypatch):
    import homework_formatter.scanner as scanner

    def explode(*args, **kwargs):
        raise NotImplementedError("unsupported zip version")

    (tmp_path / "怪.docx").write_bytes(b"PK\x03\x04 something")
    monkeypatch.setattr(scanner.zipfile, "ZipFile", explode)

    (entry,) = scan_folder(tmp_path)

    assert entry.kind is FileKind.OTHER


def test_cr05_a_file_that_cannot_be_read_is_listed_and_the_rest_carry_on(tmp_path, monkeypatch):
    from pathlib import Path

    source = tmp_path / "学生作业"
    source.mkdir()
    make_docx(source / "好的.docx", [TITLE, "张三 20240301", HEADING, BODY])
    make_docx(source / "被占着.docx", [TITLE, "李四 20240302", HEADING, BODY])
    real = Path.read_bytes

    def locked(self):
        if self.name == "被占着.docx":
            raise PermissionError("locked")
        return real(self)

    monkeypatch.setattr(Path, "read_bytes", locked)

    result = pipeline.run(source, tmp_path / "排版结果", RULES, roster=ROSTER)

    assert records(result)["被占着.docx"].outcome is Outcome.UNPROCESSED
    assert any("读不出来" in n for n in records(result)["被占着.docx"].notes)
    assert records(result)["好的.docx"].outcome is Outcome.DONE


def test_cr05_the_inspector_crashing_still_leaves_a_report(tmp_path, monkeypatch):
    def explode(*args, **kwargs):
        raise RuntimeError("checker bug")

    monkeypatch.setattr(pipeline, "check_document", explode)

    result = run_pipeline(tmp_path, {"a.docx": normal_doc()})

    record = records(result)["a.docx"]
    assert record.outcome is Outcome.REVIEW
    assert any("质检时出错" in n for n in record.notes)
    assert result.report_path.is_file()
    assert folder(result, "已完成") == []


def test_cr05_an_unexpected_crash_still_writes_a_report_saying_the_run_is_incomplete(tmp_path, monkeypatch):
    def explode(*args, **kwargs):
        raise RuntimeError("disk on fire")

    monkeypatch.setattr(pipeline, "_write_outputs", explode)

    result = run_pipeline(tmp_path, {"a.docx": normal_doc()})

    assert "disk on fire" in result.incomplete
    overview = {r[0]: r[1] for r in openpyxl.load_workbook(str(result.report_path))["总览"].iter_rows(values_only=True)}
    assert "没有做完" in next(iter(overview))


def test_cr05_a_roster_that_is_not_really_an_excel_file_is_explained(tmp_path, capsys, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "学生作业").mkdir()
    (tmp_path / "名单.xlsx").write_bytes(b"this is not an xlsx file")

    code = main(["--rules", str(RULES_FILE)])

    assert code == 2
    assert "打不开" in capsys.readouterr().out


# ====================== CR-06 Word 存盘留下的分页记号挡住了去空格和换编号 ======================


def start_with_marker(paragraph, marker_xml):
    paragraph.runs[0]._r.insert(0, parse_xml(marker_xml))


@pytest.mark.parametrize(
    "marker",
    [f'<w:lastRenderedPageBreak {nsdecls("w")}/>', f'<w:br {nsdecls("w")} w:type="page"/>'],
)
def test_cr06_page_markers_at_the_start_of_a_paragraph_do_not_get_in_the_way(marker):
    document = doc(TITLE, "张三 20240301", "1. 我发现的问题", BODY)
    heading = document.add_paragraph("2. 原因分析")
    start_with_marker(heading, marker)
    body = document.add_paragraph("　　" + BODY_2)
    start_with_marker(body, marker)

    result = fmt(document)

    texts = paragraph_texts(result)
    assert "二、原因分析" in texts and BODY_2 in texts
    assert not any("2." in t for t in texts)
    assert verify(document, result) == []
    assert check_document(result.docx_bytes, RULES, INFO) == []


def test_cr06_a_number_that_cannot_be_reached_is_left_alone_with_a_warning():
    document = doc(TITLE, "张三 20240301")
    heading = document.add_paragraph()
    heading.add_run("1")
    heading.runs[0]._r.append(parse_xml(f'<w:sym {nsdecls("w")} w:font="Symbol" w:char="F0B7"/>'))
    heading.add_run(". 我发现的问题")
    document.add_paragraph(BODY)

    result = fmt(document)

    assert not any(t.startswith("一、") for t in paragraph_texts(result))
    assert verify(document, result) == []


# ====================== CR-07 小条目被当成小标题一起重编；短行被错当小标题 ======================


def roles_of(*paragraphs):
    return [b.role for b in classify(doc(*paragraphs))]


def test_cr07_list_items_in_a_row_are_not_headings():
    roles = roles_of(TITLE, "张三 20240301", "一、常用的工具", BODY, "1. 手机", "2. 电脑", "3. 平板", BODY_2)

    assert roles == [Role.TITLE, Role.INFO, Role.HEADING, Role.BODY, Role.BODY, Role.BODY, Role.BODY, Role.BODY]


def test_cr07_a_main_heading_directly_followed_by_a_sub_heading_keeps_both():
    roles = roles_of(TITLE, "张三 20240301", "一、研究背景", "（一）问题的提出", BODY)

    assert roles[2:4] == [Role.HEADING, Role.HEADING]


@pytest.mark.parametrize("last_line", ["谢谢老师", "2024年10月8日", "高二三班第四小组"])
def test_cr07_a_short_last_line_is_not_a_heading_and_gets_no_number(last_line):
    document = doc(TITLE, "张三 20240301", "我发现的问题", BODY, "我的办法", BODY_2, last_line)

    result = fmt(document)

    assert paragraph_texts(result)[-1] == last_line
    assert result.headings == ["一、我发现的问题", "二、我的办法"]


@pytest.mark.parametrize("line", ["图1 各年级睡眠时间", "表 2 调查结果", "v = s / t", "2024年10月8日"])
def test_cr07_captions_dates_and_equations_are_not_headings(line):
    roles = roles_of(TITLE, "张三 20240301", "我发现的问题", BODY, line, BODY_2)

    assert roles[4] is Role.BODY


def test_cr07_nested_items_numbered_like_the_headings_stop_the_renumbering():
    headings = ["1. 背景", "2. 调查方法", "1. 问卷调查", "2. 个别访谈", "3. 结论"]

    assert plan_numbering(headings, RULES.unnumbered_headings).headings is None


# ====================== CR-08 小标题开头两个数字连写 ======================


def test_cr08_headings_that_start_with_a_number_as_part_of_their_words_are_not_renumbered():
    headings = ["一、二年级的情况", "三、四年级的情况", "五、六年级的情况"]
    document = doc(TITLE, "张三 20240301")
    for heading in headings:
        document.add_paragraph(heading)
        document.add_paragraph(BODY)

    result = fmt(document)

    assert result.headings == headings
    assert any("编号" in w for w in result.warnings)


# ====================== CR-09 图表比版心宽时崩溃 ======================

CHART = (
    '<w:drawing {w} xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
    '<wp:inline><wp:extent cx="5274310" cy="3076575"/><wp:docPr id="9" name="图表 1"/>'
    '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
    "<c:chart/></a:graphicData></a:graphic></wp:inline></w:drawing>"
)


def test_cr09_an_embedded_chart_slightly_wider_than_the_page_is_shrunk_not_fatal():
    document = normal_doc()
    document.add_paragraph().add_run()._r.append(parse_xml(CHART.format(w=nsdecls("w"))))

    result = fmt(document)

    extent = next(open_bytes(result.docx_bytes).element.body.iter(
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"))
    assert int(extent.get("cx")) <= RULES.text_width_twips * 635
    assert abs(int(extent.get("cx")) / int(extent.get("cy")) - 5274310 / 3076575) < 0.001
    assert any("缩小" in c for c in result.changes)


# ====================== CR-10 软回车拆段后的处理 ======================


def test_cr10_a_picture_between_soft_breaks_survives(tiny_png):
    document = normal_doc()
    paragraph = document.add_paragraph()
    paragraph.add_run("图片上面的话。").add_break()
    paragraph.add_run().add_picture(io.BytesIO(tiny_png), width=Cm(3))
    paragraph.add_run().add_break()
    paragraph.add_run("图片下面的话。")

    result = fmt(document)

    assert len(open_bytes(result.docx_bytes).inline_shapes) == 1
    assert verify(document, result) == []


def test_cr10_a_heading_and_its_text_joined_by_a_soft_break_still_pass_the_comparison():
    document = doc(TITLE, "张三 20240301")
    run = document.add_paragraph().add_run("1. 我发现的问题")
    run.add_break()
    run.add_text(BODY)
    document.add_paragraph(BODY_2)

    result = fmt(document)

    assert paragraph_texts(result)[2:4] == ["一、我发现的问题", BODY]
    assert verify(document, result) == []


def test_cr10_splitting_does_not_multiply_section_breaks():
    document = normal_doc()
    document.add_section()
    document.add_paragraph("第二节里的话。")
    # 带分节信息的那一段里正好有软回车：拆开后分节信息只能留在最后一片上
    carrier = next(p for p in document.paragraphs if p._p.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None)
    run = carrier.add_run("第一行的话，有标点。")
    run.add_break()
    run.add_text("第二行的话，有标点。")
    assert len(document.sections) == 2

    result = fmt(document)

    assert len(open_bytes(result.docx_bytes).sections) == 2
    assert verify(document, result) == []


# ====================== CR-11 删除线、隐藏文字、符号字体被清掉 ======================


def test_cr11_struck_out_and_hidden_text_stay_that_way():
    document = normal_doc()
    paragraph = document.add_paragraph("保留的话，")
    paragraph.add_run("划掉的话").font.strike = True
    paragraph.add_run("藏起来的话").font.hidden = True

    result = fmt(document)

    last = open_bytes(result.docx_bytes).paragraphs[-1]
    assert [bool(r.font.strike) for r in last.runs] == [False, True, False]
    assert [bool(r.font.hidden) for r in last.runs] == [False, False, True]
    assert check_document(result.docx_bytes, RULES, INFO) == []  # 质检也认可


def test_cr11_characters_typed_in_a_symbol_font_keep_their_font():
    document = normal_doc()
    paragraph = document.add_paragraph("我同意这个办法 ")
    tick = paragraph.add_run("")
    tick.font.name = "Wingdings"

    result = fmt(document)

    last = find(open_bytes(result.docx_bytes), "我同意这个办法 ")
    assert run_format(text_runs(last)[-1])["ascii"] == "Wingdings"
    assert run_format(text_runs(last)[0])["east_asian"] == "宋体"
    assert check_document(result.docx_bytes, RULES, INFO) == []


# ====================== CR-12 页脚里的字清掉了不留记录 ======================


def test_cr12_words_in_the_footer_are_recorded_before_being_replaced_by_the_page_number():
    document = normal_doc()
    document.sections[0].footer.paragraphs[0].text = "指导老师：王建国　本文数据来自校学生会"

    result = fmt(document)

    assert any("页脚" in c and "指导老师：王建国" in c for c in result.changes)


def test_cr12_a_students_own_page_number_in_the_footer_is_not_worth_mentioning():
    document = normal_doc()
    document.sections[0].footer.paragraphs[0].text = "第 1 页"

    assert not any("页脚" in c for c in fmt(document).changes)


def test_cr12_a_picture_in_the_header_is_flagged(tiny_png):
    document = normal_doc()
    document.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(tiny_png), width=Cm(1))

    assert any("页眉" in w and "图片" in w for w in fmt(document).warnings)


# ====================== CR-13 报告里学生写的字被当成公式 ======================


def test_cr13_a_title_starting_with_an_equals_sign_is_text_not_a_formula(tmp_path):
    result = run_pipeline(
        tmp_path,
        {"=A1+B1.docx": ['=HYPERLINK("http://example.com","点这里看成绩")', "张三 20240301", HEADING, BODY]},
    )

    sheet = openpyxl.load_workbook(str(result.report_path))["明细"]
    cells = [c for row in sheet.iter_rows(min_row=2) for c in row if isinstance(c.value, str)]
    assert any(c.value.startswith("=HYPERLINK") for c in cells)
    assert any(c.value == "=A1+B1.docx" for c in cells)
    assert all(c.data_type == "s" for c in cells)


# ====================== CR-14 写文件阶段没有保护 ======================


def test_cr14_names_that_are_not_valid_filenames_are_made_safe(tmp_path):
    roster = Roster([("20240301", "张三*"), ("20240302", "李/四")])
    from homework_formatter.decisions import Decisions

    decisions = Decisions(assigned={"a.docx": ("张三*", "20240301"), "b.docx": ("李/四", "20240302")})

    result = run_pipeline(
        tmp_path,
        {"a.docx": [TITLE, HEADING, BODY], "b.docx": [TITLE, HEADING, BODY_2]},
        roster=roster,
        decisions=decisions,
    )

    # 名字里带了不能做文件名的字符：文件照样写得出来（换成下划线），并请老师确认
    assert folder(result, "需人工确认") == ["20240301-张三_.docx", "20240302-李_四.docx"]
    assert result.incomplete is None and result.report_path.is_file()


def test_cr14_very_long_original_names_do_not_break_the_review_folder(tmp_path):
    long_name = "很长的文件名" * 30
    result = run_pipeline(
        tmp_path,
        {f"{long_name}一.docx": [TITLE, "张三 20240301", HEADING, BODY], "二.docx": [TITLE, "张三 20240301", HEADING, BODY_2]},
    )

    names = folder(result, "需人工确认")
    assert len(names) == 2 and all(len(n) < 120 for n in names)
    assert result.report_path.is_file() and result.incomplete is None


def test_cr14_one_file_failing_to_write_does_not_stop_the_rest(tmp_path, monkeypatch):
    real = pipeline._write_new

    def flaky(folder_path, name, data):
        if "李四" in name:
            raise OSError("disk full")
        return real(folder_path, name, data)

    monkeypatch.setattr(pipeline, "_write_new", flaky)

    result = run_pipeline(
        tmp_path,
        {"a.docx": [TITLE, "张三 20240301", HEADING, BODY], "b.docx": [TITLE, "李四 20240302", HEADING, BODY]},
    )

    assert folder(result, "已完成") == ["20240301-张三.docx"]
    assert records(result)["b.docx"].outcome is Outcome.UNPROCESSED
    assert any("写文件时出错" in n for n in records(result)["b.docx"].notes)


def test_cr14_an_existing_file_is_never_overwritten(tmp_path):
    first = pipeline._write_new(tmp_path, "a.docx", b"one")
    second = pipeline._write_new(tmp_path, "a.docx", b"two")

    assert (first, second) == ("a.docx", "a_2.docx")
    assert (tmp_path / "a.docx").read_bytes() == b"one"


# ====================== CR-15 LibreOffice 卡住或清理失败 ======================


@pytest.fixture
def doc_files(tmp_path):
    paths = [tmp_path / "一.doc", tmp_path / "二.doc", tmp_path / "三.doc"]
    for path in paths:
        path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + path.name.encode())
    return paths


def fake_soffice(monkeypatch, behaviour):
    """用一个替身代替 LibreOffice：behaviour(要转的副本们, 输出文件夹) 决定写出什么、返回什么。"""
    from pathlib import Path

    monkeypatch.setattr(converter, "find_soffice", lambda: Path("soffice"))

    def run(command, timeout):
        out_dir = Path(command[command.index("--outdir") + 1])
        copies = [Path(c) for c in command[command.index("--outdir") + 2 :]]
        return behaviour(copies, out_dir)

    monkeypatch.setattr(converter, "_run_soffice", run)


def test_cr15_a_timeout_fails_every_file_with_a_reason_instead_of_crashing(doc_files, monkeypatch):
    fake_soffice(monkeypatch, lambda copies, out_dir: "LibreOffice 转换超过 180 秒没有完成")

    result = converter.convert_doc_files(doc_files)

    assert result.converted == {}
    assert all("超过" in reason for reason in result.failed.values())


def test_cr15_one_bad_file_in_a_batch_does_not_condemn_the_others(doc_files, monkeypatch):
    def behaviour(copies, out_dir):
        if len(copies) > 1:  # 一起转：卡在第二份，后面的都没出来
            (out_dir / f"{copies[0].stem}.docx").write_bytes(b"converted")
            return None
        if copies[0].read_bytes().endswith("二.doc".encode()):
            return None  # 单独转第二份：还是不行
        (out_dir / f"{copies[0].stem}.docx").write_bytes(b"converted")
        return None

    fake_soffice(monkeypatch, behaviour)

    result = converter.convert_doc_files(doc_files)

    assert sorted(p.name for p in result.converted) == ["一.doc", "三.doc"]
    assert [p.name for p in result.failed] == ["二.doc"]


def test_cr15_the_fixed_install_location_is_preferred_over_searching_the_path(monkeypatch, tmp_path):
    installed = tmp_path / "soffice.exe"
    installed.write_bytes(b"")
    monkeypatch.delenv(converter.SOFFICE_ENV, raising=False)
    monkeypatch.setattr(converter, "SOFFICE_LOCATIONS", (str(installed),))
    monkeypatch.setattr(converter.shutil, "which", lambda name: pytest.fail("不该先去别处找"))

    assert converter.find_soffice() == installed


# ====================== CR-16 决定表「不要这份」填什么都算不要 ======================


def decisions_file(tmp_path, rows):
    workbook = openpyxl.Workbook()
    workbook.active.append(["原文件名", "姓名", "学号", "不要这份", "备注"])
    for row in rows:
        workbook.active.append(row)
    path = tmp_path / "老师的决定.xlsx"
    workbook.save(str(path))
    return path


def test_cr16_writing_no_means_keep_it(tmp_path):
    path = decisions_file(tmp_path, [["a.docx", "张三", "20240301", "否", None], ["b.docx", None, None, "是", None]])

    decisions = load_decisions(path)

    assert decisions.discarded == {"b.docx": ""}
    assert decisions.assigned == {"a.docx": ("张三", "20240301")}


@pytest.mark.parametrize("answer", ["√", "不要", "是", 1])
def test_cr16_clear_ways_of_saying_yes(tmp_path, answer):
    assert "a.docx" in load_decisions(decisions_file(tmp_path, [["a.docx", None, None, answer, None]])).discarded


def test_cr16_an_unclear_answer_is_queried_not_guessed(tmp_path):
    with pytest.raises(DecisionsError, match="看不懂"):
        load_decisions(decisions_file(tmp_path, [["a.docx", None, None, "也许", None]]))


def test_cr16_the_same_file_on_two_rows_is_queried(tmp_path):
    rows = [["a.docx", "张三", "20240301", None, None], ["a.docx", "李四", "20240302", None, None]]

    with pytest.raises(DecisionsError, match="两次"):
        load_decisions(decisions_file(tmp_path, rows))


# ====================== CR-17 质检只看直接格式、只查列出来的几种 ======================


@pytest.fixture
def good():
    return fmt(normal_doc()).docx_bytes


def spoil(data, change):
    document = open_bytes(data)
    change(document)
    return save_bytes(document)


def run_properties(document, text):
    return docx.text.paragraph.Paragraph(find(document, text), None).runs[0]._r.get_or_add_rPr()


def add_property(properties, tag, **attributes):
    element = OxmlElement(tag)
    for key, value in attributes.items():
        element.set(qn(key.replace("_", ":")), value)
    properties.append(element)


@pytest.mark.parametrize(
    "tag, attributes, where",
    [
        ("w:shd", {"w_val": "clear", "w_fill": "FFFF00"}, "run"),
        ("w:spacing", {"w_val": "40"}, "run"),
        ("w:em", {"w_val": "dot"}, "run"),
        ("w:bdr", {"w_val": "single"}, "run"),
        ("w:shd", {"w_val": "clear", "w_fill": "D9D9D9"}, "paragraph"),
        ("w:keepNext", {}, "paragraph"),
    ],
)
def test_cr17_any_formatting_not_on_the_allowed_list_fails_inspection(good, tag, attributes, where):
    def change(document):
        if where == "run":
            add_property(run_properties(document, BODY), tag, **attributes)
        else:
            add_property(find(document, BODY).get_or_add_pPr(), tag, **attributes)

    problems = check_document(spoil(good, change), RULES, INFO)

    assert any("不该有" in p and tag.split(":")[1] in p for p in problems), problems


def test_cr17_spacing_counted_in_lines_and_theme_fonts_fail_inspection(good):
    def lines(document):
        find(document, BODY).find(qn("w:pPr") + "/" + qn("w:spacing")).set(qn("w:beforeLines"), "100")

    def theme(document):
        run_properties(document, BODY).find(qn("w:rFonts")).set(qn("w:eastAsiaTheme"), "minorEastAsia")

    assert any("多余的设置" in p for p in check_document(spoil(good, lines), RULES, INFO))
    assert any("主题" in p for p in check_document(spoil(good, theme), RULES, INFO))


def test_cr17_formatting_hidden_in_the_default_styles_fails_inspection(good):
    def dirty_normal_style(document):
        document.styles["Normal"].font.bold = True

    def dirty_defaults(document):
        defaults = document.styles.element.find(qn("w:docDefaults"))
        add_property(defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr")), "w:b")

    for change in (dirty_normal_style, dirty_defaults):
        assert any("默认格式没有清干净" in p for p in check_document(spoil(good, change), RULES, INFO))


def test_cr17_bold_and_shading_set_in_the_students_normal_style_do_not_survive_formatting():
    document = normal_doc()
    document.styles["Normal"].font.bold = True
    document.styles["Normal"].paragraph_format.space_after = Pt(20)

    result = fmt(document)

    output = open_bytes(result.docx_bytes)
    assert output.styles["Normal"].font.bold is None
    assert output.styles["Normal"].paragraph_format.space_after is None
    assert check_document(result.docx_bytes, RULES, INFO) == []


def test_cr17_a_table_style_with_shaded_bold_rows_is_flagged():
    document = normal_doc()
    table = document.add_table(rows=2, cols=2)
    table.style = "Medium Shading 1 Accent 1"
    table.cell(0, 0).text = "日期"

    assert any("表格样式" in w for w in fmt(document).warnings)


# ====================== CR-18 页码相关的漏洞 ======================


def test_cr18_a_students_roman_page_numbers_starting_at_five_are_reset():
    document = normal_doc()
    add_property(document.sections[0]._sectPr, "w:pgNumType", w_fmt="upperRoman", w_start="5")

    result = fmt(document)

    assert open_bytes(result.docx_bytes).sections[0]._sectPr.find(qn("w:pgNumType")) is None
    assert check_document(result.docx_bytes, RULES, INFO) == []


def different_first_page(document):
    document.sections[0].different_first_page_header_footer = True


def odd_and_even(document):
    document.settings.odd_and_even_pages_header_footer = True


def gutter(document):
    document.sections[0].gutter = Cm(3)


def roman(document):
    add_property(document.sections[0]._sectPr, "w:pgNumType", w_fmt="upperRoman")


def big_red_number(document):
    for run in document.sections[0].footer.paragraphs[0].runs:
        run.font.size = Pt(36)


def second_section_with_its_own_footer(document):
    document.add_section()
    document.sections[1].footer.is_linked_to_previous = False


@pytest.mark.parametrize(
    "change, expected",
    [
        (different_first_page, "首页不同"),
        (odd_and_even, "奇偶页不同"),
        (gutter, "装订线"),
        (roman, "阿拉伯数字"),
        (big_red_number, "页码的字号"),
        (second_section_with_its_own_footer, "另有自己的页脚"),
    ],
)
def test_cr18_page_number_faults_fail_inspection(good, change, expected):
    problems = check_document(spoil(good, change), RULES, INFO)

    assert any(expected in p for p in problems), problems


# ====================== CR-19 图片和文字在同一段 ======================


def test_cr19_a_picture_sharing_a_paragraph_with_its_caption_is_flagged(tiny_png):
    document = normal_doc()
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(io.BytesIO(tiny_png), width=Cm(3))
    paragraph.add_run("图 1 各年级的睡眠时间")

    result = fmt(document)

    assert any("图片和文字写在同一段" in w for w in result.warnings)


# ====================== CR-20 重复提交漏掉了读不了的那一份 ======================


def test_cr20_a_second_submission_that_cannot_be_read_still_counts_as_a_duplicate(tmp_path):
    result = run_pipeline(
        tmp_path,
        {"张三.docx": [TITLE, "张三 20240301", HEADING, BODY], "研究性学习 张三.pdf": b"%PDF-1.7"},
    )

    record = records(result)["张三.docx"]
    assert record.outcome is Outcome.REVIEW
    assert any("研究性学习 张三.pdf" in n and "读不了" in n for n in record.notes)
    overview = {r[0]: r[1] for r in openpyxl.load_workbook(str(result.report_path))["总览"].iter_rows(values_only=True)}
    assert "张三" not in overview["交了但没法处理的学生"]  # 他有一份排好了，等老师确认


# ====================== CR-21 ~ CR-25 排版里的几处小问题 ======================


def test_cr21_the_blank_line_between_two_tables_is_kept():
    document = normal_doc()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表一"
    document.add_paragraph("")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表二"

    body = open_bytes(fmt(document).docx_bytes).element.body
    kinds = [child.tag.split("}")[1] for child in body.iterchildren()]

    assert kinds[-4:-1] == ["tbl", "p", "tbl"]


def test_cr22_numbering_the_student_switched_off_does_not_come_back():
    document = normal_doc()
    paragraph = document.add_paragraph("这一段学生取消了编号，后面是普通的话。", style="List Number")
    paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0

    result = fmt(document)

    last = find(open_bytes(result.docx_bytes), "这一段学生取消了编号，后面是普通的话。")
    assert last.find(qn("w:pPr") + "/" + qn("w:numPr")) is None
    assert not any("自动编号" in w for w in result.warnings)


def test_cr23_bookmark_and_spelling_markers_between_paragraphs_are_not_unknown_content():
    document = normal_doc()
    body = document.element.body
    for xml in (f'<w:bookmarkEnd {nsdecls("w")} w:id="0"/>', f'<w:proofErr {nsdecls("w")} w:type="spellEnd"/>'):
        body.insert(2, parse_xml(xml))

    assert fmt(document).warnings == []


def test_cr24_trailing_space_after_a_heading_does_not_eat_its_words():
    document = doc(TITLE, "张三 20240301", "1. 背景  ", BODY)

    result = fmt(document)

    assert paragraph_texts(result)[2].strip() == "一、背景"
    assert verify(document, result) == []


def test_cr25_page_break_before_a_heading_is_kept():
    document = normal_doc()
    heading = document.add_paragraph("二、原因分析")
    heading.paragraph_format.page_break_before = True
    document.add_paragraph(BODY_2)

    result = fmt(document)

    paragraph = find(open_bytes(result.docx_bytes), "二、原因分析")
    assert paragraph.find(qn("w:pPr") + "/" + qn("w:pageBreakBefore")) is not None
    assert check_document(result.docx_bytes, RULES, INFO) == []


def test_cr25_a_document_with_no_page_settings_at_all_is_still_formatted():
    document = normal_doc()
    section = document.element.body.find(qn("w:sectPr"))
    document.element.body.remove(section)

    result = fmt(document)

    assert check_document(result.docx_bytes, RULES, INFO) == []


# ====================== CR-26 规则表对老师手改不够宽容 ======================


def rules_text():
    return RULES_FILE.read_text(encoding="utf-8")


def test_cr26_a_rules_file_saved_by_notepad_with_a_byte_order_mark_loads(tmp_path):
    path = tmp_path / "规则.toml"
    path.write_bytes(b"\xef\xbb\xbf" + rules_text().encode("utf-8"))

    assert load_rules(path).body.east_asian_font == "宋体"


def test_cr26_a_rules_file_saved_in_the_wrong_encoding_is_explained(tmp_path):
    path = tmp_path / "规则.toml"
    path.write_bytes(rules_text().encode("gbk"))

    with pytest.raises(RulesError, match="UTF-8"):
        load_rules(path)


@pytest.mark.parametrize(
    "old, new, complaint",
    [
        ('"纸张宽_厘米" = 21.0', '"纸张宽_厘米" = "21"', "数字"),
        ('"首行缩进_字" = 2', '"首行缩进_字" = 2.5', "整数"),
        ('"行距_倍" = 1.5', '"行距_倍" = "1.5倍"', "数字"),
        ('"字体" = "楷体"', '"字体" = 5', "文字"),
    ],
)
def test_cr26_values_of_the_wrong_kind_are_caught_when_loading(tmp_path, old, new, complaint):
    assert old in rules_text()
    path = tmp_path / "规则.toml"
    path.write_text(rules_text().replace(old, new), encoding="utf-8")

    with pytest.raises(RulesError, match=complaint):
        load_rules(path)


def test_cr26_an_indent_written_as_two_point_zero_becomes_a_whole_number(tmp_path):
    path = tmp_path / "规则.toml"
    path.write_text(rules_text().replace('"首行缩进_字" = 2', '"首行缩进_字" = 2.0'), encoding="utf-8")

    rules = load_rules(path)
    result = fmt(normal_doc(), rules=rules)

    assert rules.body.first_line_chars == 2 and isinstance(rules.body.first_line_chars, int)
    assert paragraph_format(find(open_bytes(result.docx_bytes), BODY))["first_line_chars"] == "200"


# ====================== CR-27 内容全在内容控件里的作业被说成"空白" ======================


def test_cr27_an_assignment_wrapped_in_a_content_control_is_not_called_blank(tmp_path):
    document = docx.Document()
    control = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent>'
        f"<w:p><w:r><w:t>{TITLE}</w:t></w:r></w:p><w:p><w:r><w:t>{BODY}</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
    )
    document.element.body.insert(0, control)

    result = run_pipeline(tmp_path, {"控件.docx": document})

    record = records(result)["控件.docx"]
    assert not any("空白" in n for n in record.notes)
    assert record.outcome is Outcome.REVIEW


# ====================== CR-28 杂项 ======================


def test_cr28_a_word_temp_file_disappearing_mid_run_is_not_a_change_to_the_originals(tmp_path):
    (tmp_path / "a.docx").write_bytes(b"x")
    (tmp_path / "~$a.docx").write_bytes(b"lock")
    before = fingerprint(tmp_path)

    (tmp_path / "~$a.docx").unlink()

    assert fingerprint(tmp_path) == before == fingerprint(tmp_path)
    assert list(before) == ["a.docx"]


def test_cr28_the_review_folder_is_only_mentioned_when_it_exists(tmp_path, capsys, monkeypatch):
    monkeypatch.chdir(tmp_path)
    source = tmp_path / "学生作业"
    source.mkdir()
    make_docx(source / "a.docx", [TITLE, "张三 20240301", HEADING, BODY])
    (source / "b.pdf").write_bytes(b"%PDF-1.7")

    main(["--rules", str(RULES_FILE)])

    printed = capsys.readouterr().out
    assert "需人工确认" not in printed
    assert "没法处理" in printed


# ====================== CR-29 安全网自己的测试太薄 ======================


def table_doc():
    document = normal_doc()
    table = document.add_table(rows=2, cols=2)
    for row, values in enumerate((("日期", "分钟"), ("周一", "9"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    return document


def test_cr29_a_changed_number_inside_a_table_is_caught():
    document = table_doc()
    result = fmt(document)

    def change(d):
        d.tables[0].cell(1, 1).paragraphs[0].runs[0].text = "19"

    problems = verify_content(save_bytes(document), spoil(result.docx_bytes, change), INFO, "张三", "20240301")

    assert any("9" in p and "19" in p for p in problems)


def test_cr29_swapped_table_cells_are_caught():
    document = table_doc()
    result = fmt(document)

    def change(d):
        d.tables[0].cell(0, 0).paragraphs[0].runs[0].text = "分钟"
        d.tables[0].cell(0, 1).paragraphs[0].runs[0].text = "日期"

    assert verify_content(save_bytes(document), spoil(result.docx_bytes, change), INFO, "张三", "20240301")


def test_cr29_when_the_info_line_was_left_alone_any_change_to_it_is_caught():
    document = doc(TITLE, "姓名：某某    学号：20249999", HEADING, BODY)
    result = fmt(document, name=None, student_id=None)
    assert verify_content(save_bytes(document), result.docx_bytes, None) == []

    def change(d):
        d.paragraphs[1].runs[0].text = "姓名：别人    学号：20249999"

    assert verify_content(save_bytes(document), spoil(result.docx_bytes, change), None)


def test_cr29_old_style_pictures_and_embedded_objects_are_counted():
    for tag in ("w:pict", "w:object"):
        document = normal_doc()
        document.add_paragraph().add_run()._r.append(OxmlElement(tag))
        result = fmt(document)

        def drop(d, tag=tag):
            element = next(d.element.body.iter(qn(tag)))
            element.getparent().remove(element)

        problems = verify_content(save_bytes(document), spoil(result.docx_bytes, drop), INFO, "张三", "20240301")

        assert any("图片" in p for p in problems), tag


def test_cr29_words_run_together_are_caught():
    document = doc(TITLE, "张三 20240301", HEADING, "We can see the lunch queue from the window.")
    result = fmt(document)

    def change(d):
        run = d.paragraphs[-1].runs[0]
        run.text = run.text.replace("see the lunch", "seethelunch")

    assert verify_content(save_bytes(document), spoil(result.docx_bytes, change), INFO, "张三", "20240301")


def test_cr29_originals_changing_during_a_run_is_noticed(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    source = tmp_path / "学生作业"
    source.mkdir()
    make_docx(source / "a.docx", [TITLE, "张三 20240301", HEADING, BODY])
    real = pipeline.format_document

    def meddling(docx_bytes, name, student_id, rules):
        (source / "a.docx").write_bytes(docx_bytes + b" ")  # 模拟：运行期间有人保存了这个文件
        return real(docx_bytes, name, student_id, rules)

    monkeypatch.setattr(pipeline, "format_document", meddling)

    code = main(["--rules", str(RULES_FILE)])

    assert code == 1
    assert "发生了变化" in capsys.readouterr().out
    (run_dir,) = (tmp_path / "排版结果").iterdir()
    overview = {r[0]: r[1] for r in openpyxl.load_workbook(str(run_dir / "处理报告.xlsx"))["总览"].iter_rows(values_only=True)}
    assert "不一致" in overview["原件核对"]


def test_cr29_real_formatting_helper_dataclass_is_used():
    # FormatResult 新增了字段后，测试里手工拼装结果容易漏字段；统一用 dataclasses.replace
    result = fmt(normal_doc())

    assert dataclasses.replace(result, docx_bytes=b"x").info_rewritten is True
