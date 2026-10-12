from pathlib import Path

import docx
import openpyxl
import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
REAL_INPUT = PROJECT_ROOT / "学生作业"
REAL_ROSTER = PROJECT_ROOT / "名单.xlsx"

needs_real_samples = pytest.mark.skipif(
    not REAL_INPUT.is_dir(), reason="没有「学生作业」文件夹"
)
needs_real_roster = pytest.mark.skipif(
    not REAL_ROSTER.is_file(), reason="没有「名单.xlsx」"
)


def make_docx(path: Path, paragraphs=(), header=None) -> Path:
    """造一份最简单的 Word 文件，供测试用。"""
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if header is not None:
        document.sections[0].header.paragraphs[0].text = header
    document.save(str(path))
    return path


def make_roster(path: Path, rows, header=("学号", "姓名")) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    workbook.save(str(path))
    return path


@pytest.fixture
def tiny_png() -> bytes:
    """一张 2×2 像素的图片，用来测试带图的文档。"""
    import struct
    import zlib

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 2 for _ in range(2))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
