"""排版：在文档副本上原地换格式。内容留在原处不搬动，所以图片、表格不会丢。

做法是"先清后套"：把每一段原有的格式整个拿掉，再按规则表明确写上该有的格式，
不依赖样式继承，这样不管学生原来怎么设置的，结果都一样。
"""

from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass, field

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips
from docx.text.paragraph import Paragraph

from .identity import MESSY, SEARCH_DEPTH, identity_leftover, info_line_kind, is_info_line
from .numbering import plan_numbering, split_number
from .rules import EMU_PER_TWIP, Rules, TextStyle
from .structure import (
    Block,
    Role,
    classify,
    has_page_break,
    is_soft_break,
    looks_like_title,
)

EDGE_WHITESPACE = " \t　 "
BLACK = "000000"
SINGLE_LINE = 240
# Word 的 A4 默认值：页眉距顶 1.5 厘米，页脚距底 1.75 厘米
HEADER_DISTANCE_TWIPS = 851
FOOTER_DISTANCE_TWIPS = 992
EXCERPT_LENGTH = 14
MAX_EMPHASIS_NOTES = 3
HIGHLIGHT_NAMES = {"yellow": "黄色高亮", "green": "绿色高亮", "red": "红色高亮", "cyan": "蓝色高亮"}

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
OBJECT_TAGS = (qn("w:drawing"), qn("w:pict"), qn("w:object"))
# 这些字符格式去掉会改变意思（划掉的字变回正常、隐藏的字显出来），所以留着
MEANINGFUL_RUN_PROPERTIES = ("w:strike", "w:dstrike", "w:vanish")
# 用这些字体打出来的是符号，换成宋体就成了方框或乱码
SYMBOL_FONTS = {"symbol", "wingdings", "wingdings 2", "wingdings 3", "webdings", "mt extra", "marlett"}
PRIVATE_USE = re.compile("[-]")
# 页脚里这样的字只是学生自己写的页码，清掉不算丢内容
PAGE_NUMBER_TEXT = re.compile(r"^[\s\d\-—–·/第页共of Pageage.]*$", re.IGNORECASE)
# 段首（段尾）不占位置的标记：Word 存盘时在每页开头留的记号、分页符。去空格、换编号时跳过它们
_SKIPPED_AT_EDGE = (qn("w:lastRenderedPageBreak"),)


@dataclass
class FormatResult:
    docx_bytes: bytes
    # 除了统一格式以外，还动了哪些地方（写给老师看）
    changes: list[str] = field(default_factory=list)
    # 程序拿不准、没有动的地方：需要老师看一眼
    warnings: list[str] = field(default_factory=list)
    # 排版后的题目和小标题
    title: str | None = None
    headings: list[str] = field(default_factory=list)
    # 姓名行有没有改写成统一写法（拿不准时不改）
    info_rewritten: bool = False


def format_document(
    docx_bytes: bytes, name: str | None, student_id: str | None, rules: Rules
) -> FormatResult:
    """name 和 student_id 传 None 表示拿不准这是谁的作业：格式照排，姓名行和页眉不动。"""
    document = docx.Document(io.BytesIO(docx_bytes))
    changes: list[str] = []
    warnings: list[str] = []
    owner_known = bool(name and student_id)
    if not document.sections:
        document.element.body.get_or_add_sectPr()

    _neutralize_defaults(document, rules)
    warnings.extend(_unusual_content(document))
    _split_soft_breaks(document, changes)
    blocks = classify(document)
    _note_emphasis(blocks, changes)
    _strip_edges(blocks, changes)
    blocks = _drop_empty_paragraphs(blocks, changes)

    header_info = [t for t in _part_texts(document, "header") if info_line_kind(t)]
    rewrite_info, keep_headers = False, False
    if owner_known:
        rewrite_info, keep_headers = _check_info_lines(blocks, header_info, name, student_id, warnings)
        if rewrite_info:
            blocks = _unify_info_line(
                document, blocks, rules.info_line(name, student_id), header_info, changes
            )
    elif header_info:
        keep_headers = True
        warnings.append("页眉里有姓名学号，但拿不准这份作业是谁的，页眉没有动")

    _keep_style_numbering(document, blocks)
    if any(b.role is Role.BODY and _direct_numbering(b.element) is not None for b in blocks):
        warnings.append("正文里有 Word 自动编号或项目符号的段落，编号和缩进保持原样")
    if any(b.role is Role.BODY and next(b.element.iter(*OBJECT_TAGS), None) is not None for b in blocks):
        warnings.append("有图片和文字写在同一段里，图片没有单独居中")
    _unify_heading_numbers(blocks, rules, changes, warnings)

    styles = {
        Role.TITLE: rules.title,
        Role.INFO: rules.info,
        Role.HEADING: rules.heading,
        Role.BODY: rules.body,
    }
    for block in blocks:
        if block.role in styles:
            _apply_style(block.element, styles[block.role], rules.line_spacing_value)
        elif block.role is Role.IMAGE:
            # 图片所在的行用单倍行距：倍数行距会按图片高度成倍留白
            _reset_paragraph(block.element, "center", 0, rules.body.size_pt, SINGLE_LINE)
    if any(block.role is Role.OTHER for block in blocks):
        warnings.append(
            "有程序没有排版的内容（比如独占一行的公式、插入的符号、内容控件），原样保留"
        )

    _format_tables(document, rules, changes, warnings)
    _fit_pictures(document, rules, changes)
    _set_up_pages(document, rules, header_info, changes, warnings, keep_headers)

    buffer = io.BytesIO()
    document.save(buffer)
    return FormatResult(
        buffer.getvalue(),
        changes,
        warnings,
        title=next((_text(b.element) for b in blocks if b.role is Role.TITLE), None),
        headings=[_text(b.element) for b in blocks if b.role is Role.HEADING],
        info_rewritten=rewrite_info,
    )


def _text(paragraph) -> str:
    return Paragraph(paragraph, None).text


# 这批作业里没出现、但别的作业里可能有的东西。认出来、原样保留、提醒老师，不硬改。
_UNUSUAL = [
    ((qn("w:txbxContent"),), "有文本框，文本框的位置和里面的文字请看一下"),
    ((qn("w:commentReference"), qn("w:commentRangeStart")), "有批注，原样保留"),
    (
        (qn("w:ins"), qn("w:del"), qn("w:moveFrom"), qn("w:moveTo")),
        "有修订痕迹（修订模式下的增删），原样保留；请在 Word 里接受或拒绝修订后再定稿",
    ),
    (
        (qn("w:footnoteReference"), qn("w:endnoteReference")),
        "有脚注或尾注，脚注里的文字没有排版",
    ),
    ((_WP + "anchor",), "有浮动的图片（不是嵌在文字行里的），位置没有调整"),
]


def _unusual_content(document) -> list[str]:
    body = document.element.body
    found = [message for tags, message in _UNUSUAL if next(body.iter(*tags), None) is not None]
    for columns in body.iter(qn("w:cols")):
        if int(columns.get(qn("w:num"), "1") or 1) > 1:
            found.append("页面有分栏，分栏没有改动")
            break
    return found


# ---------- 文字上的改动（都记进 changes） ----------


def _is_on(rpr, tag: str) -> bool:
    element = rpr.find(qn(tag)) if rpr is not None else None
    return element is not None and element.get(qn("w:val")) not in ("0", "false", "off", "none")


def _run_text(run) -> str:
    return "".join(t.text or "" for t in run.iter(qn("w:t")))


def _note_emphasis(blocks: list[Block], changes: list[str]) -> None:
    """正文里学生自己加的加粗、高亮等会被统一去掉；去掉之前记下来告诉老师。"""
    notes = []
    for block in blocks:
        if block.role is not Role.BODY:
            continue
        kinds: list[str] = []
        emphasised = ""
        for run in block.element.iter(qn("w:r")):
            rpr, text = run.find(qn("w:rPr")), _run_text(run)
            if rpr is None or not text.strip():
                continue
            found = [label for tag, label in (("w:b", "加粗"), ("w:i", "斜体"), ("w:u", "下划线")) if _is_on(rpr, tag)]
            highlight = rpr.find(qn("w:highlight"))
            if highlight is not None and highlight.get(qn("w:val")) != "none":
                found.append(HIGHLIGHT_NAMES.get(highlight.get(qn("w:val")), "高亮"))
            if found:
                emphasised += text
                kinds += [k for k in found if k not in kinds]
        if kinds:
            text = emphasised.strip(EDGE_WHITESPACE)
            excerpt = text[:EXCERPT_LENGTH] + ("…" if len(text) > EXCERPT_LENGTH else "")
            notes.append(f"去掉了正文里学生自己加的{'、'.join(kinds)}：「{excerpt}」")
    if len(notes) > MAX_EMPHASIS_NOTES:
        notes = [f"正文里有 {len(notes)} 段带学生自己加的加粗、下划线或高亮，已去掉"]
    changes.extend(notes)


def _split_soft_breaks(document, changes: list[str]) -> None:
    """段落里的软回车（Shift+回车）其实是分段，拆成真正的段落。在认结构之前做，
    这样拆出来的每一行和学生直接按回车写出来的一行，后面的处理完全一样。"""
    paragraphs = [e for e in document.element.body.iterchildren() if e.tag == qn("w:p")]
    first_with_text = next((p for p in paragraphs if _text(p).strip()), None)
    converted = 0
    for paragraph in paragraphs:
        breaks = [br for br in paragraph.iter(qn("w:br")) if is_soft_break(br)]
        # 只处理最常见的情况：软回车在段落直属的文字块里。藏在超链接等结构里的不动。
        if not breaks or any(br.getparent().getparent() is not paragraph for br in breaks):
            continue
        if paragraph is first_with_text:
            # 题目自己分两行写的不拆；但题目和姓名行挤在一段里的要拆开
            lines = _text(paragraph).split("\n")
            if looks_like_title("".join(lines)) and not any(is_info_line(line) for line in lines[1:]):
                continue
        _split_paragraph(paragraph)
        converted += len(breaks)
    if converted:
        changes.append(f"把 {converted} 处软回车（Shift+回车）改成了真正的分段")


def _split_paragraph(paragraph) -> list:
    ppr = paragraph.find(qn("w:pPr"))
    segments: list[list] = [[]]
    for child in list(paragraph):
        if child is ppr:
            continue
        if child.tag != qn("w:r") or not any(is_soft_break(c) for c in child):
            segments[-1].append(child)
            continue
        rpr = child.find(qn("w:rPr"))

        def new_run():
            run = OxmlElement("w:r")
            if rpr is not None:
                run.append(copy.deepcopy(rpr))
            return run

        current = new_run()
        for node in list(child):
            if node is rpr:
                continue
            if is_soft_break(node):
                segments[-1].append(current)
                segments.append([])
                current = new_run()
            else:
                current.append(node)
        segments[-1].append(current)

    pieces = []
    anchor = paragraph
    for index, nodes in enumerate(segments):
        piece = OxmlElement("w:p")
        if ppr is not None:
            properties = copy.deepcopy(ppr)
            # 分节信息只属于最后一片，自动编号和"段前分页"只属于第一片，不能每片都抄一份
            once_only = [("w:sectPr", index == len(segments) - 1)]
            once_only += [(tag, index == 0) for tag in ("w:numPr", "w:pageBreakBefore")]
            for tag, keep in once_only:
                element = properties.find(qn(tag))
                if element is not None and not keep:
                    properties.remove(element)
            piece.append(properties)
        for node in nodes:
            piece.append(node)
        anchor.addnext(piece)
        anchor = piece
        pieces.append(piece)
    paragraph.getparent().remove(paragraph)
    return pieces


def _skipped_at_edge(node) -> bool:
    """段首段尾不占位置的标记：Word 存盘时留的"上次分页位置"记号、分页符。"""
    return node.tag in _SKIPPED_AT_EDGE or (node.tag == qn("w:br") and not is_soft_break(node))


def _strip_edge(paragraph, leading: bool) -> bool:
    """删掉段首（或段尾）手敲的空格、全角空格和 Tab。返回有没有删掉东西。"""
    removed = False
    runs = list(paragraph.iter(qn("w:r")))
    for run in runs if leading else reversed(runs):
        children = [c for c in run if c.tag != qn("w:rPr")]
        for node in children if leading else reversed(children):
            if _skipped_at_edge(node):
                continue
            if node.tag == qn("w:tab"):
                run.remove(node)
                removed = True
            elif node.tag == qn("w:t"):
                text = node.text or ""
                kept = text.lstrip(EDGE_WHITESPACE) if leading else text.rstrip(EDGE_WHITESPACE)
                removed = removed or kept != text
                if kept:
                    node.text = kept
                    return removed
                run.remove(node)
            else:
                return removed
    return removed


def _strip_edges(blocks: list[Block], changes: list[str]) -> None:
    indented = 0
    for block in blocks:
        if block.role in (Role.TITLE, Role.HEADING, Role.BODY):
            removed = _strip_edge(block.element, leading=True)
            _strip_edge(block.element, leading=False)
            indented += removed and block.role is Role.BODY
            block.text = _text(block.element)
    if indented:
        changes.append(f"把 {indented} 段开头手敲的空格或 Tab 换成了真正的首行缩进")


def _drop_empty_paragraphs(blocks: list[Block], changes: list[str]) -> list[Block]:
    kept, dropped = [], 0
    for index, block in enumerate(blocks):
        element = block.element
        if block.role is not Role.EMPTY:
            kept.append(block)
            continue
        # 带分页符或分节信息的空段落有排版作用，留着
        structural = has_page_break(element) or element.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None
        # 两个表格之间的空段落也留着：删了它，两个表格在 Word 里会连成一个
        neighbours = blocks[index - 1 : index] + blocks[index + 1 : index + 2]
        between_tables = len(neighbours) == 2 and all(b.role is Role.TABLE for b in neighbours)
        if structural or between_tables:
            kept.append(block)
        else:
            element.getparent().remove(element)
            dropped += 1
    if dropped:
        changes.append(f"删掉了 {dropped} 个空段落")
    return kept


def _set_text(paragraph, text: str) -> None:
    for child in list(paragraph):
        if child.tag != qn("w:pPr"):
            paragraph.remove(child)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    node.set(XML_SPACE, "preserve")
    run.append(node)
    paragraph.append(run)


def _check_info_lines(
    blocks: list[Block], header_info: list[str], name: str, student_id: str, warnings: list[str]
) -> tuple[bool, bool]:
    """改写姓名行之前先看清楚：只有"除了班级、姓名、学号什么都没有"的行才能整行换掉。

    返回（可不可以改写姓名行, 要不要保留页眉）。
    """
    rewrite, keep_headers = True, False
    opening = [b for b in blocks if b.role in (Role.TITLE, Role.INFO, Role.HEADING, Role.BODY)]
    for block in opening[:SEARCH_DEPTH]:
        text = block.text.strip()
        mentions_owner = name in text or student_id in text
        if block.role is Role.INFO:
            if identity_leftover(text, name, student_id):
                # 是一行干净的姓名行，但写的不全是这位学生：不能当成他的姓名行换掉
                block.role = Role.BODY
                rewrite = rewrite and not mentions_owner
                warnings.append(f"开头有一行写着别的姓名或学号（「{text}」），原样保留，请确认")
        elif info_line_kind(text) == MESSY and mentions_owner:
            rewrite = False
            warnings.append(
                f"开头有一行除了姓名学号还写了别的内容（「{text}」），没有改写成统一的姓名行，请手动整理"
            )
    for text in header_info:
        if identity_leftover(text, name, student_id):
            keep_headers = True
            warnings.append(f"页眉里除了姓名学号还有别的内容（「{text}」），页眉没有动")
    return rewrite, keep_headers


def _unify_info_line(document, blocks, text: str, header_info: list[str], changes) -> list[Block]:
    """姓名行统一写法，放在题目下面一行；学生没写就补上。"""
    existing = [b for b in blocks if b.role is Role.INFO]
    title = next((b for b in blocks if b.role is Role.TITLE), None)
    body = document.element.body

    if existing:
        info = existing[0]
        old = "　".join(b.text.strip() for b in existing)
        for extra in existing[1:]:
            extra.element.getparent().remove(extra.element)
            blocks.remove(extra)
        if old != text:
            changes.append(f"姓名行由「{old}」改成了统一写法")
        blocks.remove(info)
    else:
        info = Block(Role.INFO, OxmlElement("w:p"))
        where = "原来写在页眉里" if header_info else "原文里没有"
        changes.append(f"在题目下面补上了姓名行（{where}）")
    _set_text(info.element, text)
    info.text = text

    if title is not None:
        title.element.addnext(info.element)
        blocks.insert(blocks.index(title) + 1, info)
    else:
        body.insert(0, info.element)
        blocks.insert(0, info)
    return blocks


# ---------- 小标题编号（老师追加的要求） ----------


def _direct_numbering(paragraph):
    return paragraph.find(qn("w:pPr") + "/" + qn("w:numPr"))


def _numbering_definitions(document):
    # 不用 document.part.numbering_part：文档里没有时它会顺手新建一个
    for relationship in document.part.rels.values():
        if relationship.reltype == RT.NUMBERING:
            return relationship.target_part.element
    return None


def _shows_a_number(definitions, numbering) -> bool:
    """这个自动编号设置会不会真的显示出编号或项目符号。有些只是摆设（比如标题的大纲级别）。"""
    num_id = numbering.find(qn("w:numId"))
    level = numbering.find(qn("w:ilvl"))
    num_id = num_id.get(qn("w:val")) if num_id is not None else None
    level = level.get(qn("w:val")) if level is not None else "0"
    if num_id in (None, "0") or definitions is None:
        return False
    instance = definitions.xpath(f"w:num[@w:numId='{num_id}']/w:abstractNumId")
    if not instance:
        return False
    abstract_id = instance[0].get(qn("w:val"))
    levels = definitions.xpath(
        f"w:abstractNum[@w:abstractNumId='{abstract_id}']/w:lvl[@w:ilvl='{level}']"
    )
    if not levels:
        return True  # 查不到定义就当它会显示，宁可多提醒
    number_format, text = levels[0].find(qn("w:numFmt")), levels[0].find(qn("w:lvlText"))
    if number_format is not None and number_format.get(qn("w:val")) == "none":
        return False
    return text is None or bool((text.get(qn("w:val")) or "").strip())


def _keep_style_numbering(document, blocks: list[Block]) -> None:
    """让"段落上直接写着编号设置"等同于"这一段真的显示自动编号"。

    不显示编号的摆设去掉；跟着样式来的真编号直接写到段落上——样式马上要被拿掉，
    不这样做编号就消失了。
    """
    definitions = _numbering_definitions(document)
    styles = {s.get(qn("w:styleId")): s for s in document.styles.element.findall(qn("w:style"))}

    def numbering_of(style_id, depth=0):
        style = styles.get(style_id)
        if style is None or depth > 10:
            return None
        numbering = style.find(qn("w:pPr") + "/" + qn("w:numPr"))
        if numbering is not None:
            return numbering
        based_on = style.find(qn("w:basedOn"))
        return numbering_of(based_on.get(qn("w:val")), depth + 1) if based_on is not None else None

    for block in blocks:
        if block.element.tag != qn("w:p"):
            continue
        direct = _direct_numbering(block.element)
        if direct is not None:
            # 段落上自己写了编号设置（包括"这一段不要编号"），以它为准，不再看样式
            if not _shows_a_number(definitions, direct):
                direct.getparent().remove(direct)
            continue
        ppr = block.element.find(qn("w:pPr"))
        style = ppr.find(qn("w:pStyle")) if ppr is not None else None
        inherited = numbering_of(style.get(qn("w:val"))) if style is not None else None
        if inherited is not None and _shows_a_number(definitions, inherited):
            ppr.append(copy.deepcopy(inherited))


def _unify_heading_numbers(blocks: list[Block], rules: Rules, changes: list[str], warnings: list[str]) -> None:
    headings = [b for b in blocks if b.role is Role.HEADING]
    if not rules.unify_heading_numbers or not headings:
        return
    if any(_direct_numbering(b.element) is not None for b in headings):
        warnings.append("小标题用了 Word 的自动编号（编号不是打出来的字），编号没有改")
        return
    texts = [_text(b.element) for b in headings]
    plan = plan_numbering(texts, rules.unnumbered_headings)
    if plan.headings is None:
        warnings.append(plan.problem)
        return
    replaced_all = True
    for block, old, new in zip(headings, texts, plan.headings):
        if old.strip() == new:
            continue
        core = split_number(old)[1]
        # 要去掉的是：开头的空白 + 原来的编号。用去掉首尾空白后的长度算，免得把结尾的空白也算进去
        old_length = (len(old) - len(old.lstrip())) + (len(old.strip()) - len(core))
        if _replace_prefix(block.element, old_length, new[: len(new) - len(core)]):
            block.text = _text(block.element)
        else:
            replaced_all = False
    if not replaced_all:
        warnings.append("有小标题的编号夹在特殊内容里，没能换成「一、二、三」，请手动改")
    elif plan.note:
        changes.append(plan.note)


def _leading_text_nodes(paragraph):
    """段落开头连着的文字和 Tab（跳过不占位置的标记），遇到别的东西就停。"""
    for run in paragraph.iter(qn("w:r")):
        for node in [c for c in run if c.tag != qn("w:rPr")]:
            if _skipped_at_edge(node):
                continue
            if node.tag not in (qn("w:t"), qn("w:tab")):
                return
            yield run, node


def _replace_prefix(paragraph, old_length: int, new_prefix: str) -> bool:
    """去掉段落开头的 old_length 个字，换成 new_prefix。只动开头，后面的内容原样不碰。

    开头够不着那么多字（中间夹着别的东西）就什么都不改，返回 False。
    """
    nodes = list(_leading_text_nodes(paragraph))
    available = sum(1 if node.tag == qn("w:tab") else len(node.text or "") for _, node in nodes)
    if available < old_length:
        return False
    remaining = old_length
    for run, node in nodes:
        if remaining <= 0:
            break
        if node.tag == qn("w:tab"):
            run.remove(node)
            remaining -= 1
        else:
            text = node.text or ""
            node.text = text[remaining:]
            remaining -= min(remaining, len(text))
    first = next((node for _, node in _leading_text_nodes(paragraph) if node.tag == qn("w:t")), None)
    if first is None:
        run = OxmlElement("w:r")
        first = OxmlElement("w:t")
        first.text = ""
        run.append(first)
        ppr = paragraph.find(qn("w:pPr"))
        paragraph.insert(0 if ppr is None else 1, run)
    first.text = new_prefix + (first.text or "")
    return True


# ---------- 套格式 ----------


def _fonts_element(style: TextStyle):
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), style.western_font)
    fonts.set(qn("w:hAnsi"), style.western_font)
    fonts.set(qn("w:eastAsia"), style.east_asian_font)
    fonts.set(qn("w:cs"), style.western_font)
    # 引号、破折号这类中西文都有的符号，按中文字体排
    fonts.set(qn("w:hint"), "eastAsia")
    return fonts


def _run_properties(style: TextStyle, kept: dict | None = None):
    """一段文字该有的格式。kept 是从原格式里留下来的几样（符号字体、删除线、隐藏、上下标）。

    元素的先后顺序是 Word 文件格式规定的，不能乱。
    """
    kept = kept or {}
    rpr = OxmlElement("w:rPr")
    rpr.append(kept.get("w:rFonts", _fonts_element(style)))
    for tag in MEANINGFUL_RUN_PROPERTIES:
        if tag in kept:
            rpr.append(kept[tag])
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLACK)
    rpr.append(color)
    for tag in ("w:sz", "w:szCs"):
        size = OxmlElement(tag)
        size.set(qn("w:val"), str(round(style.size_pt * 2)))
        rpr.append(size)
    if "w:vertAlign" in kept:
        rpr.append(kept["w:vertAlign"])
    return rpr


def _reset_paragraph(paragraph, alignment: str, first_line_chars: int, size_pt: float, line: int, mark=None):
    """把段落格式整个换掉。只留下"段前分页"、自动编号和分节信息，那三样属于内容和版面结构。"""
    old = paragraph.find(qn("w:pPr"))
    page_break = numbering = section = None
    if old is not None:
        if _is_on(old, "w:pageBreakBefore"):
            page_break = old.find(qn("w:pageBreakBefore"))
        numbering, section = old.find(qn("w:numPr")), old.find(qn("w:sectPr"))
        paragraph.remove(old)

    ppr = OxmlElement("w:pPr")
    if page_break is not None:
        ppr.append(page_break)
    if numbering is not None:
        ppr.append(numbering)
    spacing = OxmlElement("w:spacing")
    for key, value in (("w:before", 0), ("w:after", 0), ("w:line", line)):
        spacing.set(qn(key), str(value))
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    if numbering is None:  # 自动编号的段落，缩进由编号决定
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "0")
        indent.set(qn("w:right"), "0")
        indent.set(qn("w:firstLineChars"), str(first_line_chars * 100))
        indent.set(qn("w:firstLine"), str(round(first_line_chars * size_pt * 20)))
        ppr.append(indent)
    justification = OxmlElement("w:jc")
    justification.set(qn("w:val"), alignment)
    ppr.append(justification)
    if mark is not None:
        ppr.append(mark)
    if section is not None:
        ppr.append(section)
    paragraph.insert(0, ppr)


def _is_symbol_run(run, old_properties) -> bool:
    fonts = old_properties.find(qn("w:rFonts")) if old_properties is not None else None
    names = [fonts.get(qn(a), "") for a in ("w:ascii", "w:hAnsi")] if fonts is not None else []
    return any(n.lower() in SYMBOL_FONTS for n in names) or bool(PRIVATE_USE.search(_run_text(run)))


def _apply_style(paragraph, style: TextStyle, line: int) -> None:
    _reset_paragraph(
        paragraph, style.alignment, style.first_line_chars, style.size_pt, line, _run_properties(style)
    )
    for run in paragraph.iter(qn("w:r")):
        old = run.find(qn("w:rPr"))
        kept = {}
        if old is not None:
            for tag in (*MEANINGFUL_RUN_PROPERTIES, "w:vertAlign"):
                element = old.find(qn(tag))
                if element is not None and (tag == "w:vertAlign" or _is_on(old, tag)):
                    kept[tag] = element
            if _is_symbol_run(run, old) and old.find(qn("w:rFonts")) is not None:
                kept["w:rFonts"] = old.find(qn("w:rFonts"))
            run.remove(old)
        run.insert(0, _run_properties(style, kept))


def _neutralize_defaults(document, rules: Rules) -> None:
    """文档的"默认格式"和"正文"样式也清干净，免得没被直接改到的地方露出旧格式。"""
    styles = document.styles.element
    old = styles.find(qn("w:docDefaults"))
    if old is not None:
        styles.remove(old)
    defaults = OxmlElement("w:docDefaults")
    run_default = OxmlElement("w:rPrDefault")
    rpr = _run_properties(rules.body)
    rpr.remove(rpr.find(qn("w:color")))
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "en-US")
    language.set(qn("w:eastAsia"), "zh-CN")
    language.set(qn("w:bidi"), "ar-SA")
    rpr.append(language)
    run_default.append(rpr)
    paragraph_default = OxmlElement("w:pPrDefault")
    paragraph_default.append(OxmlElement("w:pPr"))
    defaults.append(run_default)
    defaults.append(paragraph_default)
    styles.insert(0, defaults)

    for style in styles.findall(qn("w:style")):
        if style.get(qn("w:type")) == "paragraph" and style.get(qn("w:default")) in ("1", "true"):
            for tag in ("w:pPr", "w:rPr"):
                element = style.find(qn(tag))
                if element is not None:
                    style.remove(element)


# ---------- 表格和图片 ----------


def _format_tables(document, rules: Rules, changes: list[str], warnings: list[str]) -> None:
    shrunk = False
    fancy_styles = {
        s.get(qn("w:styleId"))
        for s in document.styles.element.findall(qn("w:style"))
        if s.find(qn("w:tblStylePr")) is not None
    }
    for table in document.tables:
        element = table._tbl
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        indent = element.tblPr.find(qn("w:tblInd"))
        if indent is not None:
            element.tblPr.remove(indent)
        shrunk = _fit_table(element, rules.text_width_twips) or shrunk
        for paragraph in element.iter(qn("w:p")):
            _apply_style(paragraph, rules.table_text, rules.line_spacing_value)
        style = element.tblPr.find(qn("w:tblStyle"))
        if style is not None and style.get(qn("w:val")) in fancy_styles:
            message = "表格用了带底色或加粗的表格样式，样式没有动，请看一下表格里的字是否清楚"
            if message not in warnings:
                warnings.append(message)
    if shrunk:
        changes.append("表格比新页面的版心宽，已按比例缩小")


def _fit_table(table, text_width: int) -> bool:
    grid = table.find(qn("w:tblGrid"))
    columns = grid.findall(qn("w:gridCol")) if grid is not None else []
    total = sum(int(c.get(qn("w:w"), 0)) for c in columns)
    if total <= text_width:
        return False
    factor = text_width / total
    widths = list(columns) + [
        w for w in table.iter(qn("w:tcW"), qn("w:tblW")) if w.get(qn("w:type")) == "dxa"
    ]
    for element in widths:
        element.set(qn("w:w"), str(int(int(element.get(qn("w:w"), 0)) * factor)))
    return True


def _fit_pictures(document, rules: Rules, changes: list[str]) -> None:
    """嵌在文字行里的图片、图表比版心宽就按比例缩小。直接改尺寸数字，图片、图表、形状都适用。"""
    text_width = rules.text_width_twips * EMU_PER_TWIP
    shrunk = False
    for inline in document.element.body.iter(_WP + "inline"):
        extent = inline.find(_WP + "extent")
        width = int(extent.get("cx", 0)) if extent is not None else 0
        if width <= text_width:
            continue
        factor = text_width / width
        sizes = [extent] + [e for e in inline.iter(_A + "ext") if e.getparent().tag == _A + "xfrm"]
        for size in sizes:
            for attribute in ("cx", "cy"):
                if size.get(attribute) is not None:
                    size.set(attribute, str(int(int(size.get(attribute)) * factor)))
        shrunk = True
    if shrunk:
        changes.append("图片比新页面的版心宽，已按比例缩小")


# ---------- 页面、页眉、页码 ----------


def _parts(document, kind: str):
    """文档里实际存在的页眉（或页脚）。"""
    for section in document.sections:
        if kind == "header":
            candidates = (section.header, section.first_page_header, section.even_page_header)
        else:
            candidates = (section.footer, section.first_page_footer, section.even_page_footer)
        for part in candidates:
            # 没有自己页眉的节，python-docx 一访问就会新建一个；先判断再读
            if not part.is_linked_to_previous:
                yield part


def _part_texts(document, kind: str) -> list[str]:
    """页眉（或页脚）里每一段的文字，包括放在表格、文本框里的。"""
    texts = []
    for part in _parts(document, kind):
        for paragraph in part._element.iter(qn("w:p")):
            text = "".join(t.text or "" for t in paragraph.iter(qn("w:t"))).strip()
            if text and text not in texts:
                texts.append(text)
    return texts


def _set_up_pages(
    document,
    rules: Rules,
    header_info: list[str],
    changes: list[str],
    warnings: list[str],
    keep_headers: bool = False,
) -> None:
    # 清掉之前先把里面的东西记下来：字写进改动，图片没法写成字，所以提醒老师
    removed_kinds = ["footer"] if keep_headers else ["header", "footer"]
    for kind, label in (("header", "页眉"), ("footer", "页脚")):
        if kind not in removed_kinds:
            continue
        texts = _part_texts(document, kind)
        if kind == "header":
            others = [t for t in texts if t not in header_info]
            if others:
                changes.append(f"清空了页眉（原内容：「{'　'.join(others)}」）")
            elif header_info:
                changes.append("清空了页眉（里面的姓名学号已写到题目下面）")
        else:
            others = [t for t in texts if not PAGE_NUMBER_TEXT.match(t)]
            if others:
                changes.append(f"页脚换成了页码（原内容：「{'　'.join(others)}」）")
        if any(next(part._element.iter(*OBJECT_TAGS), None) is not None for part in _parts(document, kind)):
            warnings.append(f"{label}里原来有图片，已随{label}一起清掉，请确认是否需要")

    was_landscape = False
    for section in document.sections:
        # 没写纸张大小的文档读出来是 None，当成 0
        was_landscape = was_landscape or (section.page_width or 0) > (section.page_height or 0)
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Twips(rules.page_width_twips)
        section.page_height = Twips(rules.page_height_twips)
        section.top_margin = section.bottom_margin = Twips(rules.margin_top_bottom_twips)
        section.left_margin = section.right_margin = Twips(rules.margin_left_right_twips)
        section.header_distance = Twips(HEADER_DISTANCE_TWIPS)
        section.footer_distance = Twips(FOOTER_DISTANCE_TWIPS)
        section.gutter = 0
        section.different_first_page_header_footer = False
        headers = (section.header, section.first_page_header, section.even_page_header)
        footers = (section.footer, section.first_page_footer, section.even_page_footer)
        for part in footers if keep_headers else headers + footers:
            part.is_linked_to_previous = True
        _reset_line_grid(section._sectPr)
        # 学生设过的页码格式（罗马数字、从第 5 页起）去掉，页码一律 1、2、3
        numbering = section._sectPr.find(qn("w:pgNumType"))
        if numbering is not None:
            section._sectPr.remove(numbering)
    document.settings.odd_and_even_pages_header_footer = False
    if was_landscape:
        changes.append("纸张由横向改成了竖向")

    footer = document.sections[0].footer
    footer.is_linked_to_previous = False
    for child in list(footer._element):
        footer._element.remove(child)
    footer._element.append(_page_number_paragraph(rules))


def _reset_line_grid(section_properties) -> None:
    # "文档网格"会让行距对齐到网格上，1.5 倍行距就不准了；去掉网格类型
    grid = section_properties.find(qn("w:docGrid"))
    if grid is not None:
        for attribute in ("w:type", "w:charSpace"):
            grid.attrib.pop(qn(attribute), None)


def _page_number_paragraph(rules: Rules):
    paragraph = OxmlElement("w:p")
    _reset_paragraph(paragraph, rules.page_number.alignment, 0, rules.page_number.size_pt, SINGLE_LINE)

    def add_run(child):
        run = OxmlElement("w:r")
        run.append(_run_properties(rules.page_number))
        run.append(child)
        paragraph.append(run)

    def field_char(kind: str):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        return element

    instruction = OxmlElement("w:instrText")
    instruction.set(XML_SPACE, "preserve")
    instruction.text = " PAGE "
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"

    add_run(field_char("begin"))
    add_run(instruction)
    add_run(field_char("separate"))
    add_run(placeholder)
    add_run(field_char("end"))
    return paragraph
