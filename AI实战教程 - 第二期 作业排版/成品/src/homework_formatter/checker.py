"""质检：重新打开排好的文件，对照规则表逐条核对。

和排版程序（formatter.py）互不调用：排版的人不能自己说"做好了"。
两边共用的只有"认结构"（哪一段是题目、小标题、正文）和规则表。
"""

from __future__ import annotations

import io
from pathlib import Path

import docx
from docx.oxml.ns import qn

from .numbering import chinese_numeral, split_number
from .rules import EMU_PER_TWIP, Rules, TextStyle
from .structure import Block, Role, classify, has_page_break, is_soft_break

EDGE_WHITESPACE = " \t　 "
BLACK = "000000"
SINGLE_LINE = 240
EXCERPT_LENGTH = 10
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
ALIGNMENT_NAMES = {"center": "居中", "left": "靠左", "right": "靠右", "both": "两端对齐", None: "没有设置"}


def check_document(
    docx_bytes: bytes,
    rules: Rules,
    info_line: str | None,
    expect_title: bool = True,
) -> list[str]:
    """返回不符合要求的地方；空列表表示通过。

    info_line 传 None 表示这份作业的姓名行没有统一（拿不准是谁的），不查第三条的写法和位置。
    expect_title 传 False 表示已经知道这份作业没有题目。
    """
    document = docx.Document(io.BytesIO(docx_bytes))
    blocks = classify(document)
    problems: list[str] = []
    problems += _check_page(document, rules)
    problems += _check_title(blocks, rules, expect_title)
    problems += _check_info(blocks, rules, info_line)
    problems += _check_headings(blocks, rules)
    problems += _check_body(blocks, rules)
    problems += _check_tables_and_pictures(document, blocks, rules)
    problems += _check_page_number(document, rules)
    problems += _check_defaults(document)
    return problems


def _check_defaults(document) -> list[str]:
    """上面各条只看"直接写在段落上的格式"。这里补查样式表：文档默认格式和「正文」样式
    如果没清干净，里面的加粗、底色、间距会透出来，而直接格式上看不到。"""
    styles = document.styles.element
    dirty = []
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is not None:
        paragraph = defaults.find(qn("w:pPrDefault") + "/" + qn("w:pPr"))
        if paragraph is not None and len(paragraph):
            dirty.append("文档默认的段落格式")
        run = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        allowed = {qn(t) for t in ("w:rFonts", "w:sz", "w:szCs", "w:lang")}
        if run is not None and any(child.tag not in allowed for child in run):
            dirty.append("文档默认的文字格式")
    for style in styles.findall(qn("w:style")):
        is_default = style.get(qn("w:type")) == "paragraph" and style.get(qn("w:default")) in ("1", "true")
        if is_default and (style.find(qn("w:pPr")) is not None or style.find(qn("w:rPr")) is not None):
            dirty.append("「正文」样式")
    if not dirty:
        return []
    return [f"默认格式没有清干净（{'、'.join(dirty)}），里面的加粗、底色或间距可能透出来"]


def check_filename(filename: str, info_line: str, rules: Rules) -> list[str]:
    """第八条：文件名是「学号-姓名.docx」，而且和姓名行说的是同一个人。"""
    path = Path(filename)
    student_id, separator, name = path.stem.partition("-")
    if path.suffix != ".docx" or not separator or not student_id.isdigit() or not name:
        return [f"第八条：文件名「{filename}」不是「学号-姓名.docx」的写法"]
    if rules.info_line(name, student_id) != info_line:
        return [f"第八条：文件名「{filename}」和姓名行「{info_line}」对不上"]
    return []


# ---------- 读出一段实际写着的格式 ----------


def _attr(element, name: str):
    return element.get(qn(name)) if element is not None else None


def _child(parent, tag: str):
    return parent.find(qn(tag)) if parent is not None else None


def _is_on(rpr, tag: str) -> bool:
    element = _child(rpr, tag)
    return element is not None and _attr(element, "w:val") not in ("0", "false", "off", "none")


def _run_text(run) -> str:
    return "".join(t.text or "" for t in run.iter(qn("w:t")))


def _excerpt(text: str) -> str:
    text = text.strip(EDGE_WHITESPACE)
    return text[:EXCERPT_LENGTH] + ("…" if len(text) > EXCERPT_LENGTH else "")


def _paragraph_problems(paragraph, style: TextStyle, line: int, indent: bool = True) -> list[str]:
    ppr = _child(paragraph, "w:pPr")
    spacing, ind, jc = _child(ppr, "w:spacing"), _child(ppr, "w:ind"), _child(ppr, "w:jc")
    problems = []
    if _child(ppr, "w:pStyle") is not None:
        problems.append("还带着 Word 的样式")
    if _attr(jc, "w:val") != style.alignment:
        actual = ALIGNMENT_NAMES.get(_attr(jc, "w:val"), _attr(jc, "w:val"))
        problems.append(f"对齐应为{ALIGNMENT_NAMES[style.alignment]}，实际是{actual}")
    if (_attr(spacing, "w:line"), _attr(spacing, "w:lineRule")) != (str(line), "auto"):
        problems.append(f"行距应为 {line / SINGLE_LINE:g} 倍，实际不是")
    if (_attr(spacing, "w:before"), _attr(spacing, "w:after")) != ("0", "0"):
        problems.append("段前或段后有多余的间距")
    if indent:
        expected = str(round(style.first_line_chars * style.size_pt * 20))
        expected_chars = str(style.first_line_chars * 100)
        if (_attr(ind, "w:firstLine"), _attr(ind, "w:firstLineChars")) != (expected, expected_chars):
            want = f"缩进 {style.first_line_chars} 个字" if style.first_line_chars else "不缩进"
            problems.append(f"首行应该{want}，实际不是")
        if _attr(ind, "w:left") != "0" or _attr(ind, "w:hanging") is not None:
            problems.append("整段有多余的缩进")
    if _child(ppr, "w:pBdr") is not None:
        problems.append("带着边框线")
    problems += _unexpected(ppr, ALLOWED_PARAGRAPH_PROPERTIES, "段落", already_reported=("w:pStyle", "w:pBdr"))
    for element, allowed in ((spacing, ALLOWED_SPACING), (ind, ALLOWED_INDENT)):
        extra = _extra_attributes(element, allowed)
        if extra:
            problems.append(f"间距或缩进里有多余的设置（{'、'.join(extra)}）")
    return problems


# 白名单：排好的文件里，段落和文字只允许带这些设置。别的任何设置（底纹、字符间距、
# 按行算的段前段后……）都算没清干净——不必一种一种去列"哪些不行"。
ALLOWED_PARAGRAPH_PROPERTIES = {
    "w:pageBreakBefore", "w:numPr", "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr",
}  # fmt: skip
ALLOWED_RUN_PROPERTIES = {
    "w:rFonts", "w:strike", "w:dstrike", "w:vanish", "w:color", "w:sz", "w:szCs", "w:vertAlign",
}  # fmt: skip
ALLOWED_SPACING = {"w:before", "w:after", "w:line", "w:lineRule"}
ALLOWED_INDENT = {"w:left", "w:right", "w:firstLine", "w:firstLineChars"}
ALLOWED_FONT_ATTRIBUTES = {"w:ascii", "w:hAnsi", "w:eastAsia", "w:cs", "w:hint"}
# 用符号字体打的字（√、α 之类），字体必须留着，不按正文字体查
SYMBOL_FONTS = {"symbol", "wingdings", "wingdings 2", "wingdings 3", "webdings", "mt extra", "marlett"}


def _short(tag: str) -> str:
    return tag.split("}")[-1]


def _unexpected(properties, allowed: set[str], kind: str, already_reported=()) -> list[str]:
    if properties is None:
        return []
    known = {qn(t) for t in allowed} | {qn(t) for t in already_reported}
    extra = [_short(child.tag) for child in properties if child.tag not in known]
    return [f"带着不该有的{kind}格式（{'、'.join(extra)}）"] if extra else []


def _extra_attributes(element, allowed: set[str]) -> list[str]:
    if element is None:
        return []
    known = {qn(a) for a in allowed}
    return [_short(name) for name in element.attrib if name not in known]


def _is_symbol_run(run, fonts) -> bool:
    names = [(_attr(fonts, a) or "").lower() for a in ("w:ascii", "w:hAnsi")]
    private_use = any("" <= c <= "" for c in _run_text(run))
    return private_use or any(n in SYMBOL_FONTS for n in names)


def _run_problems(run, style: TextStyle) -> list[str]:
    rpr = _child(run, "w:rPr")
    fonts, size = _child(rpr, "w:rFonts"), _child(rpr, "w:sz")
    problems = _unexpected(
        rpr, ALLOWED_RUN_PROPERTIES, "文字",
        already_reported=("w:b", "w:bCs", "w:i", "w:iCs", "w:u", "w:highlight", "w:rStyle"),
    )  # fmt: skip
    if _extra_attributes(fonts, ALLOWED_FONT_ATTRIBUTES):
        problems.append("字体被文档主题的字体设置盖住了")
    if _is_symbol_run(run, fonts):
        pass
    elif _attr(fonts, "w:eastAsia") != style.east_asian_font:
        problems.append(f"中文字体应为{style.east_asian_font}，实际是{_attr(fonts, 'w:eastAsia') or '没有设置'}")
    elif (_attr(fonts, "w:ascii"), _attr(fonts, "w:hAnsi")) != (style.western_font,) * 2:
        problems.append(f"英文数字的字体应为 {style.western_font}，实际是 {_attr(fonts, 'w:ascii') or '没有设置'}")
    actual_size = int(_attr(size, "w:val")) / 2 if size is not None else None
    if actual_size != style.size_pt:
        problems.append(f"字号应为 {style.size_pt:g} 磅，实际是 {actual_size if actual_size else '没有设置'}")
    if _attr(_child(rpr, "w:color"), "w:val") != BLACK:
        problems.append("字不是黑色")
    for tag, label in (("w:b", "加粗"), ("w:i", "斜体"), ("w:u", "下划线")):
        if _is_on(rpr, tag):
            problems.append(f"有{label}")
    if _attr(_child(rpr, "w:highlight"), "w:val") not in (None, "none"):
        problems.append("有高亮")
    if _child(rpr, "w:rStyle") is not None:
        problems.append("还带着 Word 的字符样式")
    return problems


def _style_problems(paragraph, style: TextStyle, line: int, label: str, indent: bool = True) -> list[str]:
    found = _paragraph_problems(paragraph, style, line, indent)
    for run in paragraph.iter(qn("w:r")):
        if _run_text(run):
            for problem in _run_problems(run, style):
                if problem not in found:
                    found.append(problem)
    return [f"{label}：{problem}" for problem in found]


def _label(rule: str, kind: str, block: Block) -> str:
    return f"{rule}：{kind}「{_excerpt(block.text)}」"


# ---------- 逐条核对 ----------


def _check_page(document, rules: Rules) -> list[str]:
    problems = []
    expected_size = (str(rules.page_width_twips), str(rules.page_height_twips))
    vertical, horizontal = str(rules.margin_top_bottom_twips), str(rules.margin_left_right_twips)
    for section in document.element.body.iter(qn("w:sectPr")):
        size, margin = _child(section, "w:pgSz"), _child(section, "w:pgMar")
        if (_attr(size, "w:w"), _attr(size, "w:h")) != expected_size or _attr(size, "w:orient") == "landscape":
            problems.append("第一条：纸张不是 A4 竖向")
        if (_attr(margin, "w:top"), _attr(margin, "w:bottom")) != (vertical, vertical):
            problems.append(f"第一条：上下页边距不是 {rules.margin_top_bottom_cm:g} 厘米")
        if (_attr(margin, "w:left"), _attr(margin, "w:right")) != (horizontal, horizontal):
            problems.append(f"第一条：左右页边距不是 {rules.margin_left_right_cm:g} 厘米")
        if _attr(margin, "w:gutter") not in (None, "0"):
            problems.append("第一条：页面多留了装订线的位置")
        if _child(section, "w:titlePg") is not None:
            problems.append("第七条：设了「首页不同」，第一页可能没有页码")
        numbering = _child(section, "w:pgNumType")
        if numbering is not None and (
            _attr(numbering, "w:fmt") not in (None, "decimal") or _attr(numbering, "w:start") is not None
        ):
            problems.append("第七条：页码不是从 1 开始的阿拉伯数字")
    return list(dict.fromkeys(problems)) or ([] if document.sections else ["第一条：找不到页面设置"])


def _check_title(blocks: list[Block], rules: Rules, expect_title: bool) -> list[str]:
    titles = [b for b in blocks if b.role is Role.TITLE]
    if not titles:
        return ["第二条：没有题目"] if expect_title else []
    problems = []
    for title in titles:
        problems += _style_problems(
            title.element, rules.title, rules.line_spacing_value, _label("第二条", "题目", title)
        )
    return problems


def _check_info(blocks: list[Block], rules: Rules, info_line: str | None) -> list[str]:
    infos = [b for b in blocks if b.role is Role.INFO]
    problems = []
    for info in infos:
        problems += _style_problems(
            info.element, rules.info, rules.line_spacing_value, _label("第三条", "姓名行", info)
        )
    if info_line is None:
        return problems
    if [b.text for b in infos] != [info_line]:
        actual = "、".join(f"「{b.text}」" for b in infos) or "没有"
        return problems + [f"第三条：姓名行应为「{info_line}」，实际是{actual}"]
    visible = [b for b in blocks if b.role is not Role.EMPTY]
    position = visible.index(infos[0])
    has_title = any(b.role is Role.TITLE for b in blocks)
    if has_title and (position == 0 or visible[position - 1].role is not Role.TITLE):
        problems.append("第三条：姓名行不在题目下面一行")
    return problems


def _check_headings(blocks: list[Block], rules: Rules) -> list[str]:
    headings = [b for b in blocks if b.role is Role.HEADING]
    problems = []
    for heading in headings:
        numbered = _has_numbering(heading.element)
        problems += _style_problems(
            heading.element,
            rules.heading,
            rules.line_spacing_value,
            _label("第四条", "小标题", heading),
            indent=not numbered,
        )
    if rules.unify_heading_numbers and not any(_has_numbering(h.element) for h in headings):
        problems += _check_heading_numbers([h.text.strip() for h in headings], rules)
    return problems


def _check_heading_numbers(headings: list[str], rules: Rules) -> list[str]:
    """老师追加的要求：编号是「一、」「二、」「三、」，从一开始连着排；参考资料等不编号。"""
    count = 0
    for heading in headings:
        style, core = split_number(heading)
        if core.rstrip("：:").strip() in rules.unnumbered_headings:
            if style is not None:
                return [f"小标题编号：「{heading}」不该有编号"]
            continue
        count += 1
        expected = f"{chinese_numeral(count)}、"
        if not heading.startswith(expected) or split_number(heading[len(expected) :])[0] is not None:
            return [f"小标题编号：「{_excerpt(heading)}」的编号应为「{expected}」"]
    return []


def _has_numbering(paragraph) -> bool:
    numbering = paragraph.find(qn("w:pPr") + "/" + qn("w:numPr"))
    return numbering is not None and _attr(_child(numbering, "w:numId"), "w:val") not in (None, "0")


def _check_body(blocks: list[Block], rules: Rules) -> list[str]:
    problems = []
    for block in blocks:
        element = block.element
        if block.role is Role.EMPTY:
            structural = has_page_break(element) or element.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None
            if not structural:
                problems.append("第五条：段落之间有空行")
        if block.role is not Role.BODY:
            continue
        label = _label("第五条", "正文", block)
        problems += _style_problems(
            element, rules.body, rules.line_spacing_value, label, indent=not _has_numbering(element)
        )
        if any(is_soft_break(br) for br in element.iter(qn("w:br"))):
            problems.append(f"{label}：里面有软回车")
        first = next((c for r in element.iter(qn("w:r")) for c in r if c.tag != qn("w:rPr")), None)
        starts_with_tab = first is not None and first.tag == qn("w:tab")
        if starts_with_tab or block.text != block.text.strip(EDGE_WHITESPACE):
            problems.append(f"{label}：开头或结尾有手敲的空格或 Tab")
    return list(dict.fromkeys(problems))


def _check_tables_and_pictures(document, blocks: list[Block], rules: Rules) -> list[str]:
    problems = []
    for index, table in enumerate(document.tables, start=1):
        element = table._tbl
        if _attr(_child(element.tblPr, "w:jc"), "w:val") != "center":
            problems.append(f"第六条：第 {index} 个表格没有居中")
        grid = element.find(qn("w:tblGrid"))
        widths = [int(c.get(qn("w:w"), 0)) for c in grid.findall(qn("w:gridCol"))] if grid is not None else []
        if sum(widths) > rules.text_width_twips:
            problems.append(f"第六条：第 {index} 个表格比页面的版心宽")
        for paragraph in element.iter(qn("w:p")):
            found = _style_problems(
                paragraph, rules.table_text, rules.line_spacing_value, f"第六条：第 {index} 个表格里的字"
            )
            problems += [p for p in found if p not in problems]

    for block in blocks:
        if block.role is Role.IMAGE:
            ppr = _child(block.element, "w:pPr")
            centered = _attr(_child(ppr, "w:jc"), "w:val") == "center"
            indented = _attr(_child(ppr, "w:ind"), "w:firstLine") not in (None, "0")
            if not centered or indented:
                problems.append("第六条：图片没有居中")
    text_width = rules.text_width_twips * EMU_PER_TWIP
    for extent in document.element.body.iter(_WP + "extent"):
        if extent.getparent().tag == _WP + "inline" and int(extent.get("cx", 0)) > text_width:
            problems.append("第六条：图片比页面的版心宽")
    return list(dict.fromkeys(problems))


def _check_page_number(document, rules: Rules) -> list[str]:
    problems = []
    section = document.sections[0]
    paragraphs = [] if section.footer.is_linked_to_previous else list(section.footer._element.iter(qn("w:p")))
    with_field = [
        p
        for p in paragraphs
        if "".join(t.text or "" for t in p.iter(qn("w:instrText"))).strip() == "PAGE"
        and [f.get(qn("w:fldCharType")) for f in p.iter(qn("w:fldChar"))] == ["begin", "separate", "end"]
    ]
    if not with_field:
        problems.append("第七条：页面底端没有页码")
    elif _attr(_child(_child(with_field[0], "w:pPr"), "w:jc"), "w:val") != rules.page_number.alignment:
        problems.append("第七条：页码没有居中")
    else:
        expected = str(round(rules.page_number.size_pt * 2))
        sizes = {_attr(_child(_child(r, "w:rPr"), "w:sz"), "w:val") for r in with_field[0].iter(qn("w:r"))}
        plain = all(
            not _unexpected(_child(r, "w:rPr"), ALLOWED_RUN_PROPERTIES, "文字")
            and _attr(_child(_child(r, "w:rPr"), "w:color"), "w:val") == BLACK
            for r in with_field[0].iter(qn("w:r"))
        )
        if sizes != {expected} or not plain:
            problems.append("第七条：页码的字号或样子不对")
    # 页码只应该有一处：第一节的普通页脚。别处另有页脚，那些页上的页码就没了或不一样
    own_footers = sum(
        not part.is_linked_to_previous
        for part_section in document.sections
        for part in (part_section.footer, part_section.first_page_footer, part_section.even_page_footer)
    )
    if own_footers > 1:
        problems.append("第七条：首页、偶数页或后面的节另有自己的页脚")
    if document.settings.odd_and_even_pages_header_footer:
        problems.append("第七条：设了「奇偶页不同」，偶数页可能没有页码")
    other_text = [p for p in paragraphs if p not in with_field and "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()]
    if other_text:
        problems.append("第七条：页脚里除了页码还有别的字")

    for part in document.sections:
        for header in (part.header, part.first_page_header, part.even_page_header):
            # 没有自己页眉的节，python-docx 一访问就会新建一个；先判断再读
            if not header.is_linked_to_previous and any(p.text.strip() for p in header.paragraphs):
                problems.append("页眉：页眉里还有字，应该是空的")
    return list(dict.fromkeys(problems))
