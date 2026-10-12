"""把处理结果写成给老师看的 Excel 表。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .decisions import Decisions
from .inventory import Inventory, Outcome, Record
from .scanner import FileKind
from .roster import Roster

OVERVIEW_SHEET = "总览"
DETAIL_SHEET = "明细"

# 要老师处理的排在前面
OUTCOME_ORDER = [
    Outcome.REVIEW,
    Outcome.UNPROCESSED,
    Outcome.DONE,
    Outcome.READY,
    Outcome.IGNORED,
]
CHANGES_EXPLAINED = (
    "所有已完成的作业都统一了纸张、页边距、字体、字号、行距、缩进，并加了页码。"
    "「做了哪些改动」一列只写除此以外动过的地方。"
)
OUTCOME_FILL = {
    Outcome.REVIEW: PatternFill("solid", fgColor="FFF2CC"),
    Outcome.UNPROCESSED: PatternFill("solid", fgColor="F8CBAD"),
    Outcome.IGNORED: PatternFill("solid", fgColor="EDEDED"),
}
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
BAD_FONT = Font(bold=True, color="C00000")

DETAIL_COLUMNS = [
    ("原文件名", 36),
    ("类型", 18),
    ("去向", 12),
    ("姓名", 10),
    ("学号", 12),
    ("姓名来自", 10),
    ("学号来自", 10),
    ("新文件名", 26),
    ("认出的题目", 34),
    ("认出的小标题", 44),
    ("做了哪些改动", 60),
    ("质检", 10),
    ("说明", 90),
]


def _detail_row(record: Record) -> list:
    return [
        record.source.path.name,
        record.source.kind.value,
        record.outcome.value,
        record.name,
        record.student_id,
        record.name_source.value if record.name_source else None,
        record.id_source.value if record.id_source else None,
        # 点名阶段显示打算用的名字；排版之后只显示真正写出去的文件
        record.output_name or (record.new_filename if record.outcome is Outcome.READY else None),
        record.structure.title if record.structure else None,
        " | ".join(record.structure.headings) or None if record.structure else None,
        "\n".join(record.changes) or None,
        None if record.quality is None else ("没通过" if record.quality else "通过"),
        "\n".join(record.notes) or None,
    ]


def _sorted(records: list[Record]) -> list[Record]:
    return sorted(
        records,
        key=lambda r: (OUTCOME_ORDER.index(r.outcome), r.student_id or "~", r.source.path.name),
    )


def _quality_summary(inventory: Inventory) -> str:
    done = inventory.count(Outcome.DONE)
    failed = sum(1 for r in inventory.records if r.quality)
    text = f"已完成的 {done} 份全部通过：重新打开每一份，逐条核对了学校的 8 条要求和小标题编号"
    if failed:
        text += f"。另有 {failed} 份没通过，没有放进「已完成」，已转入「需要您确认」"
    return text


def _people(pairs, possible_files=None) -> str:
    possible_files = possible_files or {}
    people = []
    for student_id, name in pairs:
        files = possible_files.get(student_id)
        hint = f"，可能是「{'」「'.join(files)}」，待确认" if files else ""
        people.append(f"{name}（{student_id}{hint}）")
    return "、".join(people) or "无"


MAX_CELL_LENGTH = 30000


def _as_text(sheet) -> None:
    """表里的字一律按"文字"存。

    题目、小标题、文件名都是学生写的。Excel 会把等号开头的内容当成公式去算，
    学生把题目写成「=1+1」或者一个超链接公式，就会在老师的报告里生效。
    """
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"


def _clean(value):
    """去掉 Excel 不接受的控制字符；太长的截断，免得被悄悄砍掉。"""
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(value) > MAX_CELL_LENGTH:
        value = value[:MAX_CELL_LENGTH] + "……（太长，后面的省略了）"
    return value


def write_report(
    path: Path,
    inventory: Inventory,
    *,
    input_dir: Path,
    roster: Roster | None,
    originals_untouched: bool,
    file_count: int,
    decisions: Decisions | None = None,
    incomplete: str | None = None,
) -> Path:
    workbook = openpyxl.Workbook()
    _write_overview(
        workbook.active, inventory, input_dir, roster, originals_untouched, file_count, decisions,
        incomplete,
    )
    _write_detail(workbook.create_sheet(DETAIL_SHEET), inventory)
    for sheet in workbook.worksheets:
        _as_text(sheet)
        # 打印时横向、所有列挤进一页宽，不然一行会被拆到好几张纸上
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(str(path))
    return path


def _write_overview(
    sheet, inventory, input_dir, roster, originals_untouched, file_count, decisions, incomplete=None
) -> None:
    sheet.title = OVERVIEW_SHEET
    if originals_untouched:
        originals = f"{file_count} 个原件，运行前后完全一致"
        if any(r.source.kind is FileKind.TEMP for r in inventory.records):
            originals += "（Word 的临时文件不算原件）"
    else:
        originals = "不一致！运行期间「学生作业」里的文件发生了变化，请先不要使用这次的结果"

    rows: list[tuple[str, object]] = [
        ("运行时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("作业文件夹", str(input_dir)),
        ("文件总数", len(inventory.records)),
        *[
            (outcome.value, inventory.count(outcome))
            for outcome in Outcome
            # "已完成"和"可以排版"是同一批作业排版前后的两种叫法，没有的那种不显示
            if inventory.count(outcome) or outcome not in (Outcome.DONE, Outcome.READY)
        ],
        ("原件核对", originals),
    ]
    if incomplete:
        rows.insert(
            0,
            ("这次运行没有做完", f"中途出了意外：{incomplete}。下面的结果不全，请不要当成全部；可以重新运行一次。"),
        )
    if inventory.count(Outcome.DONE):
        rows.append(("质检", _quality_summary(inventory)))
        rows.append(("关于改动", CHANGES_EXPLAINED))
    if roster is None:
        rows.append(("班级名单", "没有提供，缺的学号无法补全，也无法知道谁没交"))
    else:
        # 同一位学生另有一份已经排好的，就不算"没法处理"
        written = {r.student_id for r in inventory.records if r.output_name}
        unusable = [
            (r.student_id, r.name)
            for r in _sorted(inventory.records)
            if r.outcome is Outcome.UNPROCESSED and r.student_id and r.name
            and r.student_id not in written
        ]
        unusable = list(dict.fromkeys(unusable))
        rows += [
            ("班级名单", f"{len(roster)} 人"),
            (
                "名单上有、但没找到作业的学生",
                _people(inventory.missing_students, inventory.possible_files),
            ),
            ("交了但没法处理的学生", _people(unusable)),
        ]
    if decisions is not None:
        rows.append(("老师的决定", f"{len(decisions.filenames())} 条，已照办"))
    if inventory.unmatched_decisions:
        rows.append(
            ("决定表里写了、但文件夹里没有的文件", "、".join(inventory.unmatched_decisions))
        )
    if inventory.skipped_folders:
        rows.append(("没有进入的子文件夹", "、".join(inventory.skipped_folders)))

    for label, value in rows:
        sheet.append([label, _clean(value)])
    for row in sheet.iter_rows():
        row[0].font = Font(bold=True)
        row[0].fill = HEADER_FILL
        row[1].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if (row[0].value == "原件核对" and not originals_untouched) or row[0].value == "这次运行没有做完":
            row[1].font = BAD_FONT
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 80


def _write_detail(sheet, inventory: Inventory) -> None:
    sheet.append([title for title, _ in DETAIL_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for index, (_, width) in enumerate(DETAIL_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for record in _sorted(inventory.records):
        sheet.append([_clean(value) for value in _detail_row(record)])
        fill = OUTCOME_FILL.get(record.outcome)
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill

    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
