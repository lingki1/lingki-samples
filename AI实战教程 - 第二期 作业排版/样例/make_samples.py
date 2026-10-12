# 造一批格式五花八门的学生作业（全是样例：学校、班级、姓名、内容都是编的）
# 用法：python make_samples.py   → 在这个脚本旁边新建 project-seed/（里面是 学生作业/ 和 格式要求.docx），并重写旁边的 truth.json
# 需要 python-docx、Pillow，以及装在默认位置的 LibreOffice（两份 .doc 和一份 PDF 是用它转出来的）。
# 种子固定。标准答案 truth.json 放在 project-seed 外面，不进那场对话的文件夹。
# 放进仓库时只改了这几行说明，别的是造样例时用的原样。
import os, io, json, random, shutil, subprocess, hashlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_COLOR_INDEX
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__)); SEED = os.path.join(HERE, 'project-seed'); HW = os.path.join(SEED, '学生作业')
SOFFICE = r'C:\Program Files\LibreOffice\program\soffice.exe'
R = random.Random(25)
CLASS = '高二（3）班'

NAMES = '林知夏 周子衿 陈一诺 许嘉树 沈星河 顾念 唐小满 白鹿鸣 江晚晴 程予安 夏以宁 宋清和 叶知秋 方可 苏禾 陆景明 温言 何其多 秦朗 孟想 童年 乔木 阮青 梁思齐 安然 柏舟 邱实 谷雨 殷实 骆一鸣 关山月 齐悦 尚可 祝余 池鱼 景行'.split()
assert len(NAMES) == 36 and len(set(NAMES)) == 36

# 六个题目，每个三节，每节两段。{a}{b} 换成每个学生自己的数，免得全班一模一样。
TOPICS = [
 ('食堂排队太久，我想这样解决', [
  ('我发现的问题', ['每天中午十二点下课，全校的人几乎同时往食堂走。我连着记了五天，从排到队尾到端上饭，最短要 {a} 分钟，最长的一次等了 {b} 分钟。',
                 '排队的时间一长，午休就只剩下不到半小时，下午第一节课很多人都在打瞌睡。']),
  ('原因分析', ['我在队伍旁边数过，一个窗口打一份饭大约要四十秒，其中一半的时间花在刷卡和找零上。',
              '另外，三个年级同时下课，六个窗口里常常只开四个，人一多就全堵在门口。']),
  ('我的办法', ['第一，三个年级错开十分钟下课，高一先走，高三最后。第二，把刷卡的机器挪到队伍中间，排到窗口之前先把卡刷好。',
              '我把这个想法讲给食堂的师傅听，他说错峰以前试过，效果不错，只是后来没有坚持下来。'])]),
 ('教室的灯，白天其实可以少开一半', [
  ('我发现的问题', ['我们教室一共有 12 盏 LED 灯，从早上七点一直开到晚自习结束。可是靠窗的那两排，晴天的时候根本用不着开灯。',
                 '我数了 {a} 天，其中有 {b} 天是晴天，这些天里靠窗的六盏灯等于白白亮了一整天。']),
  ('原因分析', ['开关在教室门口，一共只有两个，一个管前半间，一个管后半间，没有办法只关靠窗的那一边。',
              '还有一个原因是习惯。第一个进教室的人顺手把两个开关都按下去，后面就再也没有人去管它了。']),
  ('我的办法', ['最省事的办法是请电工师傅改一下线路，让一个开关管靠窗的两排，另一个管靠墙的两排。',
              '线路改好之前，可以先在开关旁边贴一张提示，再排一个值日生，每天上午第二节课后看一眼天气，决定关不关。'])]),
 ('图书角的书，为什么总是有借无还', [
  ('我发现的问题', ['开学的时候，班级图书角有 {a} 本书，到了期中只剩下 {b} 本。少掉的书不是丢了，而是借走以后一直没有人还。',
                 '想看的同学借不到，最后大家干脆不去图书角了，那个书架现在成了放水杯的地方。']),
  ('原因分析', ['借书没有登记。谁拿走了哪一本，什么时候拿的，没有人说得清，想催也不知道该催谁。',
              '另外，书上没有任何记号，带回家以后很容易和自己的书混在一起，时间一长连借书的人自己都忘了。']),
  ('我的办法', ['我用 Excel 做了一张借书登记表，打印出来挂在书架旁边，借书的时候写上书名、姓名和日期。',
              '每周五放学前，图书管理员对一次表，超过两周没有还的，在班级群里提醒一声。'])]),
 ('自行车棚为什么总是停不下', [
  ('我发现的问题', ['学校的自行车棚画了大约 {a} 个车位，可是每天早上七点半以后来的人，基本上只能把车停在棚子外面的路边。',
                 '我在上周三早上数过，棚子里实际停了 {b} 辆车，中间却还空着好几块地方，只是谁也推不进去。']),
  ('原因分析', ['先到的同学喜欢把车停在离出口最近的位置，而且是横着停，一辆车占了两辆车的地方。',
              '棚子里还有十几辆落满灰的旧车，听保安叔叔说是往届毕业生留下的，一直没有人处理。']),
  ('我的办法', ['第一，在地上重新画斜向的停车线，斜着停比横着停省地方。第二，请学校贴一张通知，一个月以后没有人认领的旧车统一清理。',
              '另外可以按年级分区，高三的同学放学最晚，把离出口最近的那一片留给他们。'])]),
 ('课间十分钟，饮水机前的长队', [
  ('我发现的问题', ['我们这层楼有四个班，将近 200 个人，共用走廊尽头的两台饮水机。课间只有十分钟，排在后面的人常常接不到水就打铃了。',
                 '我记了 {a} 个课间，其中有 {b} 个课间我没有接到水。']),
  ('原因分析', ['饮水机烧水的速度跟不上。连续接上七八杯以后，出来的水就是温的，后面的人只好站在那里等它重新烧开。',
              '还有不少同学用的是一升以上的大水壶，一个人接满就要将近一分钟。']),
  ('我的办法', ['大课间和午休的时候先把水接满，不要都挤到课间十分钟。这一条不花钱，马上就能做。',
              '长期的办法是向学校申请在楼层的另一头再加一台。我问过总务处的老师，他说可以把同学们的意见写成书面材料交上去。'])]),
 ('值日表上，为什么总有人忘', [
  ('我发现的问题', ['我们班的值日表贴在后门旁边，每天安排 4 个人。这个学期到现在，我统计了 {a} 天，其中有 {b} 天至少有一个人忘了做值日。',
                 '忘了的人不是故意偷懒，被提醒以后都会马上去补，可是那时候往往已经被扣分了。']),
  ('原因分析', ['值日表是按学号排的，一轮下来要九天，隔得太久，很难记住哪一天轮到自己。',
              '表贴在教室最后面，平时没有人会特意走过去看。']),
  ('我的办法', ['把值日表改成按星期固定，比如每个星期二都是同一组人，这样就不用记日期了。',
              '每天放学前，由劳动委员在黑板的右下角写上第二天值日的四个名字，抬头就能看见。'])]),
]
REFS = ['学校总务处. 2024 年春季学期后勤工作小结（校内资料）.', '班主任提供的本班考勤与扣分记录, 2024.', '对食堂、保安室工作人员的口头采访记录, 2024 年 4 月.', '本人的现场观察记录表（见附页）.']
CN = '一二三四五六'

def fonts(run, east, latin=None, size=None, bold=None, color=None, underline=None, hl=None, italic=None):
    latin = latin or east
    run.font.name = latin
    rf = run._element.get_or_add_rPr().find(qn('w:rFonts'))
    rf.set(qn('w:eastAsia'), east); rf.set(qn('w:hAnsi'), latin)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor(*color)
    if underline: run.font.underline = True
    if hl: run.font.highlight_color = hl

def fmt(p, align=None, spacing=None, first=None, before=None, after=None):
    f = p.paragraph_format
    if align is not None: p.alignment = align
    if spacing is not None:
        if isinstance(spacing, tuple): f.line_spacing_rule = WD_LINE_SPACING.EXACTLY; f.line_spacing = Pt(spacing[1])
        else: f.line_spacing = spacing
    if first is not None: f.first_line_indent = Pt(first)
    if before is not None: f.space_before = Pt(before)
    if after is not None: f.space_after = Pt(after)

BODY_LOOKS = [  # (中文字体, 西文字体, 字号, 行距, 对齐, 颜色)
 ('宋体', 'Times New Roman', 12, 1.5, 'j', None), ('微软雅黑', 'Arial', 11, 1.0, 'l', None), ('等线', 'Calibri', 10.5, 1.15, 'l', None),
 ('楷体', 'Times New Roman', 14, 2.0, 'l', None), ('仿宋', 'Times New Roman', 16, ('exact', 28), 'j', None), ('宋体', 'Calibri', 10.5, 1.0, 'l', None),
 ('微软雅黑', 'Arial', 12, 1.5, 'l', (0x1F, 0x4E, 0x79)), ('黑体', 'Arial', 12, 1.25, 'l', None), ('华文楷体', 'Times New Roman', 13, 1.5, 'j', None),
 ('等线', 'Calibri', 12, 1.0, 'c', None),
]
INDENTS = ['proper', 'proper', 'none', 'none', 'full', 'full', 'half', 'tab']
HEAD_NUM = ['cn', 'cn', 'cn', 'num', 'paren', 'cn', 'num', 'none']
HEAD_LOOK = ['bold', 'bold', 'style', 'plain', 'bold', 'style', 'big', 'plain']
INFO = ['a', 'b', 'c', 'd', 'a', 'b', 'c', 'a']
MARGINS = [None, None, None, (1.27, 1.27), (2.0, 2.0), None, (3.0, 2.5), None]
ALIGN = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'j': WD_ALIGN_PARAGRAPH.JUSTIFY, 'c': WD_ALIGN_PARAGRAPH.CENTER}

def head_text(i, h, style):
    return {'cn': f'{CN[i]}、{h}', 'num': f'{i + 1}. {h}', 'paren': f'（{CN[i]}）{h}', 'none': h}[style]

def info_text(kind, name, sid):
    return {'a': f'{CLASS}  {name}  {sid}', 'b': f'姓名：{name}    学号：{sid}', 'c': f'班级：高二3班 姓名:{name} 学号:{sid}', 'd': f'{name} {sid}'}[kind]

def chart_png(path):
    from PIL import Image, ImageDraw
    im = Image.new('RGB', (640, 300), 'white'); d = ImageDraw.Draw(im)
    for i, v in enumerate([120, 200, 160, 90, 230]):
        d.rectangle([60 + i * 110, 270 - v, 130 + i * 110, 270], fill=(90, 120, 170))
    d.line([40, 270, 620, 270], fill='black', width=2); im.save(path)

def build(st, path):
    """按这个学生的习惯写一份作业；返回标准答案里的段落表。"""
    d = Document(); sec = d.sections[0]; tr = st['traps']; lk = st['look']; east, latin, size, spacing, al, color = lk
    if st['margins']: sec.left_margin = sec.right_margin = Cm(st['margins'][0]); sec.top_margin = sec.bottom_margin = Cm(st['margins'][1])
    if 'letter' in tr: sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    if 'landscape' in tr: sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
    if 'style-level' in tr:  # 字体设在「正文」样式上，段落里的字不单独设
        n = d.styles['Normal']; n.font.name = latin; n.font.size = Pt(size); n.element.rPr.rFonts.set(qn('w:eastAsia'), east)
    paras = []
    def body_run(p, text, **kw):
        r = p.add_run(text)
        if 'style-level' not in tr: fonts(r, east, latin, size, color=color, **kw)
        return r
    def blank():
        if 'blank-lines' in tr: d.add_paragraph()
    title, sections = TOPICS[st['topic']]
    if 'empty' in tr: d.save(path); return paras
    # 题目
    if 'no-title' not in tr:
        tl = st['title_look']
        if tl == 'style': p = d.add_paragraph(title, style='Title')
        else:
            p = d.add_paragraph(); r = p.add_run(title)
            if tl == 'big': fonts(r, '黑体', 'Arial', 22, bold=True); fmt(p, ALIGN['c'])
            elif tl == 'red': fonts(r, '微软雅黑', 'Arial', 20, bold=True, color=(0xC0, 0, 0), underline=True); fmt(p, ALIGN['c'])
            elif tl == 'left': fonts(r, east, latin, size + 2, bold=True); fmt(p, ALIGN['l'])
            else: fonts(r, east, latin, size)  # 'plain'：跟正文长得一样
        paras.append({'role': 'title', 'text': title}); blank()
    # 姓名学号
    ik = st['info']
    if ik in 'abcd':
        p = d.add_paragraph(); body_run(p, info_text(ik, st['name'], st['id'])); fmt(p, ALIGN['c' if st['title_look'] in ('big', 'red', 'style') else 'l'])
        paras.append({'role': 'info', 'text': info_text(ik, st['name'], st['id'])}); blank()
    elif ik == 'header':
        hp = sec.header.paragraphs[0]; hp.text = info_text('a', st['name'], st['id'])
    # 三节
    for i, (h, ps) in enumerate(sections):
        ht = head_text(i, h, st['head_num']); hl = st['head_look']
        if hl == 'style': p = d.add_paragraph(ht, style='Heading 2')
        else:
            p = d.add_paragraph(); r = p.add_run(ht)
            if 'style-level' not in tr or hl != 'plain': fonts(r, east, latin, size + (4 if hl == 'big' else 0), bold=(hl in ('bold', 'big')) or None, color=color)
        paras.append({'role': 'heading', 'text': ht})
        A, B = (st['a'], st['b']) if st['topic'] == 0 else (st['b'], st['a'])   # 只有第一个题目是「最短 a、最长 b」，别的都是「一共 a、其中 b」
        texts = [t.format(a=A, b=B) for t in ps]
        pre = {'full': '\u3000\u3000', 'half': '    ', 'tab': '\t'}.get(st['indent'], '')
        if 'soft-break' in tr:  # 一节只有一个段落，里面用换行（Shift+回车）隔开
            p = d.add_paragraph()
            for k, t in enumerate(texts):
                r = body_run(p, pre + t)
                if k < len(texts) - 1: r.add_break()
            fmt(p, ALIGN[al], spacing)
            paras += [{'role': 'body', 'text': t, 'soft': True} for t in texts]
        else:
            for k, t in enumerate(texts):
                p = d.add_paragraph()
                if 'emphasis' in tr and i == 2 and k == 0:  # 学生自己标了一句重点：加粗加黄底
                    cut = t.index('。') + 1; body_run(p, pre + t[:cut], bold=True, hl=WD_COLOR_INDEX.YELLOW); body_run(p, t[cut:])
                else: body_run(p, pre + t)
                fmt(p, ALIGN[al], spacing, first=(size * 2 if st['indent'] == 'proper' else None))
                paras.append({'role': 'body', 'text': t}); blank()
        if 'table' in tr and i == 0:
            rows = [('日期', '最短（分钟）', '最长（分钟）'), ('周一', '9', '21'), ('周二', '11', '18'), ('周三', '8', '24')]
            t = d.add_table(rows=len(rows), cols=3); t.style = 'Table Grid'
            for a, row in enumerate(rows):
                for b, v in enumerate(row): t.cell(a, b).text = v
            paras.append({'role': 'table', 'rows': rows})
        if 'image' in tr and i == 0:
            png = os.path.join(HERE, '_chart.png'); chart_png(png); d.add_picture(png, width=Cm(10)); os.remove(png)
            paras.append({'role': 'image'})
    if st['refs']:
        p = d.add_paragraph(); fonts(p.add_run('参考资料'), east, latin, size, bold=True); paras.append({'role': 'heading', 'text': '参考资料', 'ref_heading': True})
        for k, t in enumerate(st['refs']):
            p = d.add_paragraph(); body_run(p, f'[{k + 1}] {t}'); paras.append({'role': 'ref', 'text': f'[{k + 1}] {t}'})
    d.save(path); return paras

def convert(path, ext):
    subprocess.run([SOFFICE, '--headless', '--convert-to', ext, '--outdir', os.path.dirname(path), path], check=True, capture_output=True)
    os.remove(path); return os.path.splitext(path)[0] + '.' + ext

def spec_doc(path):
    d = Document()
    for t in ['研究性学习小报告 格式要求', '青岚中学 教务处', '', '一、页面：A4 纸，竖向。页边距上、下 2.54 厘米，左、右 3.18 厘米。',
              '二、题目：黑体，二号，居中，不加粗，不加下划线。', '三、班级、姓名、学号：写在题目下面一行，楷体，小四，居中。写法统一成「高二（3）班  姓名  学号」，中间各空两格。',
              '四、小标题：黑体，四号，顶格，不加粗。', '五、正文：宋体，小四；英文和数字用 Times New Roman。每段首行缩进两个字，1.5 倍行距，两端对齐。段与段之间不空行。全文黑色。',
              '六、表格和图片：保留，居中。', '七、页码：页面底端居中。', '八、文件名：学号-姓名.docx，例如 20240301-张三.docx。']:
        d.add_paragraph(t)
    d.save(path)

def main():
    if os.path.isdir(SEED): shutil.rmtree(SEED)
    os.makedirs(HW)
    sts = []
    for i, name in enumerate(NAMES):
        sts.append({'n': i + 1, 'name': name, 'id': str(20240301 + i), 'topic': i % 6, 'look': BODY_LOOKS[R.randrange(len(BODY_LOOKS))], 'indent': R.choice(INDENTS),
                    'head_num': R.choice(HEAD_NUM), 'head_look': R.choice(HEAD_LOOK), 'info': R.choice(INFO), 'margins': R.choice(MARGINS),
                    'title_look': R.choice(['big', 'big', 'left', 'plain', 'red', 'style', 'big']), 'a': R.randrange(8, 15), 'b': R.randrange(16, 30),
                    'refs': R.sample(REFS, 2) if R.random() < 0.3 else None, 'traps': [], 'ext': 'docx'})
    # 文件名的习惯
    pats = ['{name}.docx', '{id}{name}.docx', '{name}-研究性学习报告.docx', '高二3班 {name} 小报告.docx', '作业_{name}_{id}.docx', '{name}（最终版）.docx', '{name}改.docx', '{id}-{name}.docx', '研究性学习 {name}.docx', '{name}{id}.docx']
    for s in sts: s['file'] = R.choice(pats).format(name=s['name'], id=s['id'])
    # 指定的坑（按序号固定，方便对账）
    def trap(n, *t): sts[n - 1]['traps'] += list(t)
    trap(3, 'table'); trap(8, 'image'); trap(5, 'blank-lines'); trap(14, 'blank-lines'); trap(27, 'blank-lines')
    trap(11, 'soft-break'); trap(17, 'landscape'); trap(22, 'letter'); trap(9, 'emphasis'); trap(30, 'emphasis')
    trap(6, 'style-level'); trap(19, 'style-level'); trap(25, 'no-title'); trap(33, 'empty')
    sts[10]['indent'] = 'full'; sts[24]['info'] = 'a'; sts[32]['file'] = f"{sts[32]['name']}.docx"
    sts[1]['info'] = 'header'; sts[15]['info'] = 'header'                      # 姓名学号只写在页眉里
    sts[12]['info'] = 'none'; sts[12]['file'] = f"{sts[12]['name']}的报告.docx"  # 文档里没写，只有文件名里有名字
    sts[20]['info'] = 'none'; sts[20]['file'] = '新建 Microsoft Word 文档.docx'  # 哪儿都没写：只能人工认
    sts[28]['info'] = 'b'; sts[28]['file'] = '新建 Microsoft Word 文档 (2).docx'  # 文件名没名字，文档里有
    sts[18]['file'] = 'report_qinlang.docx'; sts[34]['file'] = '文档1.docx'
    sts[3]['head_num'] = 'none'; sts[3]['head_look'] = 'plain'                  # 小标题没编号也没加粗：最难认
    sts[6]['ext'] = 'doc'; sts[26]['ext'] = 'doc'; sts[31]['ext'] = 'pdf'       # 两份老格式、一份交成了 PDF
    for s in sts: s['file'] = os.path.splitext(s['file'])[0] + '.docx'
    assert len({s['file'] for s in sts}) == 36, '文件名撞了'
    truth = []
    for s in sts:
        path = os.path.join(HW, s['file']); paras = build(s, path)
        if s['ext'] != 'docx': path = convert(path, s['ext']); s['file'] = os.path.basename(path)
        truth.append({'file': s['file'], 'n': s['n'], 'name': s['name'], 'id': s['id'], 'kind': s['ext'], 'traps': s['traps'], 'info_where': s['info'], 'indent': s['indent'],
                      'head_num': s['head_num'], 'head_look': s['head_look'], 'title_look': s['title_look'], 'look': s['look'], 'margins': s['margins'], 'paras': paras})
    # 同一个人交了两次：第二份改了一段
    s = dict(sts[23]); first = s['file']; s['file'] = os.path.splitext(first)[0] + '(1).docx'; s['a'] += 1
    paras = build(s, os.path.join(HW, s['file']))
    truth.append({'file': s['file'], 'n': s['n'], 'name': s['name'], 'id': s['id'], 'kind': 'docx', 'traps': s['traps'] + ['resubmit'], 'dup_of': first, 'info_where': s['info'], 'indent': s['indent'],
                  'head_num': s['head_num'], 'head_look': s['head_look'], 'title_look': s['title_look'], 'look': s['look'], 'margins': s['margins'], 'paras': paras})
    # Word 没关时留下的临时文件（真文件夹里常有）
    lock = '~$' + sts[0]['file']   # Word 对短文件名是在前面加 ~$
    open(os.path.join(HW, lock), 'wb').write(b'\x05ling' + b'\x00' * 157)
    truth.append({'file': lock, 'kind': 'lock', 'traps': ['lock-file'], 'paras': []})
    spec_doc(os.path.join(SEED, '格式要求.docx'))
    for t in truth: t['sha256'] = hashlib.sha256(open(os.path.join(HW, t['file']), 'rb').read()).hexdigest()
    io.open(os.path.join(HERE, 'truth.json'), 'w', encoding='utf-8', newline='\n').write(json.dumps(truth, ensure_ascii=False, indent=1))
    kinds = {}
    for t in truth: kinds[t['kind']] = kinds.get(t['kind'], 0) + 1
    print('文件', len(truth), kinds, '| 段落', sum(len(t['paras']) for t in truth))
    for t in truth:
        if t['traps'] or t.get('info_where') in ('header', 'none') or t['kind'] != 'docx': print(f"  {t.get('n', '-'):>2} {t['file']:<34} {t['kind']:<5} 信息:{t.get('info_where', '-'):<6} {','.join(t['traps'])}")

if __name__ == '__main__': main()
